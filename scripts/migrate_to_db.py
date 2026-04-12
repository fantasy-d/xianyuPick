import json, pymysql, os, re
from pathlib import Path
from openpyxl import load_workbook

# --- 配置加载 ---
BASE_DIR = Path(__file__).resolve().parents[1]
CONFIG_PATH = BASE_DIR / "config" / "database.json"
config = json.load(open(CONFIG_PATH))
config["cursorclass"] = pymysql.cursors.DictCursor

def get_db_conn(): return pymysql.connect(**config)

def migrate():
    conn = get_db_conn()
    cursor = conn.cursor()
    
    cursor.execute("SELECT id, keyword, root_dir FROM tasks")
    tasks = cursor.fetchall()
    print(f"Found {len(tasks)} tasks to migrate.")

    for t in tasks:
        task_id = t['id']
        root_dir = Path(t['root_dir'])
        if not root_dir.is_absolute(): root_dir = BASE_DIR / root_dir
        
        # 路径兼容性尝试
        if not root_dir.exists():
            alt_dir = BASE_DIR / "outputs" / root_dir.name
            if alt_dir.exists(): root_dir = alt_dir

        xianyu_json = root_dir / "xianyu_hot_items.json"
        if not xianyu_json.exists(): continue

        print(f"\n>>> Migrating Task: {t['keyword']} ({task_id})")
        
        try:
            hot_items = json.loads(xianyu_json.read_text()).get("hot_items", [])
            cursor.execute("DELETE FROM xianyu_items WHERE task_id = %s", (task_id,))
            
            for i, item in enumerate(hot_items, start=1):
                cursor.execute("""
                    INSERT INTO xianyu_items (task_id, rank_index, title, price, image_url, want_count)
                    VALUES (%s, %s, %s, %s, %s, %s)
                """, (task_id, i, item.get('title'), item.get('price'), item.get('image_url'), item.get('want_count')))
                item_db_id = cursor.lastrowid
                
                # --- 修复：模糊匹配 Rank 文件夹 ---
                # 寻找以 Rank_{i}_ 开头的文件夹
                found_dirs = list(root_dir.glob(f"Rank_{i}_*"))
                if not found_dirs:
                    print(f"  [Warning] No Rank_{i} folder found in {root_dir}")
                    continue
                
                item_dir = found_dirs[0] # 取匹配到的第一个
                source_count = 0
                for f in item_dir.glob("*.xlsx"):
                    try:
                        wb = load_workbook(filename=f, read_only=True)
                        ws = wb.active
                        prices = [float(row[1]) for row in ws.iter_rows(min_row=2, max_col=2, values_only=True) if row[1]]
                        if prices:
                            offer_id = f.stem.split("_")[-1]
                            cursor.execute("""
                                INSERT INTO ali1688_sources (item_id, title, offer_id, min_price, sku_count, source_url)
                                VALUES (%s, %s, %s, %s, %s, %s)
                            """, (item_db_id, f.stem.replace(f"_{offer_id}", ""), offer_id, min(prices), len(prices), f"https://detail.1688.com/offer/{offer_id}.html"))
                            source_count += 1
                    except: continue
                print(f"  - Rank {i}: Imported {source_count} sources.")
            conn.commit()
        except Exception as e:
            print(f"  [Error] {task_id}: {e}")
            conn.rollback()

    conn.close()
    print("\nMigration completed!")

if __name__ == "__main__":
    migrate()
