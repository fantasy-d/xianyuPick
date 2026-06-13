import json, pymysql, os, re, csv
from pathlib import Path
from openpyxl import load_workbook

import sys
# --- 配置加载 ---
BASE_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BASE_DIR / "src"))
from xianyu_tools.config import settings

config = settings.get_database_config()
config["cursorclass"] = pymysql.cursors.DictCursor

def get_db_conn(): return pymysql.connect(**config)

def migrate():
    conn = get_db_conn()
    cursor = conn.cursor()
    
    cursor.execute("SELECT id, keyword, root_dir FROM tasks WHERE keyword = '人体工学椅'")
    tasks = cursor.fetchall()
    print(f"Targeting Task: 人体工学椅...")

    for t in tasks:
        task_id = t['id']
        # 强制修正路径：直接去找无空格的目录
        root_dir = BASE_DIR / "outputs" / "人体工学椅_20260411"
        
        if not root_dir.exists():
            print(f"Directory {root_dir} still not found. Checking all outputs...")
            # 暴力尝试：列出所有 outputs 下包含 '人体' 的目录
            dirs = [d for d in (BASE_DIR / "outputs").iterdir() if d.is_dir() and '人体' in d.name]
            if dirs: root_dir = dirs[0]
            else: continue

        xianyu_json = root_dir / "xianyu_hot_items.json"
        if not xianyu_json.exists(): continue

        print(f"\n>>> Re-Migrating Task: {t['keyword']} from {root_dir}")
        
        try:
            hot_items = json.loads(xianyu_json.read_text()).get("hot_items", [])
            # 找到数据库中的 xianyu_items ID 映射
            cursor.execute("SELECT id, rank_index FROM xianyu_items WHERE task_id = %s", (task_id,))
            db_items = {row['rank_index']: row['id'] for row in cursor.fetchall()}
            
            for i, item in enumerate(hot_items, start=1):
                item_db_id = db_items.get(i)
                if not item_db_id: continue
                
                # 模糊匹配 Rank 文件夹
                found_dirs = list(root_dir.glob(f"Rank_{i}_*"))
                if not found_dirs: continue
                
                item_dir = found_dirs[0]
                source_count = 0
                
                # 方案 A: 扫描 .xlsx
                for f in item_dir.glob("*.xlsx"):
                    try:
                        wb = load_workbook(filename=f, read_only=True)
                        ws = wb.active
                        prices = [float(row[1]) for row in ws.iter_rows(min_row=2, max_col=2, values_only=True) if row[1]]
                        if prices:
                            offer_id = f.stem.split("_")[-1]
                            cursor.execute("INSERT INTO ali1688_sources (item_id, title, offer_id, min_price, sku_count, source_url) VALUES (%s, %s, %s, %s, %s, %s)",
                                         (item_db_id, f.stem.replace(f"_{offer_id}", ""), offer_id, min(prices), len(prices), f"https://detail.1688.com/offer/{offer_id}.html"))
                            source_count += 1
                    except: continue
                
                # 方案 B: 扫描 .csv (针对早期数据)
                for f in item_dir.glob("*.csv"):
                    if f.name == 'summary.csv': continue
                    try:
                        prices = []
                        with open(f, 'r', encoding='utf-8-sig') as csvfile:
                            reader = csv.DictReader(csvfile)
                            for row in reader:
                                if row.get('价格'): prices.append(float(row['价格']))
                        if prices:
                            offer_id = f.stem.split("_")[-1]
                            cursor.execute("INSERT INTO ali1688_sources (item_id, title, offer_id, min_price, sku_count, source_url) VALUES (%s, %s, %s, %s, %s, %s)",
                                         (item_db_id, f.stem.replace(f"_{offer_id}", ""), offer_id, min(prices), len(prices), f"https://detail.1688.com/offer/{offer_id}.html"))
                            source_count += 1
                    except: continue
                    
                print(f"  - Rank {i}: Imported {source_count} sources.")
            conn.commit()
        except Exception as e:
            print(f"  [Error]: {e}")
            conn.rollback()

    conn.close()
    print("\nMigration completed!")

if __name__ == "__main__":
    migrate()
