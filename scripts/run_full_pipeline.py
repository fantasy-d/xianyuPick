import os, json, asyncio, re, csv, argparse, sys, pymysql
from datetime import datetime
from pathlib import Path

# --- 辅助 ---
async def run_command(cmd):
    process = await asyncio.create_subprocess_shell(cmd, stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.PIPE)
    stdout, stderr = await process.communicate()
    return stdout.decode('utf-8', errors='ignore')

def sanitize_dir_name(name: str) -> str:
    clean = re.sub(r'\s+', '', str(name))
    return re.sub(r'[\\/:*?"<>|]', '_', clean).strip()[:60]

# --- 核心：定义 BASE_DIR ---
BASE_DIR = Path(__file__).resolve().parents[1]

def get_db_conn():
    config_path = BASE_DIR / "config" / "database.json"
    config = json.load(open(config_path))
    config["cursorclass"] = pymysql.cursors.DictCursor
    return pymysql.connect(**config)

def extract_json(text):
    try:
        import re
        match = re.search(r'(\{.*"hot_items".*\})', text, re.DOTALL)
        if match: return match.group(1)
    except: pass
    return text

async def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--keyword", required=True)
    parser.add_argument("--task-id", required=False)
    args = parser.parse_args()
    
    keyword, task_id = args.keyword, args.task_id
    python_path = sys.executable
    
    # 状态初始化
    Task = None
    checkpoint = {}
    root_dir = None

    if task_id:
        sys.path.insert(0, str(BASE_DIR / 'src'))
        try:
            from web_api.main import Task
            conn = get_db_conn(); cursor = conn.cursor()
            cursor.execute("SELECT root_dir, checkpoint FROM tasks WHERE id = %s", (task_id,))
            row = cursor.fetchone(); conn.close()
            if row:
                root_dir = Path(row['root_dir'])
                checkpoint = json.loads(row['checkpoint']) if row['checkpoint'] else {}
        except: pass

    if not root_dir:
        root_dir = Path("outputs") / f"{sanitize_dir_name(keyword)}_{datetime.now().strftime('%Y%m%d')}"
    root_dir.mkdir(parents=True, exist_ok=True)

    # --- 1. 闲鱼扫描 (Phase 1) ---
    cur_phase = checkpoint.get("phase", 1)
    xianyu_json_path = root_dir / "xianyu_hot_items.json"
    
    if cur_phase > 1 and xianyu_json_path.exists() and xianyu_json_path.stat().st_size > 100:
        print(f"Checkpoint: Skipping Phase 1.", flush=True)
        raw_output = xianyu_json_path.read_text()
    else:
        if Task: Task.update(task_id, msg="闲鱼扫描中...", progress=10)
        cmd_xianyu = f"export PYTHONPATH=$PYTHONPATH:{BASE_DIR}/src && {python_path} scripts/run_xianyu_hot_items.py --keyword '{keyword}' --state-file xianyu_state.json --max-pages 1 --top-n 10"
        full_output = await run_command(cmd_xianyu)
        raw_output = extract_json(full_output)
        try:
            json.loads(raw_output)
            xianyu_json_path.write_text(raw_output)
            if Task: Task.update(task_id, checkpoint=json.dumps({"phase": 2, "processed_rank": 0}))
        except:
            if Task: Task.update(task_id, status="失败", msg="解析闲鱼数据失败")
            return

    try:
        hot_items = json.loads(raw_output).get("hot_items", [])
    except: return

    # --- 资产初始化 ---
    db_item_ids = {}
    if task_id:
        conn = get_db_conn(); cursor = conn.cursor()
        cursor.execute("SELECT id, rank_index FROM xianyu_items WHERE task_id = %s", (task_id,))
        db_items = cursor.fetchall()
        if not db_items:
            for i, item in enumerate(hot_items, start=1):
                cursor.execute("INSERT INTO xianyu_items (task_id, rank_index, title, price, image_url, want_count, item_url) VALUES (%s, %s, %s, %s, %s, %s, %s)",
                             (task_id, i, item.get('title'), item.get('price'), item.get('image_url'), item.get('want_count'), item.get('item_url')))
                db_item_ids[i] = cursor.lastrowid
            conn.commit()
        else:
            db_item_ids = {row['rank_index']: row['id'] for row in db_items}
        conn.close()

    # --- 2. 1688 深度验证 (Phase 2) ---
    processed_rank = checkpoint.get("processed_rank", 0)
    if Task: Task.update(task_id, msg="1688溯源中...", progress=30)
    
    for i, item in enumerate(hot_items, start=1):
        if i <= processed_rank: continue

        # 暂停自检 (即用即连模式)
        if task_id:
            try:
                _c = get_db_conn(); _cur = _c.cursor()
                _cur.execute("SELECT status FROM tasks WHERE id = %s", (task_id,))
                _s = (_cur.fetchone() or {}).get('status'); _c.close()
                if _s == "正在暂停":
                    if Task: Task.update(task_id, status="已暂停", msg=f"在第 {i} 个商品处暂停")
                    return
            except: pass

        if Task: Task.update(task_id, progress=30 + int((i/len(hot_items))*60), msg=f"处理爆款 {i}/{len(hot_items)}...")
        
        item_dir = root_dir / f"Rank_{i}_{sanitize_dir_name(item.get('title', 'item'))}"
        item_dir.mkdir(parents=True, exist_ok=True)
        
        cmd_1688 = f"export PYTHONPATH=$PYTHONPATH:{BASE_DIR}/src && {python_path} scripts/run_ali1688_slow_flow.py --image-url '{item.get('image_url')}' --output-dir '{item_dir}' --detail-top-n 10 --target-keyword '{keyword}'"
        await run_command(cmd_1688)
        
        # 资产入库 (即用即连)
        if task_id and i in db_item_ids:
            item_db_id = db_item_ids[i]
            _conn = get_db_conn(); _cursor = _conn.cursor()
            try:
                _cursor.execute("DELETE FROM ali1688_sources WHERE item_id = %s", (item_db_id,))
                # 尝试读取 summary.json
                summary_path = item_dir / "summary.json"
                if summary_path.exists():
                    results = json.loads(summary_path.read_text())
                    for res in results:
                        if res.get("status") == "dropped":
                            _cursor.execute("INSERT INTO ali1688_sources (item_id, task_id, title, offer_id, min_price, sku_count, source_url, drop_reason) VALUES (%s,%s,%s,%s,%s,%s,%s,%s)",
                                         (item_db_id, task_id, res['title'], res['offer_id'], res['min_price'], 0, res['item_url'], res['drop_reason']))
                        else:
                            # 读取详情价格
                            offer_id = res['offer_id']
                            xlsx = list(item_dir.glob(f"*_{offer_id}.xlsx"))
                            if xlsx:
                                from openpyxl import load_workbook
                                wb = load_workbook(filename=xlsx[0], read_only=True)
                                prices = [float(row[1]) for row in wb.active.iter_rows(min_row=2, max_col=2, values_only=True) if row[1]]
                                if prices:
                                    _cursor.execute("INSERT INTO ali1688_sources (item_id, task_id, title, offer_id, min_price, sku_count, source_url) VALUES (%s,%s,%s,%s,%s,%s,%s)",
                                                 (item_db_id, task_id, res['title'], offer_id, min(prices), len(prices), res['item_url']))
                # 更新 Checkpoint
                _cursor.execute("UPDATE tasks SET checkpoint = %s WHERE id = %s", (json.dumps({"phase": 2, "processed_rank": i}), task_id))
                _conn.commit()
            except Exception as e: print(f"DB Sync Error: {e}")
            finally: _conn.close()

    if Task: Task.update(task_id, status="已完成", progress=100, msg="分析完成")

if __name__ == "__main__":
    asyncio.run(main())
