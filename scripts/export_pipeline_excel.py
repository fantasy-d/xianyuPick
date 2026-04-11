#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Export the current pipeline result files to a single Excel workbook."
    )
    parser.add_argument(
        "--hot-items-file",
        required=True,
        help="Path to the xianyu hot_items JSON file.",
    )
    parser.add_argument(
        "--source-bundle-file",
        required=True,
        help="Path to the source_bundle JSON file.",
    )
    parser.add_argument(
        "--profit-analysis-file",
        required=True,
        help="Path to the profit_analysis JSON file.",
    )
    parser.add_argument(
        "--listing-candidates-file",
        required=True,
        help="Path to the listing_candidates JSON file.",
    )
    parser.add_argument(
        "--output-file",
        required=True,
        help="Path to the output .xlsx file.",
    )
    args = parser.parse_args()

    hot_payload = _load_json(Path(args.hot_items_file))
    source_bundle = _load_json(Path(args.source_bundle_file))
    profit_payload = _load_json(Path(args.profit_analysis_file))
    listing_payload = _load_json(Path(args.listing_candidates_file))

    hot_items = list(hot_payload.get("hot_items") or source_bundle.get("hot_items") or [])
    xianyu_market = dict(hot_payload.get("xianyu_market") or {})
    source_resolution = list(source_bundle.get("source_resolution") or [])
    source_items = list(source_bundle.get("source_items") or [])
    profit_analysis = list(profit_payload.get("profit_analysis") or [])
    listing_candidates = list(listing_payload.get("listing_candidates") or [])
    accepted_source_items = [
        row for row in source_items if str(row.get("candidate_status") or "accepted") == "accepted"
    ]
    filtered_source_items = [
        row for row in source_items if str(row.get("candidate_status") or "") == "filtered"
    ]

    hot_items_by_id = {
        str(row.get("hot_item_id") or ""): row for row in hot_items if isinstance(row, dict)
    }
    source_items_by_id = {
        str(row.get("source_item_id") or ""): row for row in accepted_source_items if isinstance(row, dict)
    }

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "汇总"

    _write_summary_sheet(
        summary_sheet,
        xianyu_market=xianyu_market,
        hot_items=hot_items,
        source_resolution=source_resolution,
        source_items=accepted_source_items,
        filtered_source_items=filtered_source_items,
        profit_analysis=profit_analysis,
        listing_candidates=listing_candidates,
    )

    _write_table_sheet(
        workbook,
        "筛选总览",
        _build_pipeline_overview_rows(
            hot_items=hot_items,
            source_resolution=source_resolution,
            source_items=accepted_source_items,
            filtered_source_items=filtered_source_items,
            profit_analysis=profit_analysis,
            listing_candidates=listing_candidates,
        ),
    )
    _write_table_sheet(workbook, "闲鱼热品", [_serialize_hot_item_row(row) for row in hot_items])
    _write_table_sheet(
        workbook,
        "货源匹配结果",
        [_serialize_source_resolution_row(row) for row in source_resolution],
    )
    _write_table_sheet(
        workbook,
        "1688货源候选",
        [_serialize_source_item_row(row, hot_items_by_id=hot_items_by_id) for row in source_items],
    )
    _write_table_sheet(
        workbook,
        "利润测算",
        [_serialize_profit_row(row) for row in profit_analysis],
    )
    _write_table_sheet(
        workbook,
        "最终上架候选",
        [
            _serialize_listing_candidate_row(
                row,
                hot_items_by_id=hot_items_by_id,
                source_items_by_id=source_items_by_id,
            )
            for row in listing_candidates
        ],
    )

    output_file = Path(args.output_file)
    output_file.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_file)
    print(
        json.dumps(
            {
                "output_file": str(output_file),
                "hot_items_count": len(hot_items),
                "source_resolution_count": len(source_resolution),
                "source_items_count": len(accepted_source_items),
                "filtered_source_items_count": len(filtered_source_items),
                "profit_analysis_count": len(profit_analysis),
                "listing_candidates_count": len(listing_candidates),
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


def _load_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def _write_summary_sheet(
    sheet,
    *,
    xianyu_market: dict[str, Any],
    hot_items: list[dict[str, Any]],
    source_resolution: list[dict[str, Any]],
    source_items: list[dict[str, Any]],
    filtered_source_items: list[dict[str, Any]],
    profit_analysis: list[dict[str, Any]],
    listing_candidates: list[dict[str, Any]],
) -> None:
    resolved_count = sum(1 for row in source_resolution if row.get("resolved"))
    resolution_reason_counts = Counter(
        str(row.get("resolution_reason") or "unknown") for row in source_resolution
    )
    blocked_reason_counts = Counter()
    for row in listing_candidates:
        for reason in row.get("blocked_by") or []:
            blocked_reason_counts[str(reason)] += 1

    stats = dict(xianyu_market.get("top10_price_stats") or {})

    rows = [
        ("品类关键词", xianyu_market.get("category_keyword") or ""),
        ("闲鱼搜索结果数", xianyu_market.get("result_count") or 0),
        ("筛选条件", ",".join(xianyu_market.get("filtered_by") or [])),
        ("排序方式", xianyu_market.get("sorted_by") or ""),
        ("Top数量", xianyu_market.get("top_n") or len(hot_items)),
        ("Top10最低价", stats.get("min")),
        ("Top10最高价", stats.get("max")),
        ("Top10中位价", stats.get("median")),
        ("热品数量", len(hot_items)),
        ("货源解析记录数", len(source_resolution)),
        ("成功匹配货源数", resolved_count),
        ("未匹配货源数", len(source_resolution) - resolved_count),
        ("1688有效货源条数", len(source_items)),
        ("1688已过滤候选条数", len(filtered_source_items)),
        ("利润分析条数", len(profit_analysis)),
        ("最终候选数", len(listing_candidates)),
        (
            "推荐上架数",
            sum(1 for row in listing_candidates if bool(row.get("is_recommended"))),
        ),
    ]
    sheet.append(["指标", "值"])
    for row in rows:
        sheet.append(list(row))

    sheet.append([])
    sheet.append(["货源匹配结果", "数量"])
    for reason, count in sorted(resolution_reason_counts.items()):
        sheet.append([reason, count])

    sheet.append([])
    sheet.append(["拦截原因", "数量"])
    for reason, count in sorted(blocked_reason_counts.items()):
        sheet.append([reason, count])

    _style_sheet(sheet)
    _autosize_sheet(sheet)


def _build_pipeline_overview_rows(
    *,
    hot_items: list[dict[str, Any]],
    source_resolution: list[dict[str, Any]],
    source_items: list[dict[str, Any]],
    filtered_source_items: list[dict[str, Any]],
    profit_analysis: list[dict[str, Any]],
    listing_candidates: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    resolution_by_hot_item = {
        str(row.get("hot_item_id") or ""): row for row in source_resolution if isinstance(row, dict)
    }
    source_items_by_hot_item: dict[str, list[dict[str, Any]]] = {}
    for row in source_items:
        if not isinstance(row, dict):
            continue
        hot_item_id = str(row.get("hot_item_id") or "")
        source_items_by_hot_item.setdefault(hot_item_id, []).append(row)
    filtered_source_items_by_hot_item: dict[str, list[dict[str, Any]]] = {}
    for row in filtered_source_items:
        if not isinstance(row, dict):
            continue
        hot_item_id = str(row.get("hot_item_id") or "")
        filtered_source_items_by_hot_item.setdefault(hot_item_id, []).append(row)
    profit_by_pair = {
        (str(row.get("hot_item_id") or ""), str(row.get("source_item_id") or "")): row
        for row in profit_analysis
        if isinstance(row, dict)
    }
    candidate_by_hot_item = {
        str(row.get("hot_item_id") or ""): row for row in listing_candidates if isinstance(row, dict)
    }

    rows: list[dict[str, Any]] = []
    for hot_item in hot_items:
        hot_item_id = str(hot_item.get("hot_item_id") or "")
        resolution = resolution_by_hot_item.get(hot_item_id, {})
        sources = source_items_by_hot_item.get(hot_item_id, [])
        best_source = sources[0] if sources else {}
        candidate = candidate_by_hot_item.get(hot_item_id, {})
        candidate_pair = (
            hot_item_id,
            str(candidate.get("source_item_id") or ""),
        )
        best_profit = profit_by_pair.get(candidate_pair, {})
        if not best_profit and best_source:
            best_profit = profit_by_pair.get(
                (hot_item_id, str(best_source.get("source_item_id") or "")),
                {},
            )
        rows.append(
            {
                "闲鱼商品ID": hot_item_id,
                "闲鱼标题": hot_item.get("title") or "",
                "闲鱼售价": hot_item.get("price"),
                "想要人数": hot_item.get("want_count"),
                "卖家名称": hot_item.get("seller_name") or "",
                "地区": hot_item.get("area") or "",
                "图片链接": hot_item.get("image_url") or "",
                "闲鱼链接": hot_item.get("item_url") or "",
                "是否匹配到货源": bool(resolution.get("resolved")),
                "匹配结果": resolution.get("resolution_reason") or "",
                "货源候选数": len(sources),
                "已过滤候选数": len(filtered_source_items_by_hot_item.get(hot_item_id, [])),
                "最佳货源ID": best_source.get("source_item_id") or "",
                "最佳货源标题": best_source.get("title") or "",
                "最佳货源价格": best_source.get("price"),
                "最佳货源链接": best_source.get("item_url") or "",
                "预估利润": best_profit.get("estimated_margin"),
                "毛利率": best_profit.get("gross_margin_rate"),
                "成本利润率": best_profit.get("cost_profit_rate"),
                "是否推荐上架": candidate.get("is_recommended"),
                "推荐原因": _join_list(candidate.get("reasons")),
                "拦截原因": _join_list(candidate.get("blocked_by")),
            }
        )
    return rows


def _serialize_hot_item_row(row: dict[str, Any]) -> dict[str, Any]:
    metadata = dict(row.get("metadata") or {})
    return {
        "闲鱼商品ID": row.get("hot_item_id") or "",
        "平台": row.get("platform") or "",
        "标题": row.get("title") or "",
        "售价": row.get("price"),
        "想要人数": row.get("want_count"),
        "卖家名称": row.get("seller_name") or "",
        "地区": row.get("area") or "",
        "图片链接": row.get("image_url") or "",
        "闲鱼链接": row.get("item_url") or "",
        "发布时间": metadata.get("publish_time") or "",
        "标签": _join_list(metadata.get("tags")),
    }


def _serialize_source_resolution_row(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "闲鱼商品ID": row.get("hot_item_id") or "",
        "是否匹配成功": bool(row.get("resolved")),
        "货源商品ID列表": _join_list(row.get("source_item_ids")),
        "匹配结果": row.get("resolution_reason") or "",
    }


def _serialize_source_item_row(
    row: dict[str, Any],
    *,
    hot_items_by_id: dict[str, dict[str, Any]],
) -> dict[str, Any]:
    hot_item = hot_items_by_id.get(str(row.get("hot_item_id") or ""), {})
    return {
        "闲鱼商品ID": row.get("hot_item_id") or "",
        "闲鱼商品名称": hot_item.get("title") or "",
        "货源商品ID": row.get("source_item_id") or "",
        "货源平台": row.get("source_platform") or "",
        "货源标题": row.get("title") or "",
        "货源价格": row.get("price"),
        "货源链接": row.get("item_url") or "",
        "店铺名称": row.get("shop_name") or "",
        "销量": row.get("sales"),
        "图片链接": row.get("image_url") or "",
        "运费": row.get("shipping_fee"),
        "候选状态": _candidate_status_label(row.get("candidate_status")),
        "过滤原因": _filter_reason_label(row.get("filter_reason")),
    }


def _serialize_profit_row(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "闲鱼商品ID": row.get("hot_item_id") or "",
        "货源商品ID": row.get("source_item_id") or "",
        "货源成本": row.get("source_cost"),
        "闲鱼对比价": row.get("target_xianyu_price"),
        "预估利润": row.get("estimated_margin"),
        "毛利率": row.get("gross_margin_rate"),
        "成本利润率": row.get("cost_profit_rate"),
        "风险标记": _join_list(row.get("risk_flags")),
    }


def _serialize_listing_candidate_row(
    row: dict[str, Any],
    *,
    hot_items_by_id: dict[str, dict[str, Any]],
    source_items_by_id: dict[str, dict[str, Any]],
) -> dict[str, Any]:
    hot_item = hot_items_by_id.get(str(row.get("hot_item_id") or ""), {})
    source_item = source_items_by_id.get(str(row.get("source_item_id") or ""), {})
    return {
        "闲鱼商品ID": row.get("hot_item_id") or "",
        "闲鱼商品名称": hot_item.get("title") or "",
        "闲鱼商品链接": hot_item.get("item_url") or "",
        "货源商品ID": row.get("source_item_id") or "",
        "1688商品名称": source_item.get("title") or "",
        "1688商品链接": source_item.get("item_url") or "",
        "闲鱼价格": hot_item.get("price"),
        "1688价格": source_item.get("price"),
        "预估利润": row.get("estimated_margin"),
        "毛利率": row.get("gross_margin_rate"),
        "成本利润率": row.get("cost_profit_rate"),
        "是否推荐上架": bool(row.get("is_recommended")),
        "推荐原因": _join_list(row.get("reasons")),
        "拦截原因": _join_list(row.get("blocked_by")),
    }


def _write_table_sheet(workbook: Workbook, title: str, rows: list[dict[str, Any]]) -> None:
    sheet = workbook.create_sheet(title)
    headers = _collect_headers(rows)
    if not headers:
        sheet.append(["empty"])
        _autosize_sheet(sheet)
        return
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header) for header in headers])
    _style_sheet(sheet)
    _autosize_sheet(sheet)


def _collect_headers(rows: list[dict[str, Any]]) -> list[str]:
    headers: list[str] = []
    for row in rows:
        for key in row:
            if key not in headers:
                headers.append(key)
    return headers


def _join_list(value: Any) -> str:
    if not value:
        return ""
    if isinstance(value, list):
        return ",".join(str(item) for item in value if item not in (None, ""))
    return str(value)


def _candidate_status_label(value: Any) -> str:
    mapping = {
        "accepted": "已保留",
        "filtered": "已过滤",
    }
    return mapping.get(str(value or "accepted"), str(value or ""))


def _filter_reason_label(value: Any) -> str:
    mapping = {
        "": "",
        "title_irrelevant": "商品标题与原始搜索词不相关",
    }
    return mapping.get(str(value or ""), str(value or ""))


def _autosize_sheet(sheet) -> None:
    for column_cells in sheet.columns:
        max_length = 0
        column_letter = column_cells[0].column_letter
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))
        sheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 60)


def _style_sheet(sheet) -> None:
    header_font = Font(name="PingFang SC", size=16, bold=True)
    body_font = Font(name="PingFang SC", size=15)
    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    for row_index, row in enumerate(sheet.iter_rows(), start=1):
        for cell in row:
            if row_index == 1:
                cell.font = header_font
                cell.fill = header_fill
            else:
                cell.font = body_font


if __name__ == "__main__":
    sys.exit(main())
