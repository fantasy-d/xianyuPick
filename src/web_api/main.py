from __future__ import annotations

import json, asyncio, os, uuid, pymysql, re, sys, signal, logging, subprocess, shlex
from datetime import datetime
from pathlib import Path
from typing import Any, Dict
from urllib.parse import unquote
from fastapi import FastAPI
from fastapi.responses import PlainTextResponse, FileResponse
from fastapi.staticfiles import StaticFiles

from xianyu_tools.logging_util import get_unified_logger

# --- 日志配置 ---
logger = get_unified_logger("WebAPI")

app = FastAPI(title="Xianyu-1688 Management System")
SOURCE_SKUS_CACHE: Dict[int, list] = {}


class NoCacheStaticFiles(StaticFiles):
    async def get_response(self, path, scope):
        response = await super().get_response(path, scope)
        if getattr(response, "status_code", 200) < 400:
            response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate"
            response.headers["Pragma"] = "no-cache"
            response.headers["Expires"] = "0"
        return response

from xianyu_tools.config import settings
from xianyu_tools.source_channel_config import (
    build_source_channel_account_runtime,
    get_source_channel,
    get_source_channel_account,
    normalize_active_source_account_ids,
    normalize_source_channels_config,
    strip_source_channel_runtime_fields,
)

# --- 常量 ---
BASE_DIR = Path(__file__).resolve().parents[2]
WEB_DIR = BASE_DIR / "web"
OUTPUTS_DIR = BASE_DIR / "outputs"

DEFAULT_GROSS_PROFIT_RATE = 0.3

# --- 辅助函数 ---
DB_CONFIG = settings.get_database_config()
DB_CONFIG["cursorclass"] = pymysql.cursors.DictCursor
def get_db_conn(): return pymysql.connect(**DB_CONFIG)


def _to_float(value, default: float = 0.0) -> float:
    try:
        return float(value)
    except Exception:
        return default


def _extract_xianyu_item_id(item_url: str | None) -> str | None:
    text = str(item_url or "").strip()
    if not text:
        return None
    query_match = re.search(r"(?:[?&]id=)([^&#]+)", text)
    if query_match:
        return unquote(query_match.group(1)).strip() or None
    path_match = re.search(r"/item/(\d+)", text)
    if path_match:
        return path_match.group(1)
    return None


def _normalize_gross_profit_rate(value) -> float:
    try:
        rate = float(value)
        if rate > 1:
            rate = rate / 100
        if rate <= 0:
            return DEFAULT_GROSS_PROFIT_RATE
        return min(rate, 1)
    except Exception:
        return DEFAULT_GROSS_PROFIT_RATE


def _compute_source_estimated_profit(source_row: dict, gross_profit_rate: float | None = None) -> float:
    source = source_row or {}
    rate = _normalize_gross_profit_rate(gross_profit_rate)
    return _to_float(source.get("min_price")) * rate


def _safe_int(value) -> int | None:
    try:
        if value is None or value == "":
            return None
        return int(value)
    except Exception:
        return None


def _normalize_openapi_product_status(detail_data: dict[str, Any]) -> dict[str, Any]:
    data = detail_data if isinstance(detail_data, dict) else {}
    product_status = _safe_int(data.get("product_status"))
    publish_status = _safe_int(data.get("publish_status"))

    publish_status_map = {
        -1: ("failed", "不可操作"),
        1: ("selected", "草稿箱"),
        2: ("pending", "待发布"),
        3: ("success", "销售中"),
        4: ("depublished", "已下架"),
        5: ("depublished", "已售罄"),
        9: ("failed", "商品异常"),
    }
    if publish_status in publish_status_map:
        status, label = publish_status_map[publish_status]
        return {"status": status, "label": label}

    product_status_map = {
        -1: ("deleted", "已删除"),
        21: ("pending", "待发布"),
        22: ("success", "销售中"),
        23: ("depublished", "已售罄"),
        31: ("depublished", "手动下架"),
        33: ("depublished", "售出下架"),
        36: ("depublished", "自动下架"),
    }
    if product_status in product_status_map:
        status, label = product_status_map[product_status]
        return {"status": status, "label": label}

    status_text = " ".join(
        str(data.get(key) or "")
        for key in ("status", "status_text", "status_name", "product_status_text", "publish_status_text")
    )
    if any(word in status_text for word in ("已删除", "删除")):
        return {"status": "deleted", "label": status_text.strip() or "已删除"}
    if any(word in status_text for word in ("已下架", "下架")):
        return {"status": "depublished", "label": status_text.strip() or "已下架"}
    if any(word in status_text for word in ("已上架", "上架中", "在线", "在售", "出售中", "发布成功")):
        return {"status": "success", "label": status_text.strip() or "已上架"}
    if any(word in status_text for word in ("发布中", "审核中", "同步中")):
        return {"status": "pending", "label": status_text.strip() or "发布中"}
    if any(word in status_text for word in ("失败", "异常", "驳回")):
        return {"status": "failed", "label": status_text.strip() or "发布失败"}
    if any(word in status_text for word in ("待发布", "草稿")):
        return {"status": "selected", "label": status_text.strip() or "待发布"}

    return {
        "status": "",
        "label": "",
        "raw_product_status": product_status,
        "raw_publish_status": publish_status,
    }


def _format_unix_time(value) -> str:
    try:
        timestamp = int(value or 0)
        if timestamp <= 0:
            return ""
        return datetime.fromtimestamp(timestamp).strftime("%Y-%m-%d %H:%M:%S")
    except Exception:
        return ""


def _format_cent_amount(value) -> str:
    try:
        return f"¥{float(value or 0) / 100:.2f}"
    except Exception:
        return "¥0.00"


def _parse_cent_amount_from_request(req: dict, cent_key: str, yuan_key: str, *, default: int | None = None) -> int | None:
    if not isinstance(req, dict):
        return default
    raw_cent = req.get(cent_key)
    if raw_cent not in (None, ""):
        try:
            return int(raw_cent)
        except Exception:
            return None
    raw_yuan = req.get(yuan_key)
    if raw_yuan in (None, ""):
        return default
    try:
        return int(round(float(raw_yuan) * 100))
    except Exception:
        return None


OPENAPI_ORDER_STATUS_LABELS = {
    11: "待付款",
    12: "待发货",
    21: "已发货",
    22: "已完成",
    23: "已退款",
    24: "已关闭",
}


def _format_openapi_order_status(value) -> str:
    try:
        status = int(value)
    except Exception:
        return "未知状态"
    return OPENAPI_ORDER_STATUS_LABELS.get(status, f"状态 {status}")


def _normalize_openapi_order(order_data: dict[str, Any] | None, *, include_raw: bool = False) -> dict[str, Any]:
    order = order_data if isinstance(order_data, dict) else {}
    goods = order.get("goods") if isinstance(order.get("goods"), dict) else {}
    images = goods.get("images") if isinstance(goods.get("images"), list) else []
    address_parts = [
        order.get("prov_name"),
        order.get("city_name"),
        order.get("area_name"),
        order.get("town_name"),
        order.get("address"),
    ]
    normalized = {
        "order_no": str(order.get("order_no") or ""),
        "order_status": order.get("order_status"),
        "order_status_label": _format_openapi_order_status(order.get("order_status")),
        "refund_status": order.get("refund_status"),
        "order_time": order.get("order_time"),
        "order_time_text": _format_unix_time(order.get("order_time")),
        "pay_time_text": _format_unix_time(order.get("pay_time")),
        "consign_time_text": _format_unix_time(order.get("consign_time")),
        "confirm_time_text": _format_unix_time(order.get("confirm_time")),
        "cancel_time_text": _format_unix_time(order.get("cancel_time")),
        "update_time_text": _format_unix_time(order.get("update_time")),
        "total_amount": order.get("total_amount"),
        "total_amount_text": _format_cent_amount(order.get("total_amount")),
        "pay_amount": order.get("pay_amount"),
        "pay_amount_text": _format_cent_amount(order.get("pay_amount")),
        "express_fee": order.get("express_fee"),
        "express_fee_text": _format_cent_amount(order.get("express_fee")),
        "buyer_nick": str(order.get("buyer_nick") or ""),
        "seller_name": str(order.get("seller_name") or ""),
        "seller_remark": str(order.get("seller_remark") or ""),
        "receiver_name": str(order.get("receiver_name") or ""),
        "receiver_mobile": str(order.get("receiver_mobile") or ""),
        "receiver_address": "".join(str(part or "") for part in address_parts),
        "waybill_no": str(order.get("waybill_no") or ""),
        "express_code": str(order.get("express_code") or ""),
        "express_name": str(order.get("express_name") or ""),
        "goods": {
            "title": str(goods.get("title") or ""),
            "quantity": goods.get("quantity"),
            "price": goods.get("price"),
            "price_text": _format_cent_amount(goods.get("price")),
            "product_id": str(goods.get("product_id") or ""),
            "item_id": str(goods.get("item_id") or ""),
            "outer_id": str(goods.get("outer_id") or ""),
            "sku_id": str(goods.get("sku_id") or ""),
            "sku_text": str(goods.get("sku_text") or ""),
            "image": str(images[0]) if images else "",
            "images": images,
            "service_support": str(goods.get("service_support") or ""),
        },
    }
    if include_raw:
        normalized["raw"] = order
    return normalized


def _extract_openapi_order_list(data: dict[str, Any]) -> tuple[list[dict[str, Any]], int]:
    payload = data if isinstance(data, dict) else {}
    raw_items = payload.get("list") or payload.get("orders") or payload.get("items") or []
    if not isinstance(raw_items, list):
        raw_items = []
    total = payload.get("total") or payload.get("total_count") or payload.get("count") or len(raw_items)
    try:
        total = int(total)
    except Exception:
        total = len(raw_items)
    return [_normalize_openapi_order(item) for item in raw_items], total


def _ordered_unique_string_list(values) -> list[str]:
    if not isinstance(values, list):
        return []
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        text = str(value or "").strip()
        if not text or text in seen:
            continue
        seen.add(text)
        result.append(text)
    return result


def _summarize_channel_filter_snapshot(
    snapshot: dict | None,
    *,
    channel_type: str | None = None,
    has_recorded_snapshot: bool = True,
) -> dict:
    normalized_snapshot = settings.normalize_channel_search_filter_snapshot(
        snapshot if isinstance(snapshot, dict) else {},
        channel_type=str(channel_type or "").strip(),
    )
    filter_status_map = (
        normalized_snapshot.get("filter_status_map")
        if isinstance(normalized_snapshot.get("filter_status_map"), dict)
        else {}
    )
    configured = _ordered_unique_string_list(
        normalized_snapshot.get("configured_enabled_filter_keys")
        or normalized_snapshot.get("configured_filter_keys")
        or normalized_snapshot.get("enabled_filter_keys")
    )
    query_injected = [
        key for key, meta in filter_status_map.items()
        if isinstance(meta, dict) and meta.get("status") == "query_injected_pending_verification"
    ]
    applied = [
        key for key, meta in filter_status_map.items()
        if isinstance(meta, dict) and meta.get("status") == "applied"
    ]
    unapplied = [
        key for key, meta in filter_status_map.items()
        if isinstance(meta, dict) and meta.get("status") == "unapplied"
    ]
    resolved_channel_type = str(
        normalized_snapshot.get("channel_type") or channel_type or ""
    ).strip().lower()
    unsupported = bool(resolved_channel_type) and resolved_channel_type != "ali1688"
    configured_pending_reason_set = {
        "runtime_mapping_not_implemented_yet",
        "query_candidate_not_validated",
        "semantic_combo_not_confirmed",
        "snapshot_only_until_semantics_confirmed",
        "special_panel_unmapped",
        "special_panel_entry_detected_unmapped",
    }
    configured_pending = []
    for filter_key in configured:
        filter_meta = filter_status_map.get(filter_key)
        if not isinstance(filter_meta, dict):
            configured_pending.append(filter_key)
            continue
        if filter_meta.get("status") != "unapplied":
            continue
        mapping_stage = str(filter_meta.get("mapping_stage") or "").strip()
        reason = str(filter_meta.get("reason") or "").strip()
        if mapping_stage == "snapshot_only" or reason in configured_pending_reason_set:
            configured_pending.append(filter_key)

    applied_set = set(applied)
    query_injected_set = set(query_injected)
    configured_pending_set = set(configured_pending)
    unapplied_strict = [
        filter_key for filter_key in unapplied
        if filter_key not in applied_set
        and filter_key not in query_injected_set
        and filter_key not in configured_pending_set
    ]
    top_level_query_verification_details = (
        normalized_snapshot.get("query_verification_details")
        if isinstance(normalized_snapshot.get("query_verification_details"), dict)
        else normalized_snapshot.get("verification_details")
        if isinstance(normalized_snapshot.get("verification_details"), dict)
        else {}
    )
    status_map_verification_details = {}
    for filter_key, meta in filter_status_map.items():
        if not isinstance(meta, dict):
            continue
        verification_detail = meta.get("verification_detail")
        if isinstance(verification_detail, dict):
            status_map_verification_details[filter_key] = verification_detail

    has_signal = _channel_filter_snapshot_has_signal(normalized_snapshot)
    return {
        "configured": configured,
        "configured_pending": configured_pending,
        "query_injected": query_injected,
        "applied": applied,
        "unapplied": unapplied_strict,
        "unsupported": unsupported,
        "configured_filter_count": len(configured),
        "filter_status_map": filter_status_map,
        "query_verification_details": {
            **status_map_verification_details,
            **top_level_query_verification_details,
        },
        "mapping_stage": str(normalized_snapshot.get("mapping_stage") or "").strip(),
        "mapping_notes": str(normalized_snapshot.get("mapping_notes") or "").strip(),
        "runtime_audit_stage": str(normalized_snapshot.get("runtime_audit_stage") or "").strip(),
        "runtime_audit_source": str(normalized_snapshot.get("runtime_audit_source") or "").strip(),
        "legacy_missing_snapshot": (not unsupported and not has_recorded_snapshot and not has_signal),
        "has_runtime_signal": has_signal,
    }


def _sort_detail_sources_and_groups(
    *,
    source_rows: list[dict],
    channel_groups_map: dict[str, dict],
    used_channels_map: dict[str, dict],
    gross_profit_rate: float,
) -> tuple[list[dict], list[dict], list[dict]]:
    def source_page_order(source: dict) -> tuple[int, int]:
        page_index = int(source.get("page_original_index") or 0)
        normalized_index = page_index if page_index > 0 else 1_000_000_000
        return normalized_index, int(source.get("db_id") or 0)

    def source_profit_order(source: dict) -> tuple[float, int, int]:
        page_index, db_id = source_page_order(source)
        return -_compute_source_estimated_profit(source, gross_profit_rate), page_index, db_id

    sorted_sources = sorted(
        list(source_rows or []),
        key=source_profit_order,
    )

    decorated_groups: list[dict] = []
    for group in list(channel_groups_map.values()):
        sorted_group_sources = sorted(
            list(group.get("sources") or []),
            key=source_profit_order,
        )
        best_estimated_profit = (
            max(_compute_source_estimated_profit(source, gross_profit_rate) for source in sorted_group_sources)
            if sorted_group_sources
            else float("-inf")
        )
        first_page_original_index = (
            source_page_order(sorted_group_sources[0])[0]
            if sorted_group_sources
            else 1_000_000_000
        )
        decorated_groups.append(
            {
                **group,
                "sources": sorted_group_sources,
                "source_count": len(sorted_group_sources),
                "best_estimated_profit": best_estimated_profit,
                "first_page_original_index": first_page_original_index,
            }
        )

    sorted_groups = sorted(
        decorated_groups,
        key=lambda group: (
            -float(group.get("best_estimated_profit") or 0),
            int(group.get("first_page_original_index") or 1_000_000_000),
            str(group.get("channel_label") or group.get("channel_id") or ""),
        ),
    )

    sorted_used_channels = sorted(
        list(used_channels_map.values()),
        key=lambda channel: (
            next(
                (
                    index
                    for index, group in enumerate(sorted_groups)
                    if group.get("channel_id") == channel.get("channel_id")
                ),
                len(sorted_groups),
            ),
            str(channel.get("channel_label") or channel.get("channel_id") or ""),
        ),
    )
    return sorted_sources, sorted_groups, sorted_used_channels


def _build_detail_source_sort_strategy() -> dict:
    return {
        "field": "estimated_profit",
        "order": "desc",
        "label": "预估纯利倒序",
        "description": "当前结果按预估纯利从高到低固定排序，渠道筛选仅影响当前展示范围。",
    }


def _build_detail_channel_group_sort_strategy() -> dict:
    return {
        "field": "best_estimated_profit",
        "order": "desc",
        "label": "渠道最高预估纯利倒序",
        "description": "当前渠道分组按各渠道最高预估纯利从高到低固定排序。",
    }


def _build_task_channel_summary_map(
    *,
    task_channel_map: dict[str, list[dict]],
    snapshot_rows: list[dict],
) -> dict[str, list[dict]]:
    task_channel_snapshot_groups: dict[str, dict[str, dict]] = {}
    for snapshot_row in list(snapshot_rows or []):
        task_id = str(snapshot_row.get("task_id") or "").strip()
        channel_id = str(snapshot_row.get("channel_id") or "ali1688").strip() or "ali1688"
        if not task_id:
            continue
        task_group = task_channel_snapshot_groups.setdefault(task_id, {})
        if channel_id not in task_group:
            task_group[channel_id] = {
                "channel_id": channel_id,
                "channel_type": snapshot_row.get("channel_type") or "ali1688",
                "channel_label": snapshot_row.get("channel_label") or "1688 货源渠道",
                "snapshot": {},
                "has_recorded_snapshot": False,
            }
        channel_group = task_group[channel_id]
        raw_snapshot_json = snapshot_row.get("source_filter_snapshot_json")
        has_recorded_snapshot = bool(str(raw_snapshot_json or "").strip())
        channel_group["has_recorded_snapshot"] = bool(
            channel_group.get("has_recorded_snapshot") or has_recorded_snapshot
        )
        try:
            parsed_snapshot = json.loads(raw_snapshot_json or "{}")
            if not isinstance(parsed_snapshot, dict):
                parsed_snapshot = {}
        except Exception:
            parsed_snapshot = {}
        normalized_snapshot = settings.normalize_channel_search_filter_snapshot(
            parsed_snapshot,
            channel_id=channel_id,
            channel_type=channel_group.get("channel_type"),
        )
        if (
            _channel_filter_snapshot_has_signal(normalized_snapshot)
            and not _channel_filter_snapshot_has_signal(channel_group.get("snapshot"))
        ):
            channel_group["snapshot"] = normalized_snapshot

    task_channel_summary_map: dict[str, list[dict]] = {}
    for task_id, channels in task_channel_map.items():
        channel_groups = task_channel_snapshot_groups.get(task_id, {})
        task_channel_summary_map[task_id] = []
        for channel in channels:
            group = channel_groups.get(channel.get("channel_id") or "")
            task_channel_summary_map[task_id].append(
                {
                    **channel,
                    "filter_summary": _summarize_channel_filter_snapshot(
                        (group or {}).get("snapshot"),
                        channel_type=(group or {}).get("channel_type") or channel.get("channel_type"),
                        has_recorded_snapshot=bool((group or {}).get("has_recorded_snapshot")),
                    ),
                    "has_recorded_filter_snapshot": bool((group or {}).get("has_recorded_snapshot")),
                }
            )
    return task_channel_summary_map

def init_db_schema():
    try:
        conn = get_db_conn()
        cursor = conn.cursor()
        
        # 1. 创建 ali1688_skus 表
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS ali1688_skus (
                id INT AUTO_INCREMENT PRIMARY KEY,
                source_id INT NOT NULL,
                sku_text VARCHAR(255) NOT NULL,
                price DECIMAL(10,2) NOT NULL,
                stock INT NOT NULL,
                spec_id VARCHAR(50) DEFAULT '',
                image VARCHAR(1024) DEFAULT '',
                FOREIGN KEY (source_id) REFERENCES ali1688_sources(id) ON DELETE CASCADE
            ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
        """)
        
        # 1.5 创建 llm_token_logs 表
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS llm_token_logs (
                id INT AUTO_INCREMENT PRIMARY KEY,
                task_id VARCHAR(50) DEFAULT NULL,
                feature VARCHAR(50) NOT NULL,
                model VARCHAR(100) NOT NULL,
                prompt_tokens INT DEFAULT 0,
                completion_tokens INT DEFAULT 0,
                total_tokens INT DEFAULT 0,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
        """)
        
        # 2. 丢弃不需要的库内 HTML 表 ali1688_source_htmls
        cursor.execute("DROP TABLE IF EXISTS ali1688_source_htmls;")
        
        # 3. 自愈添加 html_path 字段以关联本地 HTML 物理文件
        try:
            cursor.execute("ALTER TABLE ali1688_sources ADD COLUMN html_path VARCHAR(1024) DEFAULT '';")
        except Exception:
            pass  # 如果列已经存在则会报错，直接忽略即可
        # 3.1 自愈添加货源渠道元数据，便于决策资产库按渠道展示
        try:
            cursor.execute("ALTER TABLE ali1688_sources ADD COLUMN source_channel_id VARCHAR(100) DEFAULT 'ali1688';")
        except Exception:
            pass
        try:
            cursor.execute("ALTER TABLE ali1688_sources ADD COLUMN source_channel_type VARCHAR(100) DEFAULT 'ali1688';")
        except Exception:
            pass
        try:
            cursor.execute("ALTER TABLE ali1688_sources ADD COLUMN source_channel_label VARCHAR(255) DEFAULT '1688 货源渠道';")
        except Exception:
            pass
        try:
            cursor.execute("ALTER TABLE ali1688_sources ADD COLUMN source_account_id VARCHAR(100) DEFAULT '';")
        except Exception:
            pass
        try:
            cursor.execute("ALTER TABLE ali1688_sources ADD COLUMN source_account_label VARCHAR(255) DEFAULT '';")
        except Exception:
            pass
        for column_name, ddl in (
            ("pickup_48h_text", "ALTER TABLE ali1688_sources ADD COLUMN pickup_48h_text VARCHAR(64) DEFAULT '';"),
            ("pickup_24h_text", "ALTER TABLE ali1688_sources ADD COLUMN pickup_24h_text VARCHAR(64) DEFAULT '';"),
            ("month_dispatch_text", "ALTER TABLE ali1688_sources ADD COLUMN month_dispatch_text VARCHAR(64) DEFAULT '';"),
            ("seven_day_dispatch_text", "ALTER TABLE ali1688_sources ADD COLUMN seven_day_dispatch_text VARCHAR(64) DEFAULT '';"),
            ("listing_count_text", "ALTER TABLE ali1688_sources ADD COLUMN listing_count_text VARCHAR(64) DEFAULT '';"),
            ("distributor_count_text", "ALTER TABLE ali1688_sources ADD COLUMN distributor_count_text VARCHAR(64) DEFAULT '';"),
            ("waybill_support_text", "ALTER TABLE ali1688_sources ADD COLUMN waybill_support_text VARCHAR(64) DEFAULT '';"),
            ("settled_years_text", "ALTER TABLE ali1688_sources ADD COLUMN settled_years_text VARCHAR(64) DEFAULT '';"),
            ("company_name", "ALTER TABLE ali1688_sources ADD COLUMN company_name VARCHAR(255) DEFAULT '';"),
            ("source_filter_snapshot_json", "ALTER TABLE ali1688_sources ADD COLUMN source_filter_snapshot_json TEXT DEFAULT NULL;"),
            ("page_original_index", "ALTER TABLE ali1688_sources ADD COLUMN page_original_index INT DEFAULT 0;"),
            ("month_dispatch_count", "ALTER TABLE ali1688_sources ADD COLUMN month_dispatch_count INT DEFAULT 0;"),
            ("seven_day_dispatch_count", "ALTER TABLE ali1688_sources ADD COLUMN seven_day_dispatch_count INT DEFAULT 0;"),
            ("listing_count", "ALTER TABLE ali1688_sources ADD COLUMN listing_count INT DEFAULT 0;"),
            ("distributor_count", "ALTER TABLE ali1688_sources ADD COLUMN distributor_count INT DEFAULT 0;"),
            ("is_detail_incomplete", "ALTER TABLE ali1688_sources ADD COLUMN is_detail_incomplete TINYINT(1) DEFAULT 0;"),
            ("detail_incomplete_reason", "ALTER TABLE ali1688_sources ADD COLUMN detail_incomplete_reason VARCHAR(255) DEFAULT '';"),
            ("detail_status", "ALTER TABLE ali1688_sources ADD COLUMN detail_status VARCHAR(32) DEFAULT '';"),
        ):
            try:
                cursor.execute(ddl)
            except Exception:
                pass
        try:
            cursor.execute("""
                UPDATE ali1688_sources
                SET
                    is_detail_incomplete = 1,
                    detail_incomplete_reason = '历史数据：详情页未完整抓取，可能受风控/验证码影响',
                    detail_status = 'failed'
                WHERE COALESCE(is_detail_incomplete, 0) = 0
                  AND COALESCE(sku_count, 0) = 0
                  AND (images IS NULL OR TRIM(images) = '' OR TRIM(images) = '[]')
            """)
        except Exception:
            pass
            
        # 4. 自愈修改 publish_status 的 ENUM 增加本地选品/下架/删除状态
        try:
            cursor.execute("ALTER TABLE xianyu_published_items MODIFY COLUMN publish_status ENUM('selected','pending','success','failed','depublished','deleted') DEFAULT 'pending';")
        except Exception as alter_err:
            logger.warning(f"[DB] Failed to modify publish_status enum: {alter_err}")
            
        # 5. 自愈添加 input_type 字段以支持任务类型的细化展示
        try:
            cursor.execute("ALTER TABLE tasks ADD COLUMN input_type VARCHAR(20) DEFAULT 'keyword';")
        except Exception:
            pass
        # 6. 自愈添加 total_tokens 字段以支持 Token 的计量
        try:
            cursor.execute("ALTER TABLE tasks ADD COLUMN total_tokens INT DEFAULT 0;")
        except Exception:
            pass
        # 任务启动时的商品爬取配置快照。任务执行读取快照，不受后续系统默认配置变更影响。
        try:
            cursor.execute("ALTER TABLE tasks ADD COLUMN crawl_config_snapshot_json TEXT DEFAULT NULL;")
        except Exception:
            pass
        # 7. 物理刷新历史数据，防止 NULL 导致前端 React 渲染 crash
        try:
            cursor.execute("UPDATE tasks SET total_tokens = 0 WHERE total_tokens IS NULL;")
            cursor.execute("UPDATE tasks SET input_type = 'keyword' WHERE input_type IS NULL;")
        except Exception as update_err:
            logger.warning(f"[DB] Failed to refresh historical tasks null values: {update_err}")
            
        # 8. 自愈创建 system_configs 配置表，并进行旧配置数据的零感自动入库
        try:
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS system_configs (
                    cfg_key VARCHAR(50) PRIMARY KEY,
                    cfg_value TEXT NOT NULL
                ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
            """)
            
            cursor.execute("SELECT COUNT(*) AS total FROM system_configs")
            row = cursor.fetchone()
            total = row["total"] if isinstance(row, dict) else row[0]
            if total == 0:
                # 首次运行，将本地 config.json 的配置同步到 DB 里
                from xianyu_tools.config import settings
                llm_data = settings.get("llm")
                if llm_data:
                    cursor.execute("INSERT INTO system_configs (cfg_key, cfg_value) VALUES (%s, %s)", ("llm", json.dumps(llm_data)))
                openapi_data = settings.get("openapi")
                if openapi_data:
                    cursor.execute("INSERT INTO system_configs (cfg_key, cfg_value) VALUES (%s, %s)", ("openapi", json.dumps(openapi_data)))
                crawl_data = settings.get("crawl")
                if crawl_data:
                    cursor.execute("INSERT INTO system_configs (cfg_key, cfg_value) VALUES (%s, %s)", ("crawl", json.dumps(crawl_data)))
                source_channels_data = settings.get("source_channels")
                if source_channels_data:
                    cursor.execute("INSERT INTO system_configs (cfg_key, cfg_value) VALUES (%s, %s)", ("source_channels", json.dumps(source_channels_data)))
                logger.info("[DB] Initial configurations successfully synchronized to system_configs table.")
        except Exception as config_db_err:
            logger.error(f"[DB] Failed to initialize system_configs table or sync initial data: {config_db_err}")
            
        conn.commit()
        conn.close()
        logger.info("[DB] Schema initialization complete (ensured tasks.input_type, total_tokens, and fixed NULL values).")
    except Exception as e:
        logger.error(f"[DB] Schema initialization failed: {e}")

# 执行自动建表自愈
init_db_schema()

def sanitize_dir_name(name: str) -> str:
    clean = re.sub(r'\s+', '', str(name))
    return re.sub(r'[\\/:*?"<>|]', '_', clean).strip()[:60]


def _parse_json_dict(raw_value, fallback=None) -> dict:
    fallback = fallback if isinstance(fallback, dict) else {}
    if isinstance(raw_value, dict):
        return raw_value
    if not raw_value:
        return fallback
    try:
        parsed = json.loads(raw_value)
        return parsed if isinstance(parsed, dict) else fallback
    except Exception:
        return fallback


def build_task_crawl_config_snapshot(raw_cfg: dict | None = None) -> dict:
    cfg = raw_cfg if isinstance(raw_cfg, dict) else settings.get_crawl_config()
    return settings.normalize_crawl_config(
        cfg,
        source_channels_cfg=settings.get_source_channels_config(),
    )

def _clean_html_span(text: str) -> str:
    if not text: return ""
    text = re.sub(r'<span[^>]*?>.*?</span>', '', text, flags=re.IGNORECASE | re.DOTALL)
    text = re.sub(r'<[^>]+>', '', text)
    import html
    text = html.unescape(text)
    text = text.replace(">", ";")
    parts = [p.strip() for p in text.split(";") if p.strip()]
    return ";".join(parts)

DEFAULT_OPENAPI_ACCOUNT = {
    "id": "account-1",
    "name": "闲鱼账号 1",
    "base_url": "https://open.goofish.pro",
    "appid": "",
    "app_secret": "",
    "state_file": "xianyu_state_account-1.json",
    "default_config": {
        "user_name": "",
        "province": 110000,
        "city": 110100,
        "district": 110101,
        "item_biz_type": 2,
        "sp_biz_type": 2,
        "channel_cat_id": "",
        "stuff_status": 100,
        "express_fee": 0,
    },
}

def normalize_openapi_multi_account(raw_cfg: dict | None) -> dict:
    raw_cfg = dict(raw_cfg or {})
    accounts = raw_cfg.get("accounts")
    if isinstance(accounts, list) and accounts:
        normalized_accounts = []
        for index, account in enumerate(accounts, start=1):
            merged = json.loads(json.dumps(DEFAULT_OPENAPI_ACCOUNT, ensure_ascii=False))
            incoming = dict(account or {})
            merged.update({k: v for k, v in incoming.items() if k != "default_config"})
            merged_default = dict(DEFAULT_OPENAPI_ACCOUNT["default_config"])
            merged_default.update(dict(incoming.get("default_config") or {}))
            merged["default_config"] = merged_default
            merged["id"] = merged.get("id") or f"account-{index}"
            merged["name"] = merged.get("name") or f"闲鱼账号 {index}"
            merged["state_file"] = merged.get("state_file") or f"xianyu_state_{merged['id']}.json"
            normalized_accounts.append(merged)
        active_account_id = raw_cfg.get("active_account_id") or normalized_accounts[0]["id"]
        if not any(item["id"] == active_account_id for item in normalized_accounts):
            active_account_id = normalized_accounts[0]["id"]
        return {"active_account_id": active_account_id, "accounts": normalized_accounts}

    merged = json.loads(json.dumps(DEFAULT_OPENAPI_ACCOUNT, ensure_ascii=False))
    merged.update({k: v for k, v in raw_cfg.items() if k != "default_config"})
    merged_default = dict(DEFAULT_OPENAPI_ACCOUNT["default_config"])
    merged_default.update(dict(raw_cfg.get("default_config") or {}))
    merged["default_config"] = merged_default
    merged["state_file"] = raw_cfg.get("state_file") or "xianyu_state.json"
    return {"active_account_id": merged["id"], "accounts": [merged]}

def get_openapi_account(raw_cfg: dict | None, account_id: str | None = None) -> dict:
    normalized = normalize_openapi_multi_account(raw_cfg)
    target_id = account_id or normalized.get("active_account_id")
    selected = next((item for item in normalized["accounts"] if item["id"] == target_id), None)
    return selected or normalized["accounts"][0]


def get_request_openapi_account_id(payload: dict | None = None) -> str | None:
    if not isinstance(payload, dict):
        return None
    account_id = payload.get("account_id")
    return str(account_id).strip() if account_id else None


def hydrate_source_channel_account(channel_cfg: dict | None, account_cfg: dict | None) -> tuple[dict, dict]:
    channel = dict(channel_cfg or {})
    account = dict(account_cfg or {})
    runtime_cfg = build_source_channel_account_runtime(
        channel.get("channel_type"),
        channel.get("channel_id"),
        account.get("account_id"),
    )
    account.update(runtime_cfg)
    return channel, account


def get_source_channel_account_error(channel_cfg: dict | None, account_cfg: dict | None) -> str | None:
    channel = dict(channel_cfg or {})
    account = dict(account_cfg or {})
    channel_label = channel.get("label") or "当前货源渠道"
    account_label = account.get("label") or "当前渠道账号"

    if not channel.get("channel_id"):
        return "未找到货源渠道配置"
    if channel.get("enabled", True) is False:
        return f"{channel_label}已停用"
    if not account.get("account_id"):
        return f"{channel_label}下未配置可用账号"
    if account.get("enabled", True) is False:
        return f"{account_label}已停用"
    return None


def source_channel_supports_session_state(channel_cfg: dict | None) -> bool:
    channel = dict(channel_cfg or {})
    return channel.get("channel_type") == "ali1688"


def build_source_channel_unavailable_report(
    channel_cfg: dict | None,
    account_cfg: dict | None,
    error_message: str | None = None,
    status_text: str | None = None,
) -> dict:
    error_message = error_message or get_source_channel_account_error(channel_cfg, account_cfg) or "未配置可用账号"
    status_text = status_text or "未配置可用账号"
    return {
        "is_usable": False,
        "is_logged_in": False,
        "requires_verification": False,
        "account_name": "",
        "status_text": status_text,
        "last_checked_at": datetime.now().isoformat(),
        "error_message": error_message,
        "meta": {
            "channel_id": (channel_cfg or {}).get("channel_id") or "",
            "account_id": (account_cfg or {}).get("account_id") or "",
        },
    }


def resolve_runtime_path(raw_path: str | None) -> Path | None:
    if not raw_path:
        return None
    path = Path(str(raw_path).strip()).expanduser()
    if not path.is_absolute():
        path = (BASE_DIR / path).resolve()
    return path


def build_ali1688_session_report_path(state_file: str | None) -> Path | None:
    state_path = resolve_runtime_path(state_file)
    if not state_path:
        return None
    return state_path.with_name("session_report.json")


def load_ali1688_session_report_cache(state_file: str | None) -> dict:
    report_path = build_ali1688_session_report_path(state_file)
    if not report_path or not report_path.exists():
        return {}

    try:
        from xianyu_tools.ali1688_session import looks_like_real_account_name, normalize_account_name

        payload = json.loads(report_path.read_text(encoding="utf-8"))
        if not isinstance(payload, dict):
            return {}
        cached_name = normalize_account_name(payload.get("account_name"))
        if cached_name and not looks_like_real_account_name(cached_name):
            return {}
        payload["account_name"] = cached_name
        payload.setdefault("meta", {})
        payload["meta"]["cache_file"] = str(report_path)
        return payload
    except Exception as exc:
        return {
            "account_name": "",
            "error_message": str(exc),
            "meta": {"cache_file": str(report_path)},
        }


def save_ali1688_session_report_cache(state_file: str | None, report: dict | None) -> None:
    report_path = build_ali1688_session_report_path(state_file)
    if not report_path or not isinstance(report, dict):
        return

    account_name = str(report.get("account_name") or "").strip()
    if not account_name:
        return

    payload = {
        "account_name": account_name,
        "captured_at": report.get("last_checked_at") or datetime.now().isoformat(),
        "source": report.get("account_name_source") or report.get("source") or "realtime",
        "confidence": report.get("account_name_confidence") or report.get("confidence") or 0.0,
        "status_text": report.get("status_text") or "",
        "is_usable": bool(report.get("is_usable")),
        "is_logged_in": bool(report.get("is_logged_in") or report.get("is_usable")),
        "requires_verification": bool(report.get("requires_verification")),
        "state_file": str(resolve_runtime_path(state_file) or ""),
        "url": ((report.get("meta") or {}).get("url") or ""),
        "title": ((report.get("meta") or {}).get("title") or ""),
    }
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def inspect_ali1688_state_file_quick(state_file: str | None, fallback_account_name: str | None = None) -> dict:
    fallback_account_name = str(fallback_account_name or "").strip()
    if not state_file:
        return {
            "is_usable": False,
            "is_logged_in": False,
            "requires_verification": False,
            "account_name": fallback_account_name,
            "status_text": "未配置状态文件",
            "last_checked_at": datetime.now().isoformat(),
            "error_message": "",
            "meta": {"state_file": ""},
        }

    state_path = Path(state_file or "").expanduser()
    if not state_path.is_absolute():
        state_path = (BASE_DIR / state_path).resolve()

    if not state_path.exists():
        return {
            "is_usable": False,
            "is_logged_in": False,
            "requires_verification": False,
            "account_name": fallback_account_name,
            "status_text": "未配置状态文件" if not state_file else "状态文件不存在",
            "last_checked_at": datetime.now().isoformat(),
            "error_message": "" if not state_file else f"未找到状态文件：{state_file}",
            "meta": {"state_file": str(state_path)},
        }

    try:
        with open(state_path, "r", encoding="utf-8") as f:
            state_data = json.load(f)
        cookies = state_data.get("cookies") or []
        cookie_names = {str(item.get("name") or "") for item in cookies if isinstance(item, dict)}
        useful_cookie_names = {"cookie2", "_m_h5_tk", "_m_h5_tk_enc", "ali_apache_id", "cna"}
        has_useful_cookie = bool(cookie_names.intersection(useful_cookie_names))

        account_name = ""
        for item in cookies:
            if not isinstance(item, dict):
                continue
            name = str(item.get("name") or "")
            if name in {"tracknick", "_w_tb_nick", "loginId", "cn"}:
                raw_value = str(item.get("value") or "").strip()
                if raw_value:
                    account_name = unquote(raw_value)
                    break
        if not account_name and fallback_account_name:
            account_name = fallback_account_name

        is_usable = bool(cookies) and has_useful_cookie
        return {
            "is_usable": is_usable,
            "is_logged_in": is_usable,
            "requires_verification": False,
            "account_name": account_name,
            "status_text": "登录正常" if is_usable else "未检测到有效登录",
            "last_checked_at": datetime.now().isoformat(),
            "error_message": "",
            "meta": {
                "state_file": str(state_path),
                "cookie_count": len(cookies),
            },
        }
    except Exception as exc:
        return {
            "is_usable": False,
            "is_logged_in": False,
            "requires_verification": False,
            "account_name": fallback_account_name,
            "status_text": "检测失败",
            "last_checked_at": datetime.now().isoformat(),
            "error_message": str(exc),
            "meta": {"state_file": str(state_path)},
        }


def merge_ali1688_session_report(
    state_file: str | None,
    fallback_account_name: str | None = None,
    realtime_report: dict | None = None,
) -> dict:
    fallback_account_name = str(fallback_account_name or "").strip()
    state_report = inspect_ali1688_state_file_quick(state_file, fallback_account_name)
    cached_report = load_ali1688_session_report_cache(state_file)
    realtime_report = dict(realtime_report or {})

    merged = dict(realtime_report or state_report)
    merged.setdefault("is_usable", bool(state_report.get("is_usable")))
    merged.setdefault("status_text", realtime_report.get("status_text") or state_report.get("status_text") or "未检测")
    merged.setdefault("last_checked_at", realtime_report.get("last_checked_at") or state_report.get("last_checked_at") or datetime.now().isoformat())
    merged.setdefault("error_message", realtime_report.get("error_message") or state_report.get("error_message") or "")

    realtime_logged_in = bool(realtime_report.get("is_logged_in")) if "is_logged_in" in realtime_report else None
    cached_logged_in = bool(cached_report.get("is_logged_in")) if "is_logged_in" in cached_report else None
    state_logged_in = bool(state_report.get("is_logged_in")) if "is_logged_in" in state_report else None

    if realtime_logged_in is not None:
        is_logged_in = realtime_logged_in
    else:
        is_logged_in = any(
            value is True for value in (state_logged_in, cached_logged_in, bool(merged.get("is_usable")))
        )

    realtime_requires_verification = (
        bool(realtime_report.get("requires_verification"))
        if "requires_verification" in realtime_report
        else None
    )
    cached_requires_verification = (
        bool(cached_report.get("requires_verification"))
        if "requires_verification" in cached_report
        else None
    )
    state_requires_verification = (
        bool(state_report.get("requires_verification"))
        if "requires_verification" in state_report
        else None
    )

    if realtime_requires_verification is not None:
        requires_verification = realtime_requires_verification
    else:
        requires_verification = any(
            value is True for value in (state_requires_verification, cached_requires_verification)
        )

    merged["is_logged_in"] = is_logged_in
    merged["requires_verification"] = requires_verification

    realtime_name = str(realtime_report.get("account_name") or "").strip()
    cached_name = str(cached_report.get("account_name") or "").strip()
    state_name = str(state_report.get("account_name") or "").strip()
    final_name = realtime_name or cached_name or state_name or fallback_account_name

    if realtime_name:
        name_source = str(realtime_report.get("account_name_source") or realtime_report.get("source") or "realtime")
        name_confidence = realtime_report.get("account_name_confidence") or realtime_report.get("confidence") or 0.0
    elif cached_name:
        name_source = str(cached_report.get("source") or "cache")
        name_confidence = cached_report.get("confidence") or 0.75
    elif state_name and state_name != fallback_account_name:
        name_source = "state_file_cookie"
        name_confidence = 0.6
    elif final_name:
        name_source = "label_fallback"
        name_confidence = 0.2
    else:
        name_source = ""
        name_confidence = 0.0

    merged["account_name"] = final_name
    merged["account_name_source"] = name_source
    merged["account_name_confidence"] = name_confidence
    merged["meta"] = {
        **dict(state_report.get("meta") or {}),
        **dict(cached_report.get("meta") or {}),
        **dict(merged.get("meta") or {}),
        "state_file": str(resolve_runtime_path(state_file) or ""),
        "cache_file": str(build_ali1688_session_report_path(state_file) or ""),
    }
    if cached_report.get("captured_at"):
        merged["meta"]["cached_account_name_at"] = cached_report.get("captured_at")

    return merged


async def inspect_source_channel_account(channel_type: str, account_cfg: dict) -> dict:
    channel_type = str(channel_type or "").strip().lower()
    if channel_type == "ali1688":
        user_data_dir = account_cfg.get("user_data_dir") or ""
        if user_data_dir:
            try:
                from xianyu_tools.ali1688_session import Ali1688SessionConfig, inspect_ali1688_session

                session_result = await inspect_ali1688_session(
                    Ali1688SessionConfig(
                        user_data_dir=user_data_dir,
                        profile_directory=account_cfg.get("profile_directory") or "Default",
                        headless=True,
                    )
                )
                state = session_result.get("state") or "unknown"
                is_logged_in = state in {"search", "member", "slider"}
                requires_verification = state == "slider"
                is_usable = state in {"search", "member"}
                return {
                    "is_usable": is_usable,
                    "is_logged_in": is_logged_in,
                    "requires_verification": requires_verification,
                    "account_name": session_result.get("account_name") or "",
                    "account_name_source": session_result.get("account_name_source") or "",
                    "account_name_confidence": session_result.get("account_name_confidence") or 0.0,
                    "status_text": "登录正常" if is_usable else {
                        "profile_locked": "浏览器配置被占用",
                        "login": "跳转到了登录页",
                        "slider": "已登录，需完成滑块或风控验证",
                        "timeout": "1688 页面访问超时",
                        "unknown": "未检测到可用搜索态",
                    }.get(state, "会话不可用"),
                    "last_checked_at": datetime.now().isoformat(),
                    "error_message": (
                        "当前已有 1688 登录浏览器正在使用该账号环境，请关闭现有登录窗口后再重试。"
                        if state == "profile_locked"
                        else (
                            session_result.get("error")
                            or (
                                "1688 页面访问超时，系统已尝试自动回退检测，请稍后重试。"
                                if state == "timeout"
                                else ""
                            )
                        )
                    ),
                    "meta": session_result,
                }
            except Exception as exc:
                fallback = merge_ali1688_session_report(
                    account_cfg.get("state_file"),
                    account_cfg.get("label") or account_cfg.get("account_id") or "",
                )
                fallback["error_message"] = fallback.get("error_message") or str(exc)
                return fallback
        return merge_ali1688_session_report(
            account_cfg.get("state_file"),
            account_cfg.get("label") or account_cfg.get("account_id") or "",
        )

    return {
        "is_usable": False,
        "is_logged_in": False,
        "requires_verification": False,
        "account_name": "",
        "status_text": "暂不支持该渠道检测",
        "last_checked_at": datetime.now().isoformat(),
        "error_message": "",
        "meta": {"channel_type": channel_type},
    }

# --- 任务模型 ---
class Task:
    @staticmethod
    def update(task_id: str, **kwargs):
        if not kwargs: return
        try:
            conn = get_db_conn(); cursor = conn.cursor()
            fields = [f"{k} = %s" for k in kwargs.keys()]
            values = list(kwargs.values()); values.append(task_id)
            cursor.execute(f"UPDATE tasks SET {', '.join(fields)} WHERE id = %s", tuple(values))
            conn.commit(); conn.close()
        except Exception as e:
            logger.error(f"Failed to update task {task_id}: {e}")

    @staticmethod
    def add(keyword: str, crawl_config_snapshot: dict | None = None):
        try:
            from scripts.run_xianyu_hot_items import detect_input_type
            input_type = detect_input_type(keyword)
            snapshot = build_task_crawl_config_snapshot(crawl_config_snapshot)
            snapshot_json = json.dumps(snapshot, ensure_ascii=False)
            
            task_id, now = str(uuid.uuid4())[:8], datetime.now()
            version = now.strftime("%Y%m%d")
            root_dir = str(OUTPUTS_DIR / f"{sanitize_dir_name(keyword)}_{version}")
            conn = get_db_conn(); cursor = conn.cursor()
            cursor.execute(
                """
                INSERT INTO tasks (
                    id, keyword, status, progress, msg, created_at, root_dir, version,
                    is_deleted, input_type, crawl_config_snapshot_json
                )
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, 0, %s, %s)
                """,
                (task_id, keyword, "排队中", 0, "等待调度", now, root_dir, version, input_type, snapshot_json),
            )
            conn.commit(); conn.close()
            logger.info(f"Task added: {keyword} ({task_id}), type: {input_type}")
            return task_id
        except Exception as e:
            logger.error(f"Failed to add task for {keyword}: {e}")
            return None

# --- 后台 Worker ---
async def pipeline_worker():
    logger.info("Background Worker started.")
    while True:
        task_id, root_dir = None, None
        try:
            conn = get_db_conn(); cursor = conn.cursor()
            cursor.execute("SELECT id, keyword, root_dir, crawl_config_snapshot_json FROM tasks WHERE status = '排队中' AND is_deleted = 0 ORDER BY created_at ASC LIMIT 1")
            t_data = cursor.fetchone()
            if not t_data:
                conn.close(); await asyncio.sleep(5); continue
            
            task_id, keyword, root_dir = t_data["id"], t_data["keyword"], t_data["root_dir"]
            cursor.execute("UPDATE tasks SET status = '执行中' WHERE id = %s", (task_id,))
            conn.commit(); conn.close()
            
            logger.info(f"Worker picked up task: {keyword} ({task_id})")
            
            log_file = Path(root_dir) / "task.log"
            log_file.parent.mkdir(parents=True, exist_ok=True)
            snapshot = _parse_json_dict(t_data.get("crawl_config_snapshot_json"))
            crawl_config_arg = ""
            if snapshot:
                snapshot_file = Path(root_dir) / "task_crawl_config_snapshot.json"
                snapshot_file.write_text(json.dumps(snapshot, ensure_ascii=False, indent=2), encoding="utf-8")
                crawl_config_arg = f" --crawl-config-file {shlex.quote(str(snapshot_file))}"
            
            python_path = sys.executable
            cmd = (
                f"export PYTHONPATH=$PYTHONPATH:{shlex.quote(str(BASE_DIR / 'src'))} && "
                f"{shlex.quote(python_path)} scripts/run_full_pipeline.py "
                f"--keyword {shlex.quote(keyword)} --task-id {shlex.quote(task_id)}"
                f"{crawl_config_arg} > {shlex.quote(str(log_file))} 2>&1"
            )
            
            process = await asyncio.create_subprocess_shell(cmd, cwd=str(BASE_DIR), preexec_fn=os.setsid)
            Task.update(task_id, pgid=os.getpgid(process.pid))
            
            exit_code = await process.wait()
            logger.info(f"Task {task_id} process exited with code {exit_code}")
            
            conn = get_db_conn(); cursor = conn.cursor()
            cursor.execute("SELECT status FROM tasks WHERE id = %s", (task_id,))
            db_status = (cursor.fetchone() or {}).get('status')
            conn.close()

            if db_status == "正在暂停" or exit_code != 0:
                Task.update(task_id, status="已暂停", msg="任务已停止(强制)", pgid=None)
                logger.info(f"Task {task_id} marked as Paused.")
            elif db_status == "执行中" and exit_code == 0:
                Task.update(task_id, status="已完成", progress=100, msg="分析完成", pgid=None)
                logger.info(f"Task {task_id} completed successfully.")
            else:
                Task.update(task_id, pgid=None)

        except Exception as e:
            logger.error(f"Pipeline Worker Error: {e}")
            if task_id: Task.update(task_id, status="失败", msg=str(e), pgid=None)
            await asyncio.sleep(10)
        finally:
            if task_id: running_processes.pop(task_id, None)

running_processes = {} 
@app.on_event("startup")
async def startup():
    asyncio.create_task(pipeline_worker())


def _channel_filter_snapshot_has_signal(snapshot: dict | None) -> bool:
    if not isinstance(snapshot, dict):
        return False
    if snapshot.get("mapping_stage") and snapshot.get("mapping_stage") != "snapshot_only":
        return True
    for key in (
        "configured_filter_keys",
        "configured_enabled_filter_keys",
        "enabled_filter_keys",
        "applied_filter_keys",
        "query_injected_filter_keys",
        "unapplied_filter_keys",
    ):
        if isinstance(snapshot.get(key), list) and snapshot.get(key):
            return True
    filters = snapshot.get("configured_filters")
    if not isinstance(filters, dict):
        filters = snapshot.get("filters")
    if isinstance(filters, dict) and any(bool(value) for value in filters.values()):
        return True
    return False

# --- 路由 ---
@app.get("/api/tasks")
def list_tasks():
    try:
        conn = get_db_conn(); cursor = conn.cursor()
        cursor.execute("SELECT id, keyword, status, progress, msg, created_at, version, input_type, total_tokens, crawl_config_snapshot_json FROM tasks WHERE is_deleted = 0 ORDER BY created_at DESC")
        rows = cursor.fetchall()
        task_ids = [row["id"] for row in rows if row.get("id")]
        task_channel_map = {}
        task_channel_summary_map = {}
        if task_ids:
            placeholders = ",".join(["%s"] * len(task_ids))
            cursor.execute(
                f"""
                SELECT
                    xi.task_id AS task_id,
                    COALESCE(NULLIF(src.source_channel_id, ''), 'ali1688') AS channel_id,
                    COALESCE(NULLIF(src.source_channel_type, ''), 'ali1688') AS channel_type,
                    COALESCE(NULLIF(src.source_channel_label, ''), '1688 货源渠道') AS channel_label,
                    COUNT(*) AS source_count
                FROM ali1688_sources src
                INNER JOIN xianyu_items xi ON src.item_id = xi.id
                WHERE xi.task_id IN ({placeholders})
                GROUP BY
                    xi.task_id,
                    COALESCE(NULLIF(src.source_channel_id, ''), 'ali1688'),
                    COALESCE(NULLIF(src.source_channel_type, ''), 'ali1688'),
                    COALESCE(NULLIF(src.source_channel_label, ''), '1688 货源渠道')
                ORDER BY
                    xi.task_id ASC,
                    COALESCE(NULLIF(src.source_channel_label, ''), '1688 货源渠道') ASC
                """,
                tuple(task_ids),
            )
            for channel_row in cursor.fetchall():
                task_id = channel_row.get("task_id")
                if task_id not in task_channel_map:
                    task_channel_map[task_id] = []
                task_channel_map[task_id].append(
                    {
                        "channel_id": channel_row.get("channel_id") or "ali1688",
                        "channel_type": channel_row.get("channel_type") or "ali1688",
                        "channel_label": channel_row.get("channel_label") or "1688 货源渠道",
                        "source_count": int(channel_row.get("source_count") or 0),
                    }
                )
            cursor.execute(
                f"""
                SELECT
                    xi.task_id AS task_id,
                    COALESCE(NULLIF(src.source_channel_id, ''), 'ali1688') AS channel_id,
                    COALESCE(NULLIF(src.source_channel_type, ''), 'ali1688') AS channel_type,
                    COALESCE(NULLIF(src.source_channel_label, ''), '1688 货源渠道') AS channel_label,
                    src.source_filter_snapshot_json AS source_filter_snapshot_json
                FROM ali1688_sources src
                INNER JOIN xianyu_items xi ON src.item_id = xi.id
                WHERE xi.task_id IN ({placeholders})
                ORDER BY
                    xi.task_id ASC,
                    COALESCE(NULLIF(src.source_channel_label, ''), '1688 货源渠道') ASC,
                    src.id ASC
                """,
                tuple(task_ids),
            )
            task_channel_summary_map = _build_task_channel_summary_map(
                task_channel_map=task_channel_map,
                snapshot_rows=cursor.fetchall(),
            )
        conn.close()
        for r in rows:
            if isinstance(r['created_at'], datetime): r['created_at'] = r['created_at'].strftime("%Y-%m-%d %H:%M")
            r["crawl_config_snapshot"] = _parse_json_dict(r.pop("crawl_config_snapshot_json", None))
            r["used_channels"] = task_channel_map.get(r["id"], [])
            r["channel_summaries"] = task_channel_summary_map.get(r["id"], r["used_channels"])
        return rows
    except Exception as e:
        logger.error(f"Failed to list tasks: {e}")
        return []

@app.post("/api/tasks")
def create_task(req: dict): 
    snapshot = build_task_crawl_config_snapshot(req.get("crawl_config") if isinstance(req.get("crawl_config"), dict) else None)
    tid = Task.add(req["keyword"], snapshot)
    return {"id": tid} if tid else {"error": "Failed to create task"}

@app.post("/api/tasks/{task_id}/pause")
def pause_task(task_id: str):
    logger.info(f"Pause requested for task: {task_id}")
    conn = get_db_conn(); cursor = conn.cursor()
    cursor.execute("SELECT pgid FROM tasks WHERE id = %s", (task_id,))
    row = cursor.fetchone(); conn.close()
    if row and row['pgid']:
        try:
            os.killpg(int(row['pgid']), signal.SIGKILL)
            Task.update(task_id, status="已暂停", msg="任务已停止", pgid=None)
            logger.info(f"Sent SIGKILL to process group {row['pgid']} for task {task_id}")
        except Exception as e:
            logger.warning(f"Failed to kill process group for task {task_id}: {e}")
            Task.update(task_id, status="已暂停", msg="进程已不存在", pgid=None)
    else:
        Task.update(task_id, status="已暂停", msg="任务取消", pgid=None)
    return {"status": "ok"}

@app.post("/api/tasks/{task_id}/retry")
def retry_task(task_id: str):
    logger.info(f"Retry/Resume requested for task: {task_id}")
    conn = get_db_conn(); cursor = conn.cursor()
    cursor.execute("SELECT keyword, status, created_at, crawl_config_snapshot_json FROM tasks WHERE id = %s", (task_id,))
    row = cursor.fetchone()
    if not row: conn.close(); return {"error": "Not found"}
    is_today = row['created_at'].date() == datetime.now().date()
    if row['status'] == '已完成':
        if is_today:
            logger.info(f"Performing same-day overwrite for task {task_id}")
            
            # 物理清理对应的本地 HTML 文件
            try:
                cursor.execute("SELECT html_path FROM ali1688_sources WHERE task_id = %s", (task_id,))
                sources = cursor.fetchall()
                for s in sources:
                    if s.get("html_path"):
                        p = BASE_DIR / s["html_path"] if not Path(s["html_path"]).is_absolute() else Path(s["html_path"])
                        if p.exists() and p.is_file():
                            p.unlink()
                            logger.info(f"[Cleanup] Physically deleted local HTML: {p}")
            except Exception as cleanup_err:
                logger.error(f"[Cleanup] Failed to clean HTML files for task {task_id}: {cleanup_err}")

            cursor.execute("DELETE FROM xianyu_items WHERE task_id = %s", (task_id,))
            cursor.execute("DELETE FROM ali1688_sources WHERE task_id = %s", (task_id,))
            cursor.execute("UPDATE tasks SET status = '排队中', msg = '同日重扫中...', progress = 0, pgid = NULL, checkpoint = NULL WHERE id = %s", (task_id,))
            conn.commit(); conn.close()
            return {"status": "ok", "action": "overwritten_today"}
        else:
            logger.info(f"Creating new version for legacy task {task_id}")
            snapshot = _parse_json_dict(row.get("crawl_config_snapshot_json"))
            conn.close(); new_id = Task.add(row['keyword'], snapshot if snapshot else None)
            return {"status": "ok", "new_id": new_id, "action": "created_new_day"}
    else:
        logger.info(f"Resuming task {task_id} from checkpoint")
        cursor.execute("UPDATE tasks SET status = '排队中', msg = '准备恢复...', pgid = NULL WHERE id = %s", (task_id,))
        conn.commit(); conn.close(); return {"status": "ok", "action": "resumed"}

@app.delete("/api/tasks/{task_id}")
def delete_task(task_id: str):
    logger.info(f"Full delete requested for task: {task_id}")
    pause_task(task_id)
    conn = get_db_conn(); cursor = conn.cursor()
    
    # 物理清理对应的本地 HTML 文件（双重保障）
    try:
        cursor.execute("SELECT html_path FROM ali1688_sources WHERE task_id = %s", (task_id,))
        sources = cursor.fetchall()
        for s in sources:
            if s.get("html_path"):
                p = BASE_DIR / s["html_path"] if not Path(s["html_path"]).is_absolute() else Path(s["html_path"])
                if p.exists() and p.is_file():
                    p.unlink()
                    logger.info(f"[Cleanup] Physically deleted local HTML: {p}")
    except Exception as cleanup_err:
        logger.warning(f"[Cleanup] Failed to clean HTML files for deleted task {task_id}: {cleanup_err}")

    cursor.execute("SELECT root_dir FROM tasks WHERE id = %s", (task_id,))
    row = cursor.fetchone()
    if row and row['root_dir']:
        folder_path = Path(row['root_dir'])
        if folder_path.exists() and "outputs" in folder_path.parts:
            try:
                import shutil
                shutil.rmtree(folder_path)
                logger.info(f"Physical folder deleted: {folder_path}")
            except Exception as e:
                logger.error(f"Failed to delete folder {folder_path}: {e}")
                
    # 从爆款和 1688 货源表中硬删除数据
    try:
        cursor.execute("DELETE FROM xianyu_items WHERE task_id = %s", (task_id,))
        cursor.execute("DELETE FROM ali1688_sources WHERE task_id = %s", (task_id,))
        conn.commit()
        logger.info(f"[DB] Cleared database records for task {task_id}")
    except Exception as db_err:
        logger.error(f"[DB] Failed to clear records for task {task_id}: {db_err}")

    Task.update(task_id, status="已删除", msg="任务及物理文件已清理", is_deleted=1, pgid=None)
    conn.close(); return {"status": "ok"}

@app.get("/api/task_details/{task_id}")
def get_task_details(task_id: str):
    try:
        conn = get_db_conn(); cursor = conn.cursor()
        crawl_cfg = settings.get_crawl_config()
        gross_profit_rate = _normalize_gross_profit_rate(crawl_cfg.get("gross_profit_rate"))
        # 1. 找到该任务下的所有闲鱼爆款
        cursor.execute("SELECT * FROM xianyu_items WHERE task_id = %s ORDER BY rank_index ASC", (task_id,))
        db_items = cursor.fetchall()
        
        details = []
        for item in db_items:
            item_db_id = item['id'] # 闲鱼商品的唯一主键
            # 2. 根据该主键去 1688 货源表里捞数据
            cursor.execute("""
                SELECT * FROM ali1688_sources
                WHERE item_id = %s
                ORDER BY
                    COALESCE(NULLIF(source_channel_label, ''), '1688 货源渠道') ASC,
                    min_price ASC
            """, (item_db_id,))
            sources_rows = cursor.fetchall()
            source_ids = [s['id'] for s in sources_rows]
            latest_status_map = {}

            if source_ids:
                placeholders = ",".join(["%s"] * len(source_ids))
                cursor.execute(
                    f"""
                    SELECT p.source_db_id, p.publish_status, p.published_url
                    FROM xianyu_published_items p
                    INNER JOIN (
                        SELECT source_db_id, MAX(created_at) AS max_time
                        FROM xianyu_published_items
                        WHERE source_db_id IN ({placeholders})
                        GROUP BY source_db_id
                    ) latest
                        ON p.source_db_id = latest.source_db_id
                       AND p.created_at = latest.max_time
                    """,
                    tuple(source_ids),
                )
                latest_status_map = {
                    row["source_db_id"]: {
                        "publish_status": row["publish_status"],
                        "published_url": row.get("published_url"),
                    }
                    for row in cursor.fetchall()
                }
            
            sources_data = []
            channel_groups_map = {}
            used_channels_map = {}
            for s in sources_rows:
                # 安全解析图片 JSON
                try: imgs = json.loads(s['images']) if s['images'] else []
                except: imgs = []
                raw_filter_snapshot_json = s.get('source_filter_snapshot_json')
                has_recorded_filter_snapshot = bool(str(raw_filter_snapshot_json or "").strip())
                try:
                    filter_snapshot = json.loads(raw_filter_snapshot_json or "{}")
                    if not isinstance(filter_snapshot, dict):
                        filter_snapshot = {}
                except Exception:
                    filter_snapshot = {}
                filter_snapshot = settings.normalize_channel_search_filter_snapshot(
                    filter_snapshot,
                    channel_id=s.get('source_channel_id') or 'ali1688',
                    channel_type='ali1688',
                )
                latest_status = latest_status_map.get(s['id'], {})
                channel_id = s.get('source_channel_id') or 'ali1688'
                channel_type = s.get('source_channel_type') or 'ali1688'
                channel_label = s.get('source_channel_label') or '1688 货源渠道'
                account_id = s.get('source_account_id') or ''
                account_label = s.get('source_account_label') or ''
                filter_snapshot = settings.normalize_channel_search_filter_snapshot(
                    filter_snapshot,
                    channel_id=channel_id,
                    channel_type=channel_type,
                )
                
                source_row = {
                    "db_id": s['id'],
                    "title": s['title'],
                    "min_price": float(s['min_price']) if s['min_price'] else 0,
                    "sku_count": s['sku_count'],
                    "url": s['source_url'],
                    "images": imgs,
                    "drop_reason": s['drop_reason'],
                    "source_channel_id": channel_id,
                    "source_channel_type": channel_type,
                    "source_channel_label": channel_label,
                    "source_account_id": account_id,
                    "source_account_label": account_label,
                    "pickup_48h_text": s.get('pickup_48h_text') or '',
                    "pickup_24h_text": s.get('pickup_24h_text') or '',
                    "month_dispatch_text": s.get('month_dispatch_text') or '',
                    "seven_day_dispatch_text": s.get('seven_day_dispatch_text') or '',
                    "listing_count_text": s.get('listing_count_text') or '',
                    "distributor_count_text": s.get('distributor_count_text') or '',
                    "waybill_support_text": s.get('waybill_support_text') or '',
                    "settled_years_text": s.get('settled_years_text') or '',
                    "company_name": s.get('company_name') or '',
                    "page_original_index": int(s.get('page_original_index') or 0),
                    "month_dispatch_count": int(s.get('month_dispatch_count') or 0),
                    "seven_day_dispatch_count": int(s.get('seven_day_dispatch_count') or 0),
                    "listing_count": int(s.get('listing_count') or 0),
                    "distributor_count": int(s.get('distributor_count') or 0),
                    "is_detail_incomplete": bool(s.get('is_detail_incomplete')),
                    "detail_incomplete_reason": s.get('detail_incomplete_reason') or '',
                    "detail_status": s.get('detail_status') or '',
                    "source_filter_snapshot": filter_snapshot,
                    "source_filter_summary": _summarize_channel_filter_snapshot(
                        filter_snapshot,
                        channel_type=channel_type,
                        has_recorded_snapshot=has_recorded_filter_snapshot,
                    ),
                    "has_recorded_filter_snapshot": has_recorded_filter_snapshot,
                    "publish_status": latest_status.get("publish_status", "none"),
                    "published_url": latest_status.get("published_url", ""),
                }
                source_row["estimated_profit"] = _compute_source_estimated_profit(source_row, gross_profit_rate)
                sources_data.append(source_row)

                if channel_id not in channel_groups_map:
                    channel_groups_map[channel_id] = {
                        "channel_id": channel_id,
                        "channel_type": channel_type,
                        "channel_label": channel_label,
                        "source_count": 0,
                        "account_ids": [],
                        "account_labels": [],
                        "source_filter_snapshot": settings.normalize_channel_search_filter_snapshot(
                            {},
                            channel_id=channel_id,
                            channel_type=channel_type,
                        ),
                        "filter_summary": _summarize_channel_filter_snapshot(
                            {},
                            channel_type=channel_type,
                            has_recorded_snapshot=False,
                        ),
                        "has_recorded_filter_snapshot": False,
                        "sources": [],
                    }
                group = channel_groups_map[channel_id]
                group["sources"].append(source_row)
                group["source_count"] += 1
                group["has_recorded_filter_snapshot"] = bool(
                    group.get("has_recorded_filter_snapshot") or has_recorded_filter_snapshot
                )
                if (
                    _channel_filter_snapshot_has_signal(filter_snapshot)
                    and not _channel_filter_snapshot_has_signal(group.get("source_filter_snapshot"))
                ):
                    group["source_filter_snapshot"] = filter_snapshot
                if account_id and account_id not in group["account_ids"]:
                    group["account_ids"].append(account_id)
                if account_label and account_label not in group["account_labels"]:
                    group["account_labels"].append(account_label)

                if channel_id not in used_channels_map:
                    used_channels_map[channel_id] = {
                        "channel_id": channel_id,
                        "channel_type": channel_type,
                        "channel_label": channel_label,
                    }

            sorted_sources, sorted_channel_groups, sorted_used_channels = _sort_detail_sources_and_groups(
                source_rows=sources_data,
                channel_groups_map=channel_groups_map,
                used_channels_map=used_channels_map,
                gross_profit_rate=gross_profit_rate,
            )
            for group in sorted_channel_groups:
                group["filter_summary"] = _summarize_channel_filter_snapshot(
                    group.get("source_filter_snapshot"),
                    channel_type=group.get("channel_type"),
                    has_recorded_snapshot=bool(group.get("has_recorded_filter_snapshot")),
                )
            
            details.append({
                "rank": item['rank_index'],
                "xianyu_item": {
                    "db_id": item_db_id,
                    "title": item['title'],
                    "price": float(item['price']),
                    "image_url": item['image_url'],
                    "want_count": item['want_count'],
                    "item_url": item['item_url']
                },
                "source_sort_strategy": _build_detail_source_sort_strategy(),
                "channel_group_sort_strategy": _build_detail_channel_group_sort_strategy(),
                "sources": sorted_sources,
                "used_channels": sorted_used_channels,
                "channel_groups": sorted_channel_groups,
            })
            
        conn.close()
        return {"task_id": task_id, "details": details}
    except Exception as e:
        logger.error(f"Failed to fetch details for task {task_id}: {e}")
        return {"error": str(e)}

def load_source_skus_from_excel(source_id: int) -> list:
    """从本地 excel 中加载商品的默认规格数据"""
    try:
        conn = get_db_conn(); cursor = conn.cursor()
        cursor.execute("SELECT * FROM ali1688_sources WHERE id = %s", (source_id,))
        source = cursor.fetchone()
        if not source:
            conn.close()
            return []
            
        task_id = source['task_id']
        item_id = source['item_id']
        offer_id = source['offer_id']
        
        cursor.execute("SELECT rank_index, title FROM xianyu_items WHERE id = %s", (item_id,))
        item = cursor.fetchone()
        if not item:
            conn.close()
            return []
            
        rank = item['rank_index']
        item_title = item['title']
        
        cursor.execute("SELECT root_dir FROM tasks WHERE id = %s", (task_id,))
        task = cursor.fetchone()
        root_dir_db = task['root_dir'] if task else None
        conn.close()
        
        skus = []
        if root_dir_db:
            root_path = Path(root_dir_db)
            if not root_path.is_absolute():
                root_path = BASE_DIR / root_path
                
            clean_title = sanitize_dir_name(item_title)
            dir_name = f"Rank_{rank}_{clean_title}"
            source_dir = root_path / dir_name
            
            # 兼容相对路径回退
            if not source_dir.exists():
                rel_path = Path(root_dir_db).name
                source_dir = BASE_DIR / "outputs" / rel_path / dir_name
                
            if source_dir.exists():
                xlsx_files = list(source_dir.glob(f"*_{offer_id}.xlsx"))
                if xlsx_files:
                    from openpyxl import load_workbook
                    wb = load_workbook(filename=xlsx_files[0], read_only=True)
                    ws = wb.active
                    for idx, r in enumerate(ws.iter_rows(values_only=True)):
                        if idx == 0:
                            continue
                        if not r or len(r) < 2 or r[0] is None:
                            continue
                        img_val = str(r[4]) if len(r) > 4 and r[4] is not None else ""
                        is_valid_img = img_val.startswith("http") or img_val.startswith("//") or "alicdn.com" in img_val
                        skus.append({
                            "sku_text": _clean_html_span(str(r[0])),
                            "price": float(r[1]) if r[1] is not None else 0.0,
                            "stock": int(r[2]) if r[2] is not None else 0,
                            "spec_id": str(r[3]) if len(r) > 3 and r[3] is not None else "",
                            "image": img_val if is_valid_img else ""
                        })
                    wb.close()
        return skus
    except Exception as e:
        logger.error(f"Failed to load skus from excel for source {source_id}: {e}")
        return []


def load_source_skus_from_db(source_id: int):
    cached = SOURCE_SKUS_CACHE.get(source_id)
    if cached is not None:
        return cached

    try:
        conn = get_db_conn()
        cursor = conn.cursor()
        cursor.execute("SELECT sku_text, price, stock, spec_id, image FROM ali1688_skus WHERE source_id = %s", (source_id,))
        rows = cursor.fetchall()
        conn.close()
        
        if rows:
            skus = []
            for r in rows:
                skus.append({
                    "sku_text": _clean_html_span(r["sku_text"]),
                    "price": float(r["price"]),
                    "stock": int(r["stock"]),
                    "spec_id": r["spec_id"],
                    "image": r["image"]
                })
            SOURCE_SKUS_CACHE[source_id] = skus
            return skus
    except Exception as e:
        logger.error(f"Failed to load skus from DB for source {source_id}: {e}")
        
    # 兼容回退读取 Excel 物理文件
    logger.warning(f"[Fallback] DB skus empty or failed for source {source_id}. Loading from Excel...")
    skus = load_source_skus_from_excel(source_id)
    SOURCE_SKUS_CACHE[source_id] = skus
    return skus


@app.get("/api/source_skus/{source_id}")
def get_source_skus(source_id: int):
    logger.info(f"Fetching SKU list for source_id: {source_id}")
    skus = load_source_skus_from_db(source_id)
    return {"skus": skus}


@app.post("/api/selection/batch")
async def batch_add_to_selection(req: dict = {}):
    logger.info(f"Selection batch add request: {req}")
    source_ids = req.get("source_ids", [])
    if not source_ids:
        return {"success": [], "failed": [{"source_id": 0, "msg": "未选中任何货源"}]}

    conn = get_db_conn(); cursor = conn.cursor()
    success_list = []
    failed_list = []

    for sid in source_ids:
        cursor.execute("""
            SELECT s.id, s.task_id, xi.item_url
            FROM ali1688_sources s
            LEFT JOIN xianyu_items xi ON s.item_id = xi.id
            WHERE s.id = %s
        """, (sid,))
        source = cursor.fetchone()
        if not source:
            failed_list.append({"source_id": sid, "msg": "货源数据在数据库中不存在"})
            continue
        ref_xianyu_item_id = _extract_xianyu_item_id(source.get("item_url"))

        cursor.execute("""
            SELECT id, publish_status, xianyu_item_id
            FROM xianyu_published_items
            WHERE source_db_id = %s
            ORDER BY created_at DESC LIMIT 1
        """, (sid,))
        latest = cursor.fetchone()
        latest_status = (latest or {}).get("publish_status")

        if latest_status in ("selected", "success"):
            latest_xianyu_item_id = (latest or {}).get("xianyu_item_id")
            if latest_status == "selected" and not latest_xianyu_item_id and ref_xianyu_item_id:
                cursor.execute(
                    "UPDATE xianyu_published_items SET xianyu_item_id = %s WHERE id = %s",
                    (ref_xianyu_item_id, latest["id"]),
                )
                latest_xianyu_item_id = ref_xianyu_item_id
            success_list.append({
                "source_id": sid,
                "status": latest_status,
                "xianyu_item_id": latest_xianyu_item_id,
            })
            continue

        cursor.execute("""
            INSERT INTO xianyu_published_items (task_id, source_db_id, xianyu_item_id, publish_status, publish_msg, published_url)
            VALUES (%s, %s, %s, %s, %s, %s)
        """, (source["task_id"], sid, ref_xianyu_item_id, "selected", "已加入选品", None))
        success_list.append({"source_id": sid, "status": "selected", "xianyu_item_id": ref_xianyu_item_id})

    conn.commit()
    conn.close()
    return {"success": success_list, "failed": failed_list}


@app.post("/api/publish/batch")
async def batch_publish_to_xianyu(req: dict = {}):
    logger.info(f"OpenAPI batch publish request: {req}")
    source_ids = req.get("source_ids", [])
    account_id = get_request_openapi_account_id(req)
    if not source_ids:
        return {"success": [], "failed": [{"source_id": 0, "msg": "未选中任何商品"}]}

    custom_configs = req.get("custom_configs", {})
    conn = get_db_conn(); cursor = conn.cursor()

    items_to_publish = []
    items_to_republish = []
    failed_items = []

    for sid in source_ids:
        cursor.execute("SELECT * FROM ali1688_sources WHERE id = %s", (sid,))
        source = cursor.fetchone()
        if not source:
            failed_items.append({"source_id": sid, "msg": "货源数据在数据库中不存在"})
            continue

        cursor.execute("""
            SELECT publish_status
            FROM xianyu_published_items
            WHERE source_db_id = %s
            ORDER BY created_at DESC
            LIMIT 1
        """, (sid,))
        latest_publish = cursor.fetchone()
        latest_status = latest_publish.get("publish_status") if latest_publish else None
        if latest_status == "deleted":
            failed_items.append({"source_id": sid, "msg": "商品已删除，无法重新上架，请重新加入选品后发布"})
            continue

        if latest_status != "success":
            cursor.execute("""
                SELECT xianyu_item_id, task_id
                FROM xianyu_published_items
                WHERE source_db_id = %s AND publish_status = 'depublished' AND xianyu_item_id IS NOT NULL
                ORDER BY created_at DESC
                LIMIT 1
            """, (sid,))
            depublished_publish = cursor.fetchone()
            if depublished_publish and depublished_publish.get("xianyu_item_id"):
                items_to_republish.append({
                    "source_id": sid,
                    "product_id": depublished_publish["xianyu_item_id"],
                    "task_id": depublished_publish.get("task_id") or source["task_id"],
                })
                continue

        images = json.loads(source['images'] or "[]")
        custom_info = custom_configs.get(str(sid)) or {}
        custom_title = custom_info.get("title") or source['title']

        # 获取默认规格并进行加价 30 自愈处理
        skus = load_source_skus_from_db(sid)
        sku_items = []
        sku_images = []

        if skus:
            for s in skus:
                sku_items.append({
                    "sku_text": s["sku_text"],
                    "price": round(s["price"] + 30.0, 2),
                    "stock": min(9999, int(s["stock"]) or 1)
                })
            
            seen_img_skus = set()
            for s in skus:
                if s.get("image"):
                    first_attr = s["sku_text"].split(';')[0]
                    if first_attr not in seen_img_skus:
                        seen_img_skus.add(first_attr)
                        sku_images.append({
                            "src": s["image"],
                            "width": 800,
                            "height": 800,
                            "sku_text": first_attr
                        })
            
            final_price = min(s['price'] for s in sku_items)
        else:
            custom_price = custom_info.get("price")
            if custom_price is not None:
                final_price = float(custom_price)
            else:
                final_price = float(source['min_price']) + 30

        item_data = {
            "source_id": sid,
            "title": custom_title[:60],
            "description": f"【精选货源】\n{custom_title}\n品质保障，欢迎选购。",
            "price": final_price,
            "images": images,
        }
        if sku_items:
            item_data["sku_items"] = sku_items
        if sku_images:
            item_data["sku_images"] = sku_images

        items_to_publish.append(item_data)

    conn.close()

    # 调用批量上架自愈核心
    from xianyu_tools.xianyu_adapter.publisher_v3 import PublisherV3
    try:
        publisher = PublisherV3(account_id=account_id)
        batch_res = publisher.publish_items_batch(items_to_publish)
        
        # 将结果写回数据库记录并整合返回
        conn = get_db_conn(); cursor = conn.cursor()
        final_success = []
        final_failed = failed_items

        for item in items_to_republish:
            sid = item["source_id"]
            pid = item["product_id"]
            try:
                listing_res = publisher.commit_publish(pid)
                if listing_res.get("status") == "success":
                    pub_url = f"https://www.goofish.com/item?id={pid}"
                    cursor.execute(
                        "INSERT INTO xianyu_published_items (task_id, source_db_id, xianyu_item_id, publish_status, publish_msg, published_url) VALUES (%s,%s,%s,%s,%s,%s)",
                        (item["task_id"], sid, pid, 'success', '已重新上架下架商品', pub_url),
                    )
                    final_success.append({
                        "source_id": sid,
                        "product_id": pid,
                        "published_url": pub_url,
                        "status": "success",
                        "mode": "republish"
                    })
                else:
                    final_failed.append({
                        "source_id": sid,
                        "msg": f"重新上架失败: {listing_res.get('msg', '未知错误')}",
                        "status": "failed"
                    })
            except Exception as e:
                logger.error(f"Republish depublished product failed for source {sid}: {e}")
                final_failed.append({
                    "source_id": sid,
                    "msg": f"重新上架异常: {e}",
                    "status": "failed"
                })
        
        for succ in batch_res.get("success", []):
            sid = succ["source_id"]
            pid = succ["product_id"]
            pub_url = f"https://www.goofish.com/item?id={pid}"
            
            cursor.execute("SELECT task_id FROM ali1688_sources WHERE id = %s", (sid,))
            src = cursor.fetchone()
            task_id = src['task_id'] if src else ""
            
            cursor.execute("INSERT INTO xianyu_published_items (task_id, source_db_id, xianyu_item_id, publish_status, publish_msg, published_url) VALUES (%s,%s,%s,%s,%s,%s)",
                          (task_id, sid, pid, 'success', None, pub_url))
            final_success.append({
                "source_id": sid,
                "product_id": pid,
                "published_url": pub_url,
                "status": "success"
            })
            
        for fail in batch_res.get("failed", []):
            sid = fail["source_id"]
            msg = fail["msg"]
            
            cursor.execute("SELECT task_id FROM ali1688_sources WHERE id = %s", (sid,))
            src = cursor.fetchone()
            task_id = src['task_id'] if src else ""
            
            cursor.execute("INSERT INTO xianyu_published_items (task_id, source_db_id, xianyu_item_id, publish_status, publish_msg, published_url) VALUES (%s,%s,%s,%s,%s,%s)",
                          (task_id, sid, None, 'failed', msg, None))
            final_failed.append({
                "source_id": sid,
                "msg": msg,
                "status": "failed"
            })
            
        conn.commit()
        conn.close()
        
        return {"success": final_success, "failed": final_failed}
        
    except Exception as e:
        logger.error(f"Batch publisher global failure: {e}")
        return {"error": str(e), "success": [], "failed": failed_items}

@app.post("/api/publish/{source_id}")
async def publish_to_xianyu(source_id: int, req: dict = {}):
    logger.info(f"OpenAPI publish request for source_id: {source_id}, custom={req}")
    account_id = get_request_openapi_account_id(req)
    conn = get_db_conn(); cursor = conn.cursor()
    cursor.execute("SELECT * FROM ali1688_sources WHERE id = %s", (source_id,))
    source = cursor.fetchone()
    if not source: conn.close(); return {"error": "Source not found"}
    cursor.execute("""
        SELECT publish_status
        FROM xianyu_published_items
        WHERE source_db_id = %s
        ORDER BY created_at DESC
        LIMIT 1
    """, (source_id,))
    latest_publish = cursor.fetchone()
    latest_status = latest_publish.get("publish_status") if latest_publish else None
    if latest_status == "deleted":
        conn.close()
        return {"status": "failed", "msg": "商品已删除，无法重新上架，请重新加入选品后发布"}

    depublished_publish = None
    if latest_status != "success":
        cursor.execute("""
            SELECT xianyu_item_id, task_id
            FROM xianyu_published_items
            WHERE source_db_id = %s AND publish_status = 'depublished' AND xianyu_item_id IS NOT NULL
            ORDER BY created_at DESC
            LIMIT 1
        """, (source_id,))
        depublished_publish = cursor.fetchone()
    from xianyu_tools.xianyu_adapter.publisher_v3 import PublisherV3
    if depublished_publish and depublished_publish.get("xianyu_item_id"):
        try:
            publisher = PublisherV3(account_id=account_id)
            xianyu_item_id = depublished_publish["xianyu_item_id"]
            result = publisher.commit_publish(xianyu_item_id)
            if result.get("status") == "success":
                pub_url = f"https://www.goofish.com/item?id={xianyu_item_id}"
                cursor.execute("INSERT INTO xianyu_published_items (task_id, source_db_id, xianyu_item_id, publish_status, publish_msg, published_url) VALUES (%s,%s,%s,%s,%s,%s)",
                             (depublished_publish.get('task_id') or source['task_id'], source_id, xianyu_item_id, 'success', '已重新上架下架商品', pub_url))
                conn.commit(); conn.close(); return {"status": "success", "xianyu_item_id": xianyu_item_id, "published_url": pub_url, "mode": "republish"}
            conn.close()
            return {"status": "failed", "msg": f"重新上架失败: {result.get('msg', '未知错误')}"}
        except Exception as e:
            logger.error(f"Republish depublished product failed for source {source_id}: {e}")
            conn.close()
            return {"status": "failed", "msg": f"重新上架异常: {e}"}

    images = json.loads(source['images'] or "[]")

    # 支持前端传入自定义标题和价格，否则使用默认值
    custom_title = req.get("title") or source['title']
    
    sku_items = req.get("sku_items")
    if sku_items:
        # 如果有多规格，主商品价格自动校准为多规格中的最低价
        final_price = min(float(item['price']) for item in sku_items)
    else:
        custom_price = req.get("price")
        if custom_price is not None:
            final_price = float(custom_price)
        else:
            final_price = float(source['min_price']) + 30

    item_data = {
        "title": custom_title[:60],
        "description": f"【精选货源】\n{custom_title}\n品质保障，欢迎选购。",
        "price": final_price,
        "images": images,
    }
    sku_images = req.get("sku_images")
    if sku_items:
        item_data["sku_items"] = sku_items
    if sku_images:
        item_data["sku_images"] = sku_images
    try:
        publisher = PublisherV3(account_id=account_id)
        result = publisher.publish_item(item_data)
        pub_url = f"https://www.goofish.com/item?id={result.get('xianyu_item_id')}" if result.get('status') == 'success' else None
        cursor.execute("INSERT INTO xianyu_published_items (task_id, source_db_id, xianyu_item_id, publish_status, publish_msg, published_url) VALUES (%s,%s,%s,%s,%s,%s)",
                     (source['task_id'], source_id, result.get('xianyu_item_id'), result['status'], result.get('msg'), pub_url))
        conn.commit(); conn.close(); return {**result, "published_url": pub_url}
    except Exception as e:
        logger.error(f"Publisher Error: {e}"); conn.close(); return {"status": "failed", "msg": str(e)}

@app.get("/api/published_status/{source_id}")
def get_published_status(source_id: int):
    conn = get_db_conn(); cursor = conn.cursor()
    cursor.execute("SELECT publish_status, published_url FROM xianyu_published_items WHERE source_db_id = %s ORDER BY created_at DESC LIMIT 1", (source_id,))
    res = cursor.fetchone(); conn.close()
    return res if res else {"publish_status": "none"}

@app.post("/api/selection/sync_status/batch")
async def batch_sync_selection_status(req: dict = {}):
    logger.info(f"OpenAPI batch selection status sync request: {req}")
    source_ids = req.get("source_ids", [])
    account_id = get_request_openapi_account_id(req)
    if not source_ids:
        return {"success": [], "failed": [{"source_id": 0, "msg": "未选中任何选品"}]}

    from xianyu_tools.xianyu_adapter.publisher_v3 import PublisherV3
    try:
        publisher = PublisherV3(account_id=account_id)
    except Exception as e:
        logger.error(f"Failed to initialize PublisherV3 for status sync: {e}")
        return {"success": [], "failed": [{"source_id": sid, "msg": f"初始化发布器失败: {e}"} for sid in source_ids]}

    conn = get_db_conn(); cursor = conn.cursor()
    success_list = []
    failed_list = []

    for sid in source_ids:
        cursor.execute("""
            SELECT publish_status, xianyu_item_id, task_id, published_url
            FROM xianyu_published_items
            WHERE source_db_id = %s
            ORDER BY created_at DESC LIMIT 1
        """, (sid,))
        row = cursor.fetchone()

        if not row:
            failed_list.append({"source_id": sid, "msg": "未找到选品记录"})
            continue

        local_status = row.get("publish_status")
        xianyu_item_id = row.get("xianyu_item_id")
        if local_status == "selected":
            failed_list.append({"source_id": sid, "msg": "该选品尚未发布，无需同步闲鱼状态"})
            continue
        if not xianyu_item_id:
            failed_list.append({"source_id": sid, "msg": "缺少闲管家商品 ID，无法查询商品详情"})
            continue

        try:
            result = publisher.query_product_detail(xianyu_item_id)
        except Exception as e:
            logger.error(f"Selection status sync failed for source {sid}: {e}")
            failed_list.append({"source_id": sid, "msg": f"查询商品详情异常: {e}"})
            continue

        if result.get("status") != "success":
            failed_list.append({"source_id": sid, "msg": result.get("msg") or "查询商品详情失败"})
            continue

        detail_data = result.get("data") or {}
        normalized = _normalize_openapi_product_status(detail_data)
        remote_status = normalized.get("status")
        if remote_status not in {"selected", "pending", "success", "failed", "depublished", "deleted"}:
            failed_list.append({
                "source_id": sid,
                "msg": (
                    "查询成功，但暂未识别闲管家状态"
                    f"（product_status={normalized.get('raw_product_status')}, "
                    f"publish_status={normalized.get('raw_publish_status')}）"
                )
            })
            continue

        msg = f"已同步选品状态：{normalized.get('label') or remote_status}"
        published_url = row.get("published_url")
        if remote_status == "success":
            published_url = published_url or f"https://www.goofish.com/item?id={xianyu_item_id}"
        if remote_status in {"depublished", "deleted"}:
            published_url = None

        cursor.execute("""
            INSERT INTO xianyu_published_items (task_id, source_db_id, xianyu_item_id, publish_status, publish_msg, published_url)
            VALUES (%s, %s, %s, %s, %s, %s)
        """, (row.get("task_id"), sid, xianyu_item_id, remote_status, msg, published_url))
        success_list.append({
            "source_id": sid,
            "status": remote_status,
            "publish_status": remote_status,
            "msg": msg,
            "published_url": published_url,
        })

    conn.commit()
    conn.close()
    return {"success": success_list, "failed": failed_list}

@app.post("/api/depublish/batch")
async def batch_depublish_from_xianyu(req: dict = {}):
    logger.info(f"OpenAPI batch depublish request: {req}")
    source_ids = req.get("source_ids", [])
    account_id = get_request_openapi_account_id(req)
    if not source_ids:
        return {"success": [], "failed": [{"source_id": 0, "msg": "未选中任何商品"}]}

    from xianyu_tools.xianyu_adapter.publisher_v3 import PublisherV3
    try:
        publisher = PublisherV3(account_id=account_id)
    except Exception as e:
        logger.error(f"Failed to initialize PublisherV3: {e}")
        return {"success": [], "failed": [{"source_id": sid, "msg": f"初始化发布器失败: {e}"} for sid in source_ids]}

    conn = get_db_conn(); cursor = conn.cursor()
    success_list = []
    failed_list = []

    for sid in source_ids:
        # 1. 查找此货源最近成功的上架记录
        cursor.execute("""
            SELECT xianyu_item_id, task_id 
            FROM xianyu_published_items 
            WHERE source_db_id = %s AND publish_status = 'success' 
            ORDER BY created_at DESC LIMIT 1
        """, (sid,))
        row = cursor.fetchone()
        
        if not row:
            failed_list.append({"source_id": sid, "msg": "未找到该商品的成功发布记录，无法执行下架"})
            continue
            
        xianyu_item_id = row['xianyu_item_id']
        task_id = row['task_id']
        
        # 2. 执行下架
        try:
            result = publisher.depublish_item(xianyu_item_id)
            is_success = result.get("status") == "success"
            is_invalid_state = "不满足下架条件" in result.get("msg", "")
            
            if is_success or is_invalid_state:
                # 3. 记账
                msg = '已下架' if is_success else f"已下架 ({result.get('msg')})"
                cursor.execute("""
                    INSERT INTO xianyu_published_items (task_id, source_db_id, xianyu_item_id, publish_status, publish_msg, published_url)
                    VALUES (%s, %s, %s, %s, %s, %s)
                """, (task_id, sid, xianyu_item_id, 'depublished', msg, None))
                success_list.append({"source_id": sid})
            else:
                failed_list.append({"source_id": sid, "msg": result.get("msg", "下架失败")})
        except Exception as e:
            logger.error(f"Batch depublish failed for source {sid}: {e}")
            failed_list.append({"source_id": sid, "msg": f"下架异常: {e}"})

    conn.commit()
    conn.close()
    return {"success": success_list, "failed": failed_list}

@app.post("/api/depublish/{source_id}")
async def depublish_from_xianyu(source_id: int, req: dict = {}):
    logger.info(f"OpenAPI depublish request for source_id: {source_id}")
    account_id = get_request_openapi_account_id(req)
    conn = get_db_conn(); cursor = conn.cursor()
    
    # 1. 查找此货源最近成功的上架记录
    cursor.execute("""
        SELECT xianyu_item_id, task_id 
        FROM xianyu_published_items 
        WHERE source_db_id = %s AND publish_status = 'success' 
        ORDER BY created_at DESC LIMIT 1
    """, (source_id,))
    row = cursor.fetchone()
    
    if not row:
        conn.close()
        return {"status": "failed", "msg": "未找到该商品的成功发布记录，无法执行下架"}
        
    xianyu_item_id = row['xianyu_item_id']
    task_id = row['task_id']
    
    # 2. 调用 PublisherV3 执行下架
    from xianyu_tools.xianyu_adapter.publisher_v3 import PublisherV3
    try:
        publisher = PublisherV3(account_id=account_id)
        result = publisher.depublish_item(xianyu_item_id)
        
        is_success = result.get("status") == "success"
        is_invalid_state = "不满足下架条件" in result.get("msg", "")
        
        if is_success or is_invalid_state:
            # 3. 在发布表插入已下架状态，完成流水记账
            msg = '已下架' if is_success else f"已下架 ({result.get('msg')})"
            cursor.execute("""
                INSERT INTO xianyu_published_items (task_id, source_db_id, xianyu_item_id, publish_status, publish_msg, published_url)
                VALUES (%s, %s, %s, %s, %s, %s)
            """, (task_id, source_id, xianyu_item_id, 'depublished', msg, None))
            conn.commit()
            conn.close()
            
            return_msg = "下架成功" if is_success else f"已从本地强行下架/同步状态 (闲鱼提示: {result.get('msg')})"
            return {"status": "success", "msg": return_msg}
        else:
            conn.close()
            return {"status": "failed", "msg": result.get("msg", "下架失败")}
            
    except Exception as e:
        logger.error(f"Depublisher Error: {e}")
        conn.close()
        return {"status": "failed", "msg": str(e)}

@app.post("/api/delete/batch")
async def batch_delete_from_xianyu(req: dict = {}):
    logger.info(f"OpenAPI batch delete request: {req}")
    source_ids = req.get("source_ids", [])
    account_id = get_request_openapi_account_id(req)
    if not source_ids:
        return {"success": [], "failed": [{"source_id": 0, "msg": "未选中任何商品"}]}

    conn = get_db_conn(); cursor = conn.cursor()
    success_list = []
    failed_list = []
    publisher = None

    for sid in source_ids:
        # 1. 查找此货源最新的一条发布流水记录
        cursor.execute("""
            SELECT publish_status, xianyu_item_id, task_id 
            FROM xianyu_published_items 
            WHERE source_db_id = %s 
            ORDER BY created_at DESC LIMIT 1
        """, (sid,))
        row = cursor.fetchone()

        if not row:
            failed_list.append({"source_id": sid, "msg": "商品未发布，无法删除"})
            continue

        status = row['publish_status']
        xianyu_item_id = row['xianyu_item_id']
        task_id = row['task_id']

        if status in ('failed', 'selected'):
            cursor.execute("""
                INSERT INTO xianyu_published_items (task_id, source_db_id, xianyu_item_id, publish_status, publish_msg, published_url)
                VALUES (%s, %s, %s, %s, %s, %s)
            """, (task_id, sid, xianyu_item_id, 'deleted', '已清理本地选品记录' if status == 'selected' else '已清理失败发布记录', None))
            success_list.append({"source_id": sid, "mode": "local_cleanup"})
            continue

        if status != 'depublished':
            failed_list.append({"source_id": sid, "msg": f"商品状态为 {status}，只有已下架或同步失败商品可以删除"})
            continue

        # 2. 已下架商品调用 PublisherV3 执行云端删除
        try:
            if publisher is None:
                from xianyu_tools.xianyu_adapter.publisher_v3 import PublisherV3
                publisher = PublisherV3(account_id=account_id)
            result = publisher.delete_item(xianyu_item_id)
            if result.get("status") == "success":
                # 3. 记账
                cursor.execute("""
                    INSERT INTO xianyu_published_items (task_id, source_db_id, xianyu_item_id, publish_status, publish_msg, published_url)
                    VALUES (%s, %s, %s, %s, %s, %s)
                """, (task_id, sid, xianyu_item_id, 'deleted', '已删除', None))
                success_list.append({"source_id": sid, "mode": "cloud_delete"})
            else:
                failed_list.append({"source_id": sid, "msg": result.get("msg", "删除失败")})
        except Exception as e:
            logger.error(f"Batch delete failed for source {sid}: {e}")
            failed_list.append({"source_id": sid, "msg": f"删除异常: {e}"})

    conn.commit()
    conn.close()
    return {"success": success_list, "failed": failed_list}

@app.post("/api/delete/{source_id}")
async def delete_from_xianyu(source_id: int, req: dict = {}):
    logger.info(f"OpenAPI delete request for source_id: {source_id}")
    account_id = get_request_openapi_account_id(req)
    conn = get_db_conn(); cursor = conn.cursor()

    # 1. 查找此货源最新的一条发布流水记录
    cursor.execute("""
        SELECT publish_status, xianyu_item_id, task_id 
        FROM xianyu_published_items 
        WHERE source_db_id = %s 
        ORDER BY created_at DESC LIMIT 1
    """, (source_id,))
    row = cursor.fetchone()

    if not row:
        conn.close()
        return {"status": "failed", "msg": "商品未发布，无法删除"}

    status = row['publish_status']
    xianyu_item_id = row['xianyu_item_id']
    task_id = row['task_id']

    if status in ('failed', 'selected'):
        cursor.execute("""
            INSERT INTO xianyu_published_items (task_id, source_db_id, xianyu_item_id, publish_status, publish_msg, published_url)
            VALUES (%s, %s, %s, %s, %s, %s)
        """, (task_id, source_id, xianyu_item_id, 'deleted', '已清理本地选品记录' if status == 'selected' else '已清理失败发布记录', None))
        conn.commit()
        conn.close()
        return {"status": "success", "msg": "已清理本地选品记录" if status == 'selected' else "已清理失败发布记录", "mode": "local_cleanup"}

    if status != 'depublished':
        conn.close()
        return {"status": "failed", "msg": f"商品当前状态为 {status}，只有已下架或同步失败商品可以删除"}

    # 2. 已下架商品调用 PublisherV3 执行云端删除
    from xianyu_tools.xianyu_adapter.publisher_v3 import PublisherV3
    try:
        publisher = PublisherV3(account_id=account_id)
        result = publisher.delete_item(xianyu_item_id)

        if result.get("status") == "success":
            # 3. 记账
            cursor.execute("""
                INSERT INTO xianyu_published_items (task_id, source_db_id, xianyu_item_id, publish_status, publish_msg, published_url)
                VALUES (%s, %s, %s, %s, %s, %s)
            """, (task_id, source_id, xianyu_item_id, 'deleted', '已删除', None))
            conn.commit()
            conn.close()
            return {"status": "success", "msg": "删除成功", "mode": "cloud_delete"}
        else:
            conn.close()
            return {"status": "failed", "msg": result.get("msg", "删除失败")}

    except Exception as e:
        logger.error(f"Publisher Delete Error: {e}")
        conn.close()
        return {"status": "failed", "msg": str(e)}

@app.get("/api/xianyu_products")
def get_xianyu_products(
    page: int = 1, 
    limit: int = 10, 
    keyword: str = "", 
    sort_by: str = "publish_time", 
    sort_order: str = "desc",
    publish_status: str = "",
    min_source_price: float = None,
    max_source_price: float = None,
    min_ref_price: float = None,
    max_ref_price: float = None
):
    offset = (page - 1) * limit
    conn = get_db_conn(); cursor = conn.cursor()

    # 联表查询最新一条发布记录，且最新状态非 'deleted'
    query_base = """
        FROM xianyu_published_items p
        INNER JOIN (
            SELECT source_db_id, MAX(created_at) as max_time
            FROM xianyu_published_items
            GROUP BY source_db_id
        ) latest ON p.source_db_id = latest.source_db_id AND p.created_at = latest.max_time
        INNER JOIN ali1688_sources s ON p.source_db_id = s.id
        LEFT JOIN xianyu_items xi ON s.item_id = xi.id
        WHERE p.publish_status IN ('selected', 'success', 'depublished', 'pending', 'failed')
    """

    params = []
    if keyword:
        query_base += " AND (s.title LIKE %s OR xi.title LIKE %s)"
        params.extend([f"%{keyword}%", f"%{keyword}%"])

    if publish_status:
        if publish_status == 'success':
            query_base += " AND p.publish_status IN ('success', 'done')"
        else:
            query_base += " AND p.publish_status = %s"
            params.append(publish_status)

    if min_source_price is not None:
        query_base += " AND s.min_price >= %s"
        params.append(min_source_price)
    if max_source_price is not None:
        query_base += " AND s.min_price <= %s"
        params.append(max_source_price)

    if min_ref_price is not None:
        query_base += " AND xi.price >= %s"
        params.append(min_ref_price)
    if max_ref_price is not None:
        query_base += " AND xi.price <= %s"
        params.append(max_ref_price)

    count_query = f"SELECT COUNT(*) as count {query_base}"
    cursor.execute(count_query, tuple(params))
    total_count = cursor.fetchone()['count']

    # 排序字段映射防御 SQL 注入
    sort_mapping = {
        "title": "s.title",
        "xianyu_item_id": "p.xianyu_item_id",
        "publish_status": "p.publish_status",
        "source_price": "s.min_price",
        "ref_price": "xi.price",
        "publish_time": "p.created_at"
    }
    order_field = sort_mapping.get(sort_by, "p.created_at")
    order_direction = "DESC" if sort_order.lower() == "desc" else "ASC"

    data_query = f"""
        SELECT 
            p.id as publish_id,
            p.task_id,
            p.source_db_id,
            p.xianyu_item_id,
            p.publish_status,
            p.publish_msg,
            p.published_url,
            p.created_at as publish_time,
            s.title as source_title,
            s.source_url as source_url,
            s.images as source_images,
            s.min_price as source_price,
            s.sku_count as source_sku_count,
            s.is_detail_incomplete as source_is_detail_incomplete,
            s.detail_incomplete_reason as source_detail_incomplete_reason,
            s.detail_status as source_detail_status,
            xi.title as ref_title,
            xi.price as ref_price,
            xi.want_count as ref_want_count
        {query_base}
        ORDER BY {order_field} {order_direction}
        LIMIT %s OFFSET %s
    """
    params.extend([limit, offset])
    cursor.execute(data_query, tuple(params))
    rows = cursor.fetchall()
    conn.close()

    items = []
    for r in rows:
        images_list = []
        if r['source_images']:
            try:
                images_list = json.loads(r['source_images'])
            except Exception:
                pass

        items.append({
            "publish_id": r["publish_id"],
            "task_id": r["task_id"],
            "source_db_id": r["source_db_id"],
            "xianyu_item_id": r["xianyu_item_id"],
            "publish_status": r["publish_status"],
            "publish_msg": r["publish_msg"],
            "published_url": r["published_url"],
            "publish_time": r["publish_time"].strftime("%Y-%m-%d %H:%M:%S") if r["publish_time"] else "",
            "source_title": r["source_title"],
            "source_url": r["source_url"],
            "source_image": images_list[0] if images_list else "",
            "source_price": float(r["source_price"]) if r["source_price"] is not None else 0.0,
            "source_sku_count": r["source_sku_count"],
            "source_is_detail_incomplete": bool(r.get("source_is_detail_incomplete")),
            "source_detail_incomplete_reason": r.get("source_detail_incomplete_reason") or "",
            "source_detail_status": r.get("source_detail_status") or "",
            "ref_title": r["ref_title"] or "",
            "ref_price": float(r["ref_price"]) if r["ref_price"] is not None else 0.0,
            "ref_want_count": r["ref_want_count"] or 0
        })

    return {
        "items": items,
        "total": total_count,
        "page": page,
        "limit": limit
    }

@app.get("/api/orders")
def get_openapi_orders(
    page: int = 1,
    limit: int = 20,
    order_status: str = "",
    account_id: str = "",
):
    from xianyu_tools.xianyu_adapter.publisher_v3 import PublisherV3
    try:
        safe_page = max(1, int(page or 1))
        safe_limit = min(max(1, int(limit or 20)), 100)
        publisher = PublisherV3(account_id=get_request_openapi_account_id({"account_id": account_id}))
        result = publisher.query_order_list(
            page_no=safe_page,
            page_size=safe_limit,
            order_status=order_status,
        )
        if result.get("status") != "success":
            return {
                "status": "failed",
                "msg": result.get("msg") or "订单列表查询失败",
                "items": [],
                "total": 0,
                "page": safe_page,
                "limit": safe_limit,
            }
        items, total = _extract_openapi_order_list(result.get("data") or {})
        return {
            "status": "success",
            "items": items,
            "total": total,
            "page": safe_page,
            "limit": safe_limit,
        }
    except Exception as e:
        logger.error(f"OpenAPI order list failed: {e}")
        return {"status": "failed", "msg": f"订单列表查询异常: {e}", "items": [], "total": 0, "page": page, "limit": limit}


@app.get("/api/orders/{order_no}")
def get_openapi_order_detail(order_no: str, account_id: str = ""):
    from xianyu_tools.xianyu_adapter.publisher_v3 import PublisherV3
    try:
        publisher = PublisherV3(account_id=get_request_openapi_account_id({"account_id": account_id}))
        result = publisher.query_order_detail(order_no)
        if result.get("status") != "success":
            return {"status": "failed", "msg": result.get("msg") or "订单详情查询失败"}
        return {
            "status": "success",
            "order": _normalize_openapi_order(result.get("data") or {}, include_raw=True),
        }
    except Exception as e:
        logger.error(f"OpenAPI order detail failed for {order_no}: {e}")
        return {"status": "failed", "msg": f"订单详情查询异常: {e}"}


@app.post("/api/orders/{order_no}/ship")
def ship_openapi_order(order_no: str, req: dict = {}):
    from xianyu_tools.xianyu_adapter.publisher_v3 import PublisherV3
    try:
        payload = dict(req or {})
        payload["order_no"] = str(order_no or "").strip()
        publisher = PublisherV3(account_id=get_request_openapi_account_id(payload))
        detail_result = publisher.query_order_detail(order_no)
        if detail_result.get("status") != "success":
            return {"status": "failed", "msg": detail_result.get("msg") or "订单详情查询失败，无法确认是否可发货"}
        order_status = (detail_result.get("data") or {}).get("order_status")
        if int(order_status or 0) != 12:
            return {"status": "failed", "msg": f"当前订单状态为 {_format_openapi_order_status(order_status)}，只有待发货订单才能物流发货"}
        result = publisher.ship_order(payload)
        if result.get("status") != "success":
            return {"status": "failed", "msg": result.get("msg") or "订单物流发货失败"}
        return {"status": "success", "msg": "订单物流发货成功", "raw": result.get("raw")}
    except Exception as e:
        logger.error(f"OpenAPI order ship failed for {order_no}: {e}")
        return {"status": "failed", "msg": f"订单物流发货异常: {e}"}


@app.post("/api/orders/{order_no}/modify_price")
def modify_openapi_order_price(order_no: str, req: dict = {}):
    from xianyu_tools.xianyu_adapter.publisher_v3 import PublisherV3
    try:
        payload = dict(req or {})
        order_price = _parse_cent_amount_from_request(payload, "order_price", "order_price_yuan")
        express_fee = _parse_cent_amount_from_request(payload, "express_fee", "express_fee_yuan", default=0)
        if order_price is None:
            return {"status": "failed", "msg": "请填写有效的订单价格"}
        if express_fee is None:
            return {"status": "failed", "msg": "请填写有效的运费"}
        publisher = PublisherV3(account_id=get_request_openapi_account_id(payload))
        detail_result = publisher.query_order_detail(order_no)
        if detail_result.get("status") != "success":
            return {"status": "failed", "msg": detail_result.get("msg") or "订单详情查询失败，无法确认是否可改价"}
        order_status = (detail_result.get("data") or {}).get("order_status")
        if int(order_status or 0) != 11:
            return {"status": "failed", "msg": f"当前订单状态为 {_format_openapi_order_status(order_status)}，只有待付款订单才能修改价格"}
        result = publisher.modify_order_price(order_no, order_price, express_fee)
        if result.get("status") != "success":
            return {"status": "failed", "msg": result.get("msg") or "订单修改价格失败"}
        return {"status": "success", "msg": "订单修改价格成功", "raw": result.get("raw")}
    except Exception as e:
        logger.error(f"OpenAPI order modify price failed for {order_no}: {e}")
        return {"status": "failed", "msg": f"订单修改价格异常: {e}"}


@app.get("/api/tasks/{task_id}/logs", response_class=PlainTextResponse)
def get_logs(task_id: str):
    try:
        conn = get_db_conn(); cursor = conn.cursor()
        cursor.execute("SELECT root_dir FROM tasks WHERE id = %s", (task_id,))
        res = cursor.fetchone(); conn.close()
        if not res: return "Task not found"
        log_path = Path(res['root_dir']) / "task.log"
        return log_path.read_text(encoding='utf-8', errors='ignore') if log_path.exists() else "No logs yet"
    except Exception as e:
        logger.error(f"Failed to fetch logs for task {task_id}: {e}")
        return f"Error reading logs: {e}"

@app.get("/api/token/stats")
def get_token_stats():
    try:
        conn = get_db_conn()
        cursor = conn.cursor()
        
        # 1. 汇总数据
        cursor.execute("""
            SELECT 
                COUNT(*) as total_calls,
                IFNULL(SUM(prompt_tokens), 0) as total_prompt_tokens,
                IFNULL(SUM(completion_tokens), 0) as total_completion_tokens,
                IFNULL(SUM(total_tokens), 0) as total_tokens,
                COUNT(DISTINCT model) as model_count,
                COUNT(DISTINCT feature) as feature_count
            FROM llm_token_logs
        """)
        summary = cursor.fetchone()
        if not summary or summary.get("total_calls") == 0:
            summary = {
                "total_calls": 0, "total_prompt_tokens": 0, "total_completion_tokens": 0, 
                "total_tokens": 0, "model_count": 0, "feature_count": 0
            }
        
        # 2. 按模型统计
        cursor.execute("""
            SELECT 
                model,
                COUNT(*) as calls,
                IFNULL(SUM(prompt_tokens), 0) as prompt_tokens,
                IFNULL(SUM(completion_tokens), 0) as completion_tokens,
                IFNULL(SUM(total_tokens), 0) as total_tokens
            FROM llm_token_logs
            GROUP BY model
            ORDER BY total_tokens DESC
        """)
        by_model = cursor.fetchall()
        
        # 3. 按功能统计
        cursor.execute("""
            SELECT 
                feature,
                COUNT(*) as calls,
                IFNULL(SUM(prompt_tokens), 0) as prompt_tokens,
                IFNULL(SUM(completion_tokens), 0) as completion_tokens,
                IFNULL(SUM(total_tokens), 0) as total_tokens
            FROM llm_token_logs
            GROUP BY feature
            ORDER BY total_tokens DESC
        """)
        by_feature = cursor.fetchall()
        
        # 4. 明细日志 (关联任务关键词)
        cursor.execute("""
            SELECT 
                l.id,
                l.task_id,
                t.keyword as task_keyword,
                l.feature,
                l.model,
                l.prompt_tokens,
                l.completion_tokens,
                l.total_tokens,
                DATE_FORMAT(l.created_at, '%Y-%m-%d %H:%i:%s') as created_at
            FROM llm_token_logs l
            LEFT JOIN tasks t ON l.task_id = t.id
            ORDER BY l.id DESC
        """)
        recent_logs = cursor.fetchall()
        
        conn.close()
        return {
            "status": "success",
            "summary": summary,
            "by_model": by_model,
            "by_feature": by_feature,
            "recent_logs": recent_logs
        }
    except Exception as e:
        logger.error(f"Failed to fetch token stats: {e}")
        return {"status": "error", "message": str(e)}

@app.get("/api/sys/status")
def system_status(): 
    try:
        conn = get_db_conn(); cursor = conn.cursor()
        cursor.execute("SELECT COUNT(*) as count FROM tasks WHERE status = '执行中' AND is_deleted = 0")
        active_count = cursor.fetchone()['count']; conn.close()
        runtime_cfg = settings.get_active_ali1688_runtime_config()
        if runtime_cfg.get("account_id"):
            session_report = inspect_ali1688_state_file_quick(
                runtime_cfg.get("state_file"),
                runtime_cfg.get("label") or runtime_cfg.get("account_id") or "",
            )
            login_status = "有效" if session_report.get("is_usable") else "无效"
        else:
            login_status = "未配置"
        return {"1688_login": login_status, "active_workers": active_count, "db_type": "MySQL"}
    except Exception as e:
        logger.error(f"Failed to get system status: {e}")
        return {"error": str(e)}

@app.get("/api/system/configs")
def get_system_configs():
    try:
        from xianyu_tools.config import settings
        from xianyu_tools.xianyu_adapter.state_inspector import inspect_state_file

        openapi_cfg = normalize_openapi_multi_account(settings.get_openapi_raw_config())
        source_channels_cfg = normalize_source_channels_config(settings.get_source_channels_raw_config())
        for account in openapi_cfg["accounts"]:
            try:
                session_report = inspect_state_file(account.get("state_file") or "xianyu_state.json")
                account_name = session_report.get("account_name") or ""
                is_usable = bool(session_report.get("is_usable"))
            except Exception:
                account_name = ""
                is_usable = False
            account["default_config"]["user_name"] = account_name
            account["session_report"] = {
                "account_name": account_name,
                "is_usable": is_usable,
            }

        for channel in source_channels_cfg["channels"]:
            for account in channel.get("accounts", []):
                if channel.get("channel_type") == "ali1688":
                    account["session_report"] = merge_ali1688_session_report(
                        account.get("state_file"),
                        account.get("label") or account.get("account_id") or "",
                    )
                else:
                    account["session_report"] = {
                        "is_usable": False,
                        "account_name": "",
                        "status_text": "暂不支持该渠道检测",
                        "last_checked_at": "",
                        "error_message": "",
                        "meta": {"channel_type": channel.get("channel_type")},
                    }

        return {
            "status": "success",
            "data": {
                "llm": settings.get_llm_config(),
                "openapi": openapi_cfg,
                "crawl": settings.get_crawl_config(),
                "source_channels": source_channels_cfg
            }
        }
    except Exception as e:
        logger.error(f"Failed to fetch system configs API: {e}")
        return {"status": "failed", "msg": str(e)}

@app.post("/api/system/configs")
async def update_system_configs(payload: dict):
    try:
        from xianyu_tools.config import settings
        from xianyu_tools.xianyu_adapter.state_inspector import inspect_state_file
        import json

        def summarize_crawl_selection_adjustments(raw_crawl: dict, normalized_crawl: dict) -> str:
            raw_mode = raw_crawl.get("source_channel_selection_mode") if isinstance(raw_crawl, dict) else None
            normalized_mode = normalized_crawl.get("source_channel_selection_mode") if isinstance(normalized_crawl, dict) else None
            raw_entries = raw_crawl.get("enabled_source_channels") if isinstance(raw_crawl, dict) else []
            normalized_entries = normalized_crawl.get("enabled_source_channels") if isinstance(normalized_crawl, dict) else []
            raw_entries = raw_entries if isinstance(raw_entries, list) else []
            normalized_entries = normalized_entries if isinstance(normalized_entries, list) else []

            raw_map = {}
            for item in raw_entries:
                if not isinstance(item, dict):
                    continue
                channel_id = (item.get("channel_id") or "").strip()
                if not channel_id:
                    continue
                account_ids = item.get("account_ids") if isinstance(item.get("account_ids"), list) else []
                raw_map[channel_id] = {
                    "enabled": item.get("enabled") is not False,
                    "account_ids": [str(account_id).strip() for account_id in account_ids if str(account_id).strip()],
                }

            normalized_map = {}
            for item in normalized_entries:
                if not isinstance(item, dict):
                    continue
                channel_id = (item.get("channel_id") or "").strip()
                if not channel_id:
                    continue
                account_ids = item.get("account_ids") if isinstance(item.get("account_ids"), list) else []
                normalized_map[channel_id] = {
                    "enabled": item.get("enabled") is not False,
                    "account_ids": [str(account_id).strip() for account_id in account_ids if str(account_id).strip()],
                }

            removed_channels = sorted(channel_id for channel_id in raw_map.keys() if channel_id not in normalized_map)
            removed_accounts = []
            for channel_id, raw_item in raw_map.items():
                raw_accounts = raw_item.get("account_ids") or []
                normalized_accounts = set((normalized_map.get(channel_id) or {}).get("account_ids") or [])
                diff_accounts = [account_id for account_id in raw_accounts if account_id not in normalized_accounts]
                if diff_accounts:
                    removed_accounts.append(f"{channel_id}: {', '.join(diff_accounts)}")

            if raw_mode == "custom_selected" and normalized_mode == "custom_selected" and (removed_channels or removed_accounts):
                summary_parts = []
                if removed_channels:
                    summary_parts.append(f"无效渠道已自动移除：{'、'.join(removed_channels)}")
                if removed_accounts:
                    summary_parts.append(f"不可用账号已自动收敛：{'；'.join(removed_accounts)}")
                return "；".join(summary_parts)

            return ""
        
        llm_cfg = payload.get("llm", [])
        openapi_cfg = normalize_openapi_multi_account(payload.get("openapi", {}))
        raw_crawl_cfg = payload.get("crawl", {}) if isinstance(payload.get("crawl", {}), dict) else {}
        crawl_cfg = raw_crawl_cfg
        source_channels_cfg = normalize_source_channels_config(payload.get("source_channels", {}))
        source_channels_storage_cfg = strip_source_channel_runtime_fields(source_channels_cfg)
        crawl_cfg = settings.normalize_crawl_config(crawl_cfg, source_channels_cfg=source_channels_cfg)
        
        logger.info(
            "[OpenAPI Save] active_account_id=%s accounts=%s",
            openapi_cfg.get("active_account_id"),
            [item.get("id") for item in openapi_cfg.get("accounts", [])]
        )
        for account in openapi_cfg["accounts"]:
            account.pop("session_report", None)
            try:
                session_report = inspect_state_file(account.get("state_file") or "xianyu_state.json")
                account_name = session_report.get("account_name") or ""
            except Exception:
                account_name = ""
            account["default_config"]["user_name"] = account_name
            
        # --- 校验并规范化 crawl 配置 ---
        source_limit = crawl_cfg.get("source_limit_1688", 10)
        try:
            source_limit = int(source_limit)
            if source_limit <= 0:
                source_limit = 10
            elif source_limit > 100:
                source_limit = 100
        except (ValueError, TypeError):
            source_limit = 10
        crawl_cfg["source_limit_1688"] = source_limit

        gross_profit_rate = crawl_cfg.get("gross_profit_rate", DEFAULT_GROSS_PROFIT_RATE)
        crawl_cfg["gross_profit_rate"] = _normalize_gross_profit_rate(gross_profit_rate)

        # 过滤模型子集，确保其中每一个都存在于 llm 配置的白名单中
        models_subset = crawl_cfg.get("source_filter_models", [])
        if not isinstance(models_subset, list):
            models_subset = []
        valid_models = []
        for cfg in llm_cfg:
            raw_models = cfg.get("models") or cfg.get("model") or []
            model_arr = raw_models if isinstance(raw_models, list) else [raw_models]
            for m in model_arr:
                if m and isinstance(m, str) and m not in valid_models:
                    valid_models.append(m)
        filtered_models = [m for m in models_subset if m in valid_models]
        crawl_cfg["source_filter_models"] = filtered_models
        crawl_cfg = settings.normalize_crawl_config(crawl_cfg, source_channels_cfg=source_channels_cfg)

        for channel in source_channels_storage_cfg["channels"]:
            for account in channel.get("accounts", []):
                account.pop("session_report", None)
        
        # 1. 动态加载本地已有的 config.json 文件，以完整保留原有 database 配置！
        config_file = settings.config_file
        existing_config = {}
        if config_file.exists():
            try:
                with open(config_file, "r", encoding="utf-8") as f:
                    existing_config = json.load(f)
            except Exception:
                pass
                
        db_cfg = existing_config.get("database", {})
        # 如果读取失败或者原本没有，则采用 settings 的默认后备数据库配置
        if not db_cfg:
            db_cfg = settings.get_database_config()
            
        new_config_data = {
            "database": db_cfg,
            "llm": llm_cfg,
            "openapi": openapi_cfg,
            "crawl": crawl_cfg,
            "source_channels": source_channels_storage_cfg
        }
        
        config_file.parent.mkdir(parents=True, exist_ok=True)
        with open(config_file, "w", encoding="utf-8") as f:
            json.dump(new_config_data, f, ensure_ascii=False, indent=4)
            
        # 2. 同步写入数据库表中
        conn = get_db_conn()
        cursor = conn.cursor()
        
        # 确保表一定存在
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS system_configs (
                cfg_key VARCHAR(50) PRIMARY KEY,
                cfg_value TEXT NOT NULL
            ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4;
        """)
        
        cursor.execute("INSERT INTO system_configs (cfg_key, cfg_value) VALUES (%s, %s) ON DUPLICATE KEY UPDATE cfg_value = VALUES(cfg_value)", ("llm", json.dumps(llm_cfg)))
        cursor.execute("INSERT INTO system_configs (cfg_key, cfg_value) VALUES (%s, %s) ON DUPLICATE KEY UPDATE cfg_value = VALUES(cfg_value)", ("openapi", json.dumps(openapi_cfg)))
        cursor.execute("INSERT INTO system_configs (cfg_key, cfg_value) VALUES (%s, %s) ON DUPLICATE KEY UPDATE cfg_value = VALUES(cfg_value)", ("crawl", json.dumps(crawl_cfg)))
        cursor.execute("INSERT INTO system_configs (cfg_key, cfg_value) VALUES (%s, %s) ON DUPLICATE KEY UPDATE cfg_value = VALUES(cfg_value)", ("source_channels", json.dumps(source_channels_storage_cfg)))
        conn.commit()
        conn.close()
        
        # 3. 清理 settings 的数据库缓存，并重新触发 load()
        settings._db_cache.clear()
        settings.load()

        adjustment_summary = summarize_crawl_selection_adjustments(raw_crawl_cfg, crawl_cfg)
        success_msg = "配置已保存并同步成功"
        if adjustment_summary:
            success_msg = f"{success_msg}。{adjustment_summary}"

        return {"status": "success", "msg": success_msg}
    except Exception as e:
        logger.error(f"Failed to save system configs API: {e}")
        return {"status": "failed", "msg": f"保存配置发生异常: {str(e)}"}


@app.post("/api/system/source_channel_status/check")
async def check_source_channel_status(payload: dict = None):
    try:
        sync_source_channel_login_state()
        payload = payload or {}
        incoming_channel = payload.get("channel")
        incoming_account = payload.get("account")

        if isinstance(incoming_channel, dict) and isinstance(incoming_account, dict):
            channel = {
                "channel_id": incoming_channel.get("channel_id") or "ali1688",
                "channel_type": incoming_channel.get("channel_type") or "ali1688",
                "label": incoming_channel.get("label") or "货源渠道",
            }
            account = dict(incoming_account)
        else:
            raw_cfg = settings.get_source_channels_raw_config()
            normalized = normalize_source_channels_config(raw_cfg)
            channel_id = payload.get("channel_id")
            account_id = payload.get("account_id")
            channel = get_source_channel(normalized, channel_id)
            account = get_source_channel_account(normalized, channel.get("channel_id"), account_id)
        account_error = get_source_channel_account_error(channel, account)
        if account_error:
            return {
                "status": "failed",
                "msg": account_error,
                "data": {
                    "channel_id": channel.get("channel_id"),
                    "account_id": account.get("account_id"),
                    "report": build_source_channel_unavailable_report(channel, account, account_error),
                },
            }
        if not source_channel_supports_session_state(channel):
            return {
                "status": "success",
                "data": {
                    "channel_id": channel.get("channel_id"),
                    "account_id": account.get("account_id"),
                    "report": build_source_channel_unavailable_report(
                        channel,
                        account,
                        "当前渠道暂未接入会话状态检测",
                        "待接入",
                    ),
                },
            }
        channel, account = hydrate_source_channel_account(channel, account)
        is_currently_logging_in = (
            source_channel_login_process is not None
            and source_channel_logging_in_channel_id == (channel.get("channel_id") or "")
            and source_channel_logging_in_account_id == (account.get("account_id") or "")
        )
        if is_currently_logging_in:
            return {
                "status": "failed",
                "msg": "当前账号正在进行 1688 登录，请先完成或关闭登录窗口后再检测状态",
                "data": {
                    "channel_id": channel.get("channel_id"),
                    "account_id": account.get("account_id"),
                    "report": merge_ali1688_session_report(
                        account.get("state_file"),
                        account.get("label") or account.get("account_id") or "",
                    ),
                },
            }
        realtime_report = await inspect_source_channel_account(channel.get("channel_type"), account)
        report = merge_ali1688_session_report(
            account.get("state_file"),
            account.get("label") or account.get("account_id") or "",
            realtime_report=realtime_report,
        )
        if (realtime_report.get("meta") or {}).get("state") != "profile_locked":
            save_ali1688_session_report_cache(account.get("state_file"), report)
        return {
            "status": "success",
            "data": {
                "channel_id": channel.get("channel_id"),
                "account_id": account.get("account_id"),
                "report": report
            }
        }
    except Exception as e:
        logger.error(f"Failed to check source channel status: {e}")
        return {"status": "failed", "msg": str(e)}


@app.get("/api/system/source_channels/active")
def get_active_source_channel_runtime():
    try:
        runtime_cfg = settings.get_active_ali1688_runtime_config()
        if runtime_cfg.get("account_id"):
            report = merge_ali1688_session_report(
                runtime_cfg.get("state_file"),
                runtime_cfg.get("label") or runtime_cfg.get("account_id") or "",
            )
        else:
            report = build_source_channel_unavailable_report(
                {"channel_id": runtime_cfg.get("channel_id"), "label": "1688 货源渠道"},
                {},
                runtime_cfg.get("error_message"),
            )
        return {
            "status": "success",
            "data": {
                **runtime_cfg,
                "session_report": report,
            },
        }
    except Exception as e:
        logger.error(f"Failed to fetch active source channel runtime: {e}")
        return {"status": "failed", "msg": str(e)}


@app.get("/api/system/source_channel_login_status")
def get_source_channel_login_status(channel_id: str = None, account_id: str = None):
    try:
        sync_source_channel_login_state()

        raw_cfg = settings.get_source_channels_raw_config()
        normalized = normalize_source_channels_config(raw_cfg)
        channel = get_source_channel(normalized, channel_id)
        account = get_source_channel_account(normalized, channel.get("channel_id"), account_id)
        account_error = get_source_channel_account_error(channel, account)
        if account_error:
            return {
                "status": "success",
                "data": {
                    "channel_id": channel.get("channel_id"),
                    "account_id": account.get("account_id"),
                    "report": build_source_channel_unavailable_report(channel, account, account_error),
                    "is_logging_in": False,
                    "err_msg": "",
                }
            }
        if not source_channel_supports_session_state(channel):
            return {
                "status": "success",
                "data": {
                    "channel_id": channel.get("channel_id"),
                    "account_id": account.get("account_id"),
                    "report": build_source_channel_unavailable_report(
                        channel,
                        account,
                        "当前渠道暂未接入会话状态检测",
                        "待接入",
                    ),
                    "is_logging_in": False,
                    "err_msg": "",
                }
            }
        channel, account = hydrate_source_channel_account(channel, account)
        report = merge_ali1688_session_report(
            account.get("state_file"),
            account.get("label") or account.get("account_id") or "",
        )

        is_currently_logging_in = (
            source_channel_login_process is not None
            and source_channel_logging_in_channel_id == (channel.get("channel_id") or "")
            and source_channel_logging_in_account_id == (account.get("account_id") or "")
        )
        if is_currently_logging_in:
            report = {
                **report,
                "error_message": "",
            }

        return {
            "status": "success",
            "data": {
                "channel_id": channel.get("channel_id"),
                "account_id": account.get("account_id"),
                "report": report,
                "is_logging_in": is_currently_logging_in,
                "login_started_at": source_channel_login_started_at,
                "err_msg": source_channel_login_err_msg,
                "login_result": source_channel_login_result,
                "login_result_finished_at": source_channel_login_result_finished_at,
            }
        }
    except Exception as e:
        logger.error(f"Failed to check source channel login status: {e}")
        return {"status": "failed", "msg": str(e)}


@app.post("/api/system/source_channel_login_trigger")
async def trigger_source_channel_login(payload: dict = None):
    global source_channel_login_process, source_channel_logging_in_channel_id, source_channel_logging_in_account_id
    global source_channel_login_err_msg, source_channel_login_log_path
    global source_channel_login_result, source_channel_login_result_finished_at, source_channel_login_started_at

    sync_source_channel_login_state()
    if source_channel_login_process is not None:
        return {"status": "failed", "msg": "当前已有渠道账号登录任务在执行，请勿重复操作"}

    try:
        payload = payload or {}
        incoming_channel = payload.get("channel")
        incoming_account = payload.get("account")

        raw_cfg = settings.get_source_channels_raw_config()
        normalized = normalize_source_channels_config(raw_cfg)

        requested_channel_id = (
            payload.get("channel_id")
            or (incoming_channel.get("channel_id") if isinstance(incoming_channel, dict) else None)
        )
        requested_account_id = (
            payload.get("account_id")
            or (incoming_account.get("account_id") if isinstance(incoming_account, dict) else None)
        )

        channel = get_source_channel(normalized, requested_channel_id)
        account = get_source_channel_account(normalized, channel.get("channel_id"), requested_account_id)

        if isinstance(incoming_channel, dict):
            channel = {
                **channel,
                "channel_id": channel.get("channel_id") or incoming_channel.get("channel_id") or "ali1688",
                "channel_type": channel.get("channel_type") or incoming_channel.get("channel_type") or "ali1688",
                "label": channel.get("label") or incoming_channel.get("label") or "货源渠道",
            }
        if isinstance(incoming_account, dict):
            account = {
                **account,
                "account_id": account.get("account_id") or incoming_account.get("account_id"),
                "label": account.get("label") or incoming_account.get("label"),
                "notes": incoming_account.get("notes", account.get("notes") or ""),
            }
        account_error = get_source_channel_account_error(channel, account)
        if account_error:
            return {"status": "failed", "msg": account_error}
        channel, account = hydrate_source_channel_account(channel, account)

        if channel.get("channel_type") != "ali1688":
            return {"status": "failed", "msg": "当前仅支持为 1688 渠道触发登录"}

        state_file = account.get("state_file") or "state/source_channels/ali1688/ali1688-account-1/storage_state.json"
        user_data_dir = account.get("user_data_dir") or ""
        profile_directory = account.get("profile_directory") or ""
        log_dir = BASE_DIR / "tmp" / "source-channel-login-logs"
        log_dir.mkdir(parents=True, exist_ok=True)
        safe_account_id = re.sub(r"[^a-zA-Z0-9_-]", "_", str(account.get("account_id") or "ali1688-account"))
        source_channel_login_log_path = str((log_dir / f"{safe_account_id}.log").resolve())

        cmd = [
            sys.executable,
            "-u",
            str((BASE_DIR / "scripts" / "refresh_1688_state.py").resolve()),
            "--state-file",
            state_file,
        ]
        if user_data_dir:
            cmd.extend(["--user-data-dir", user_data_dir])
        if profile_directory:
            cmd.extend(["--profile-directory", profile_directory])

        log_fd = os.open(source_channel_login_log_path, os.O_CREAT | os.O_WRONLY | os.O_TRUNC, 0o644)

        source_channel_login_process = subprocess.Popen(
            cmd,
            cwd=str(BASE_DIR),
            env={**os.environ, "PYTHONPATH": str(BASE_DIR / "src"), "PYTHONUNBUFFERED": "1"},
            stdout=log_fd,
            stderr=subprocess.STDOUT,
        )
        os.close(log_fd)
        source_channel_logging_in_channel_id = channel.get("channel_id") or ""
        source_channel_logging_in_account_id = account.get("account_id") or ""
        source_channel_login_err_msg = ""
        source_channel_login_result = ""
        source_channel_login_result_finished_at = ""
        source_channel_login_started_at = datetime.now().isoformat()

        return {
            "status": "success",
            "msg": "已启动 1688 登录页，请在弹出窗口中完成扫码登录。登录成功后系统会自动进行渠道预热并短暂跳转页面，请勿立即关闭浏览器。",
        }
    except Exception as e:
        logger.error(f"Failed to trigger source channel login: {e}")
        source_channel_login_process = None
        source_channel_logging_in_channel_id = ""
        source_channel_logging_in_account_id = ""
        source_channel_login_err_msg = str(e)
        source_channel_login_result = "failed"
        source_channel_login_result_finished_at = datetime.now().isoformat()
        source_channel_login_started_at = ""
        return {"status": "failed", "msg": str(e)}

@app.get("/api/system/regions")
def get_system_regions():
    import urllib.request
    import urllib.parse
    import openpyxl
    import os
    import json
    from pathlib import Path
    
    json_path = Path("config/goofish_regions.json")
    if json_path.exists():
        try:
            with open(json_path, "r", encoding="utf-8") as f:
                data = json.load(f)
            return {"status": "success", "data": data}
        except Exception as e:
            logger.error(f"Failed to read regions cache: {e}")
            
    # 如果缓存不存在，尝试自动从闲管家下载解析并缓存
    url_base = "https://file.goofish.pro/doc/"
    filename = "闲管家省市区.xlsx"
    encoded_url = url_base + urllib.parse.quote(filename)
    dest_path = Path("config/goofish_regions.xlsx")
    
    try:
        dest_path.parent.mkdir(parents=True, exist_ok=True)
        req = urllib.request.Request(
            encoded_url, 
            headers={'User-Agent': 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36'}
        )
        with urllib.request.urlopen(req) as response:
            with open(dest_path, "wb") as f:
                f.write(response.read())
                
        wb = openpyxl.load_workbook(dest_path, read_only=True)
        sheet = wb.active
        
        regions = {}
        # Header: ('省份ID', '省份名称', '城市ID', '城市名称', '地区ID', '地区名称')
        for r_idx, row in enumerate(sheet.iter_rows(values_only=True)):
            if r_idx == 0:
                continue
            if not row or len(row) < 6:
                continue
                
            prov_id, prov_name, city_id, city_name, dist_id, dist_name = row
            if not prov_id or not prov_name:
                continue
                
            prov_id = int(prov_id)
            prov_name = str(prov_name).strip()
            
            if prov_id not in regions:
                regions[prov_id] = {
                    "name": prov_name,
                    "code": prov_id,
                    "cities": {}
                }
                
            if not city_id or not city_name:
                continue
                
            city_id = int(city_id)
            city_name = str(city_name).strip()
            
            if city_id not in regions[prov_id]["cities"]:
                regions[prov_id]["cities"][city_id] = {
                    "name": city_name,
                    "code": city_id,
                    "districts": {}
                }
                
            if not dist_id or not dist_name:
                continue
                
            dist_id = int(dist_id)
            dist_name = str(dist_name).strip()
            
            regions[prov_id]["cities"][city_id]["districts"][dist_id] = {
                "name": dist_name,
                "code": dist_id
            }

        sorted_regions = []
        for p_id in sorted(regions.keys()):
            p_data = regions[p_id]
            sorted_cities = []
            for c_id in sorted(p_data["cities"].keys()):
                c_data = p_data["cities"][c_id]
                sorted_districts = []
                for d_id in sorted(c_data["districts"].keys()):
                    sorted_districts.append(c_data["districts"][d_id])
                c_data["districts"] = sorted_districts
                sorted_cities.append(c_data)
            p_data["cities"] = sorted_cities
            sorted_regions.append(p_data)
            
        with open(json_path, "w", encoding="utf-8") as f:
            json.dump(sorted_regions, f, ensure_ascii=False, indent=4)
            
        return {"status": "success", "data": sorted_regions}
    except Exception as e:
        logger.error(f"Failed to auto download/parse goofish regions: {e}")
        return {"status": "failed", "msg": f"获取省市区失败: {str(e)}"}
    finally:
        if dest_path.exists():
            os.remove(dest_path)

@app.get("/api/system/openapi_categories")
def get_openapi_categories(group: str = None, query: str = None, cat_id: str = None, account_id: str = None):
    import json
    from pathlib import Path
    
    cache_path = Path("config/xianyu_categories.json")
    if not cache_path.exists():
        try:
            from xianyu_tools.xianyu_adapter.publisher_v3 import PublisherV3
            pub = PublisherV3(account_id=account_id)
            categories = pub._load_categories()
        except Exception as e:
            logger.error(f"Failed to trigger categories download: {e}")
            categories = []
    else:
        try:
            with open(cache_path, "r", encoding="utf-8") as f:
                categories = json.load(f)
        except Exception as e:
            logger.error(f"Failed to read category cache: {e}")
            categories = []
            
    if not categories:
        return {"status": "failed", "msg": "获取类目失败，缓存为空且同步异常"}
        
    # 提取所有大类分组 sp_biz_name
    groups_set = set()
    for item in categories:
        g_name = item.get("sp_biz_name")
        if g_name:
            groups_set.add(g_name)
    groups = sorted(list(groups_set))
    
    # 根据 group, query, cat_id 进行过滤
    filtered = []
    query_lower = query.lower().strip() if query else None
    
    seen_ids = set()
    for item in categories:
        item_group = item.get("sp_biz_name")
        item_name = item.get("channel_cat_name", "")
        item_id = item.get("channel_cat_id")
        
        # 0. 精准过滤 ID (如果提供了 cat_id)
        if cat_id and item_id != cat_id:
            continue
            
        # 1. 过滤大类
        if not cat_id and group and item_group != group:
            continue
            
        # 2. 模糊匹配关键字
        if not cat_id and query_lower and query_lower not in item_name.lower():
            continue
            
        # 3. 对发布类目 ID 进行去重，避免相同类目由于被重复分发在多个大类下而导致前端显示重复项
        if item_id in seen_ids:
            continue
        seen_ids.add(item_id)
            
        filtered.append({
            "id": item_id,
            "name": f"{item_group} - {item_name}" if item_group else item_name,
            "group": item_group
        })
        
    # 限制返回的前 100 条
    limited_data = filtered[:100]
    
    return {
        "status": "success",
        "data": {
            "groups": groups,
            "categories": limited_data,
            "total_matches": len(filtered),
            "is_truncated": len(filtered) > 100
        }
    }

# --- 闲鱼登录态自动捕获逻辑与接口 ---
is_logging_in = False
logging_in_account_id = ""
login_err_msg = ""
source_channel_login_process = None
source_channel_logging_in_channel_id = ""
source_channel_logging_in_account_id = ""
source_channel_login_err_msg = ""
source_channel_login_log_path = ""
source_channel_login_result = ""
source_channel_login_result_finished_at = ""
source_channel_login_started_at = ""


def sync_source_channel_login_state():
    global source_channel_login_process, source_channel_logging_in_channel_id, source_channel_logging_in_account_id
    global source_channel_login_err_msg, source_channel_login_log_path
    global source_channel_login_result, source_channel_login_result_finished_at, source_channel_login_started_at

    if source_channel_login_process is None:
        return

    return_code = source_channel_login_process.poll()
    if return_code is None:
        return

    if return_code != 0 and source_channel_login_log_path:
        try:
            log_text = Path(source_channel_login_log_path).read_text(encoding="utf-8", errors="ignore").strip()
            if log_text:
                if "浏览器已关闭，登录流程已取消" in log_text or "浏览器已关闭，操作取消" in log_text:
                    source_channel_login_err_msg = "你已关闭 1688 登录浏览器，本次登录已取消。"
                    source_channel_login_result = "cancelled"
                elif "ERR_TUNNEL_CONNECTION_FAILED" in log_text:
                    source_channel_login_err_msg = "1688 登录页打开失败，当前浏览器网络代理环境不可用，请稍后重试。"
                    source_channel_login_result = "failed"
                elif "无法打开 1688 登录页" in log_text:
                    source_channel_login_err_msg = "\n".join(log_text.splitlines()[-3:])
                    source_channel_login_result = "failed"
                else:
                    source_channel_login_err_msg = "\n".join(log_text.splitlines()[-10:])
                    source_channel_login_result = "failed"
            else:
                source_channel_login_err_msg = f"1688 登录进程异常退出，退出码：{return_code}"
                source_channel_login_result = "failed"
        except Exception:
            source_channel_login_err_msg = f"1688 登录进程异常退出，退出码：{return_code}"
            source_channel_login_result = "failed"
    elif return_code == 0:
        source_channel_login_err_msg = ""
        source_channel_login_result = "success"

    source_channel_login_result_finished_at = datetime.now().isoformat()

    source_channel_login_process = None
    source_channel_logging_in_channel_id = ""
    source_channel_logging_in_account_id = ""
    source_channel_login_started_at = ""

async def run_xianyu_login_capture(account_id: str | None = None):
    global is_logging_in, logging_in_account_id, login_err_msg
    is_logging_in = True
    logging_in_account_id = account_id or ""
    login_err_msg = ""
    
    from xianyu_tools.xianyu_adapter.browser_transport import PlaywrightBrowserTransport, PlaywrightBrowserConfig, default_desktop_context_options
    from xianyu_tools.config import settings
    account_cfg = get_openapi_account(settings.get_openapi_raw_config(), account_id)
    state_file = account_cfg.get("state_file") or f"xianyu_state_{account_cfg.get('id') or 'default'}.json"
    
    transport = PlaywrightBrowserTransport(
        config=PlaywrightBrowserConfig(
            headless=False,
            browser_channel="chrome",
            launch_args=["--start-maximized"],
            context_options=default_desktop_context_options(),
        )
    )
    
    try:
        async with transport._playwright_context() as playwright:
            browser = await playwright.chromium.launch(
                channel="chrome",
                headless=False,
                args=["--start-maximized"],
            )
            try:
                context = await transport._new_context(browser)
                page = await context.new_page()
                
                logger.info("[Xianyu Login] Opening Goofish homepage for QR login...")
                await page.goto(
                    "https://www.goofish.com/",
                    wait_until="domcontentloaded",
                    timeout=60000,
                )

                await page.wait_for_timeout(2500)

                # 直接点击闲鱼首页头部右上角登录入口，优先避开页面其它同名文案。
                login_selectors = [
                    'div[class*="user-order-container"] a:has-text("登录")',
                    'a:has-text("登录")',
                    "text=立即登录",
                    "text=登录",
                ]
                login_opened = False
                for selector in login_selectors:
                    try:
                        locator = page.locator(selector).first
                        if await locator.count():
                            logger.info(f"[Xianyu Login] Trying QR login opener selector: {selector}")
                            await locator.click(timeout=3000)
                            await page.wait_for_timeout(1200)
                            if await page.locator("text=手机扫码安全登录").count():
                                logger.info(f"[Xianyu Login] Opened QR login modal via selector: {selector}")
                                login_opened = True
                                break
                    except Exception:
                        continue

                if not login_opened:
                    logger.warning("[Xianyu Login] QR login modal was not auto-opened; page remains on Goofish homepage.")
                
                logged_in = False
                # 轮询 10 分钟，登录进行中时用更短间隔尽快感知成功状态
                for _ in range(1200):
                    await asyncio.sleep(0.5)
                    
                    if page.is_closed():
                        break
                        
                    cookies = await context.cookies("https://www.goofish.com/")
                    cookie_names = {c["name"] for c in cookies}
                    
                    # 检查是否包含 unb 或者是 tracknick，表明已登录成功
                    if "unb" in cookie_names or "tracknick" in cookie_names:
                        logger.info("[Xianyu Login] Login detected! Capturing state snapshot...")
                        await asyncio.sleep(0.3)
                        
                        from xianyu_tools.xianyu_adapter.state_exporter import PlaywrightStateExporter, StateExportConfig, build_snapshot
                        exporter = PlaywrightStateExporter(
                            config=StateExportConfig(
                                output_file=state_file,
                                browser_channel="chrome",
                                headless=False,
                            )
                        )
                        page_data = await exporter._capture_page_data(page)
                        headers = await exporter._capture_headers(page)
                        final_cookies = await context.cookies("https://www.goofish.com/")
                        
                        snapshot = build_snapshot(page.url, page_data, headers, final_cookies)
                        
                        # 自动保存到本地文件
                        import json
                        with open(state_file, "w", encoding="utf-8") as f:
                            json.dump(snapshot, f, ensure_ascii=False, indent=2)
                        
                        logger.info("[Xianyu Login] State file successfully captured and saved.")
                        is_logging_in = False
                        logged_in = True
                        break
                
                if not logged_in and not page.is_closed():
                    logger.warning("[Xianyu Login] Session ended or timed out without login success.")
            except Exception as inner_e:
                logger.error(f"[Xianyu Login] Error in login runner: {inner_e}")
                login_err_msg = str(inner_e)
            finally:
                try:
                    if not page.is_closed():
                        await page.close()
                except Exception:
                    pass
                await browser.close()
    except Exception as e:
        logger.error(f"[Xianyu Login] Failed to launch playwright browser: {e}")
        login_err_msg = str(e)
    finally:
        is_logging_in = False
        logging_in_account_id = ""

@app.get("/api/system/xianyu_login_status")
def get_xianyu_login_status(account_id: str = None):
    try:
        from xianyu_tools.config import settings
        from xianyu_tools.xianyu_adapter.state_inspector import inspect_state_file
        account_cfg = get_openapi_account(settings.get_openapi_raw_config(), account_id)
        state_file = account_cfg.get("state_file") or "xianyu_state.json"
        report = inspect_state_file(state_file)
        return {
            "status": "success",
            "data": {
                "report": report,
                "is_logging_in": is_logging_in and logging_in_account_id == (account_cfg.get("id") or ""),
                "err_msg": login_err_msg
            }
        }
    except Exception as e:
        logger.error(f"Failed to check xianyu login status: {e}")
        return {"status": "failed", "msg": str(e)}

@app.post("/api/system/xianyu_login_trigger")
async def trigger_xianyu_login(payload: dict = None):
    global is_logging_in
    if is_logging_in:
        return {"status": "failed", "msg": "当前正在执行登录，请勿重复操作"}

    try:
        from xianyu_tools.config import settings
        raw_cfg = settings.get_openapi_raw_config()
        requested_account_id = (payload or {}).get("account_id")
        account_cfg = get_openapi_account(raw_cfg, requested_account_id or (raw_cfg or {}).get("active_account_id"))
        account_id = account_cfg.get("id")
    except Exception:
        account_id = None

    asyncio.create_task(run_xianyu_login_capture(account_id))
    return {"status": "success", "msg": "已启动闲鱼登录浏览器，请使用手机扫码完成安全登录"}

app.mount("/", NoCacheStaticFiles(directory=str(WEB_DIR), html=True), name="web")
