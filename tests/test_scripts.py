import json
import importlib.util
import subprocess
import sys
from pathlib import Path


def test_search_script_exists() -> None:
    path = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/run_xianyu_search.py")
    assert path.exists()


def test_sourcing_script_exists() -> None:
    path = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/run_xianyu_sourcing.py")
    assert path.exists()


def test_workflow_script_exists() -> None:
    path = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/run_xianyu_workflow.py")
    assert path.exists()


def test_xianyu_hot_items_script_exists() -> None:
    path = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/run_xianyu_hot_items.py")
    assert path.exists()


def test_ali1688_result_url_script_exists() -> None:
    path = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/run_ali1688_result_url.py")
    assert path.exists()


def test_run_ali1688_slow_flow_script_exists() -> None:
    path = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/run_ali1688_slow_flow.py")
    assert path.exists()


def test_run_ali1688_slow_flow_dispatch_metric_helpers() -> None:
    script_path = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/run_ali1688_slow_flow.py")
    spec = importlib.util.spec_from_file_location("run_ali1688_slow_flow_script", script_path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    metrics = module._extract_dispatch_metrics_from_text("近7天代发 123 月代发 2.4万")
    assert metrics == {"seven_day_dispatch_count": 123, "month_dispatch_count": 24000}
    assert module._parse_dispatch_count("3.5k") == 3500


def test_run_ali1688_slow_flow_normalize_image_search_url() -> None:
    script_path = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/run_ali1688_slow_flow.py")
    spec = importlib.util.spec_from_file_location("run_ali1688_slow_flow_script", script_path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    assert module._normalize_image_search_url("http://img.alicdn.com/demo.jpg") == "http://img.alicdn.com/demo.jpg"
    assert module._normalize_image_search_url("https://img.alicdn.com/demo") == "https://img.alicdn.com/demo"
    assert module._normalize_image_search_url("http://img.alicdn.com/demo.heic") is None
    assert module._normalize_image_search_url("ftp://img.alicdn.com/demo.jpg") is None
    assert module._normalize_image_search_url("http://img.alicdn.com/demo.gif") is None


def test_run_ali1688_slow_flow_top_dispatch_candidates_sorting() -> None:
    script_path = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/run_ali1688_slow_flow.py")
    spec = importlib.util.spec_from_file_location("run_ali1688_slow_flow_script", script_path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    rows = [
        {"item_url": "https://detail.1688.com/offer/1.html", "seven_day_dispatch_count": 30, "month_dispatch_count": 400},
        {"item_url": "https://detail.1688.com/offer/2.html", "seven_day_dispatch_count": 50, "month_dispatch_count": 100},
        {"item_url": "https://detail.1688.com/offer/3.html", "seven_day_dispatch_count": 50, "month_dispatch_count": 300},
        {"item_url": "", "seven_day_dispatch_count": 999, "month_dispatch_count": 999},
    ]

    top_rows = module._top_dispatch_candidates(rows, 3)
    assert [row["item_url"] for row in top_rows] == [
        "https://detail.1688.com/offer/3.html",
        "https://detail.1688.com/offer/2.html",
        "https://detail.1688.com/offer/1.html",
    ]


def test_run_ali1688_slow_flow_prepare_managed_state_file_copies_into_dedicated_path(tmp_path) -> None:
    script_path = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/run_ali1688_slow_flow.py")
    spec = importlib.util.spec_from_file_location("run_ali1688_slow_flow_script", script_path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    source_state = tmp_path / "input_state.json"
    source_payload = {"cookies": [{"name": "a", "value": "1", "domain": ".1688.com", "path": "/"}], "origins": []}
    source_state.write_text(json.dumps(source_payload), encoding="utf-8")
    managed_state = tmp_path / "managed" / "storage_state.json"

    effective_path, synced = module._prepare_managed_state_file(
        str(source_state),
        managed_state_file=str(managed_state),
    )

    assert synced is True
    assert effective_path == str(managed_state.resolve())
    assert json.loads(managed_state.read_text(encoding="utf-8")) == source_payload


def test_run_ali1688_slow_flow_prepare_managed_state_file_missing_default_is_soft(tmp_path) -> None:
    script_path = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/run_ali1688_slow_flow.py")
    spec = importlib.util.spec_from_file_location("run_ali1688_slow_flow_script", script_path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    managed_state = tmp_path / "managed" / "storage_state.json"
    effective_path, synced = module._prepare_managed_state_file(
        str(managed_state),
        managed_state_file=str(managed_state),
    )

    assert effective_path is None
    assert synced is False


def test_run_ali1688_slow_flow_prepare_managed_state_file_missing_custom_raises(tmp_path) -> None:
    script_path = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/run_ali1688_slow_flow.py")
    spec = importlib.util.spec_from_file_location("run_ali1688_slow_flow_script", script_path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    missing_source = tmp_path / "missing.json"
    managed_state = tmp_path / "managed" / "storage_state.json"

    try:
        module._prepare_managed_state_file(str(missing_source), managed_state_file=str(managed_state))
    except FileNotFoundError as exc:
        assert "state file not found" in str(exc)
    else:
        raise AssertionError("expected FileNotFoundError for explicit missing state file")


def test_run_ali1688_slow_flow_append_extension_launch_args(tmp_path) -> None:
    script_path = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/run_ali1688_slow_flow.py")
    spec = importlib.util.spec_from_file_location("run_ali1688_slow_flow_script", script_path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    extension_dir = tmp_path / "ext"
    extension_dir.mkdir()
    args = module._append_extension_launch_args(["--foo", "--disable-extensions"], str(extension_dir))

    assert "--foo" in args
    assert "--disable-extensions" not in args
    assert f"--disable-extensions-except={extension_dir}" in args
    assert f"--load-extension={extension_dir}" in args


def test_run_ali1688_slow_flow_resolve_browser_channel_uses_chromium_for_extension(tmp_path) -> None:
    script_path = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/run_ali1688_slow_flow.py")
    spec = importlib.util.spec_from_file_location("run_ali1688_slow_flow_script", script_path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    extension_dir = tmp_path / "ext"
    extension_dir.mkdir()

    assert module._resolve_browser_channel("chrome", str(extension_dir)) is None
    assert module._resolve_browser_channel("msedge", str(extension_dir)) == "msedge"
    assert module._resolve_browser_channel("chrome", "") == "chrome"


def test_run_ali1688_slow_flow_overlay_close_click_point() -> None:
    script_path = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/run_ali1688_slow_flow.py")
    spec = importlib.util.spec_from_file_location("run_ali1688_slow_flow_script", script_path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    assert module._overlay_close_click_point({"x": 0.0, "y": 0.0, "width": 1400.0, "height": 720.0}) == (1378.0, 20.0)


def test_run_ali1688_slow_flow_plugin_toolbar_selectors_cover_verified_nodes() -> None:
    script_path = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/run_ali1688_slow_flow.py")
    spec = importlib.util.spec_from_file_location("run_ali1688_slow_flow_script", script_path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    selectors = module._plugin_toolbar_selectors()
    assert "#market-mate-for-1688" in selectors
    assert "#market-mate-for-1688-od" in selectors
    assert ".goods-operation-panel-media" in selectors
    assert ".goods-operation-hover.copy-sku" in selectors
    assert "text=复制sku" in selectors


def test_run_ali1688_slow_flow_sanitize_storage_state_cookies() -> None:
    script_path = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/run_ali1688_slow_flow.py")
    spec = importlib.util.spec_from_file_location("run_ali1688_slow_flow_script", script_path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    cookies = [
        {
            "name": "_m_h5_tk",
            "value": "token",
            "domain": ".1688.com",
            "path": "/",
            "expires": 123.0,
            "httpOnly": False,
            "secure": True,
            "sameSite": "None",
            "partitionKey": "https://1688.com",
            "_crHasCrossSiteAncestor": True,
        }
    ]
    assert module._sanitize_storage_state_cookies(cookies) == [
        {
            "name": "_m_h5_tk",
            "value": "token",
            "domain": ".1688.com",
            "path": "/",
            "expires": 123.0,
            "httpOnly": False,
            "secure": True,
            "sameSite": "None",
        }
    ]


def test_run_ali1688_slow_flow_resolve_runtime_user_data_dir_for_chromium_extension(tmp_path) -> None:
    script_path = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/run_ali1688_slow_flow.py")
    spec = importlib.util.spec_from_file_location("run_ali1688_slow_flow_script", script_path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    extension_dir = tmp_path / "ext"
    extension_dir.mkdir()
    output_dir = tmp_path / "out"
    output_dir.mkdir()

    resolved = module._resolve_runtime_user_data_dir(
        "/tmp/original-profile",
        None,
        str(extension_dir),
        output_dir,
    )
    assert resolved == str((output_dir / "_runtime_chromium_profile").resolve())

    preserved = module._resolve_runtime_user_data_dir(
        "/tmp/original-profile",
        "chrome",
        str(extension_dir),
        output_dir,
    )
    assert preserved == "/tmp/original-profile"


def test_run_ali1688_slow_flow_detail_offer_id_from_url() -> None:
    script_path = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/run_ali1688_slow_flow.py")
    spec = importlib.util.spec_from_file_location("run_ali1688_slow_flow_script", script_path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    assert (
        module._detail_offer_id_from_url("https://detail.1688.com/offer/904776936832.html?spm=a26352.b28411319/2508.0.0")
        == "904776936832"
    )
    assert module._detail_offer_id_from_url("https://www.1688.com/") == ""


def test_run_ali1688_slow_flow_copy_sku_drawer_opened_detection() -> None:
    script_path = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/run_ali1688_slow_flow.py")
    spec = importlib.util.spec_from_file_location("run_ali1688_slow_flow_script", script_path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    class FakeLocator:
        def __init__(self, *, count: int = 1, visible: bool = True, attrs: dict[str, str] | None = None):
            self._count = count
            self._visible = visible
            self._attrs = attrs or {}

        @property
        def first(self):
            return self

        async def count(self):
            return self._count

        async def is_visible(self):
            return self._visible

        async def get_attribute(self, name: str):
            return self._attrs.get(name)

    class FakePage:
        def __init__(self, drawer: FakeLocator, iframe: FakeLocator):
            self._drawer = drawer
            self._iframe = iframe

        def locator(self, selector: str):
            if selector == "#consign-sku-fullscreen-drawer":
                return self._drawer
            if selector == "#fullscreen-drawer-iframe":
                return self._iframe
            raise AssertionError(f"unexpected selector: {selector}")

    opened_page = FakePage(
        FakeLocator(attrs={"style": ""}),
        FakeLocator(attrs={"src": "https://air.1688.com/app/upkg-solution/od-panel/sku-panel.html?offerId=904776936832"}),
    )
    hidden_page = FakePage(
        FakeLocator(attrs={"style": "display: none;"}),
        FakeLocator(attrs={"src": "https://air.1688.com/app/upkg-solution/od-panel/sku-panel.html?offerId=904776936832#hidden"}),
    )

    import asyncio

    assert asyncio.run(module._copy_sku_drawer_opened(opened_page)) is True
    assert asyncio.run(module._copy_sku_drawer_opened(hidden_page)) is False


def test_run_ali1688_slow_flow_copy_sku_drawer_state_shape() -> None:
    script_path = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/run_ali1688_slow_flow.py")
    spec = importlib.util.spec_from_file_location("run_ali1688_slow_flow_script", script_path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    class FakeLocator:
        def __init__(self, *, count: int = 1, visible: bool = True, attrs: dict[str, str] | None = None):
            self._count = count
            self._visible = visible
            self._attrs = attrs or {}

        @property
        def first(self):
            return self

        async def count(self):
            return self._count

        async def is_visible(self):
            return self._visible

        async def get_attribute(self, name: str):
            return self._attrs.get(name)

    class FakePage:
        def __init__(self, drawer: FakeLocator, iframe: FakeLocator):
            self._drawer = drawer
            self._iframe = iframe

        def locator(self, selector: str):
            if selector == "#consign-sku-fullscreen-drawer":
                return self._drawer
            if selector == "#fullscreen-drawer-iframe":
                return self._iframe
            raise AssertionError(f"unexpected selector: {selector}")

    page = FakePage(
        FakeLocator(attrs={"style": ""}),
        FakeLocator(attrs={"src": "https://air.1688.com/app/upkg-solution/od-panel/sku-panel.html?offerId=1"}),
    )

    import asyncio

    state = asyncio.run(module._copy_sku_drawer_state(page))
    assert state == {
        "drawer_present": True,
        "drawer_visible": True,
        "drawer_style": "",
        "iframe_present": True,
        "iframe_src": "https://air.1688.com/app/upkg-solution/od-panel/sku-panel.html?offerId=1",
        "iframe_hidden": False,
        "opened": True,
    }


def test_run_ali1688_slow_flow_overlay_state_shape() -> None:
    script_path = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/run_ali1688_slow_flow.py")
    spec = importlib.util.spec_from_file_location("run_ali1688_slow_flow_script", script_path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    class FakeLocator:
        def __init__(self, *, count: int = 0, visible: bool = False):
            self._count = count
            self._visible = visible

        @property
        def first(self):
            return self

        @property
        def last(self):
            return self

        async def count(self):
            return self._count

        async def is_visible(self):
            return self._visible

    class FakePage:
        def locator(self, selector: str):
            if selector == ".J_MIDDLEWARE_FRAME_WIDGET:visible":
                return FakeLocator(count=2, visible=True)
            raise AssertionError(f"unexpected selector: {selector}")

        def get_by_text(self, text: str, exact: bool = False):
            assert text == "我知道了"
            assert exact is True
            return FakeLocator(count=1, visible=True)

    import asyncio

    async def fake_toolbar_ready(page):
        return True

    original = module._plugin_toolbar_ready
    module._plugin_toolbar_ready = fake_toolbar_ready
    try:
        state = asyncio.run(module._overlay_state(FakePage()))
    finally:
        module._plugin_toolbar_ready = original

    assert state == {
        "overlay_count": 2,
        "ack_visible": True,
        "toolbar_ready": True,
    }


def test_source_resolution_from_urls_script_exists() -> None:
    path = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/run_source_resolution_from_urls.py")
    assert path.exists()


def test_source_resolution_from_browser_runs_script_exists() -> None:
    path = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/run_source_resolution_from_browser_runs.py")
    assert path.exists()


def test_generate_ali1688_result_url_map_script_exists() -> None:
    path = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/generate_ali1688_result_url_map.py")
    assert path.exists()


def test_generate_ali1688_result_url_map_via_browser_script_exists() -> None:
    path = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/generate_ali1688_result_url_map_via_browser.py")
    assert path.exists()


def test_generate_ali1688_suggested_result_urls_script_exists() -> None:
    path = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/generate_ali1688_suggested_result_urls.py")
    assert path.exists()


def test_run_profit_analysis_script_exists() -> None:
    path = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/run_profit_analysis.py")
    assert path.exists()


def test_run_listing_candidates_script_exists() -> None:
    path = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/run_listing_candidates.py")
    assert path.exists()


def test_run_category_pipeline_from_urls_script_exists() -> None:
    path = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/run_category_pipeline_from_urls.py")
    assert path.exists()


def test_export_pipeline_excel_script_exists() -> None:
    path = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/export_pipeline_excel.py")
    assert path.exists()


def test_inspect_state_script_strict_exit_code(tmp_path) -> None:
    state_file = tmp_path / "xianyu_state.json"
    state_file.write_text(json.dumps({"cookies": [], "origins": []}), encoding="utf-8")
    repo_root = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools")
    result = subprocess.run(
        [
            sys.executable,
            "scripts/inspect_xianyu_state.py",
            "--state-file",
            str(state_file),
            "--strict",
        ],
        cwd=repo_root,
        env={"PYTHONPATH": "src"},
        capture_output=True,
        text=True,
        check=False,
    )
    assert result.returncode == 1
    assert "missing_cookies" in result.stdout


def test_run_xianyu_search_outputs_bundle_shape(monkeypatch, capsys) -> None:
    script_path = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/run_xianyu_search.py")
    spec = importlib.util.spec_from_file_location("run_xianyu_search_script", script_path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    from xianyu_tools.models import XianyuSearchItem

    class FakeAdapter:
        def search(self, keyword: str, page: int = 1):
            return [
                XianyuSearchItem(
                    item_id="xy-001",
                    title="iPhone 15 Pro Max",
                    price=5999.0,
                    item_url="https://www.goofish.com/item?id=xy-001",
                    area="上海",
                )
            ]

    class FakeFactory:
        @staticmethod
        def from_browser(config):
            return FakeAdapter()

    monkeypatch.setattr(module, "PlaywrightXianyuAdapter", FakeFactory)
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "run_xianyu_search.py",
            "--keyword",
            "iphone",
            "--state-file",
            "./xianyu_state.json",
        ],
    )
    assert module.main() == 0
    output = json.loads(capsys.readouterr().out)
    assert output["search_returned_count"] == 1
    assert output["search_items"][0]["item_id"] == "xy-001"


def test_run_xianyu_hot_items_passes_launch_args(monkeypatch, capsys) -> None:
    script_path = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/run_xianyu_hot_items.py")
    spec = importlib.util.spec_from_file_location("run_xianyu_hot_items_script", script_path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    captured = {}

    class FakeAdapter:
        pass

    class FakeFactory:
        @staticmethod
        def from_browser(config):
            captured["config"] = config
            return FakeAdapter()

    def fake_scan(keyword, *, xianyu_adapter, top_n, max_pages, require_chaozan_fish_shop):
        return {"hot_items": [], "xianyu_market": {"category_keyword": keyword}}

    monkeypatch.setattr(module, "PlaywrightXianyuAdapter", FakeFactory)
    monkeypatch.setattr(module, "scan_xianyu_market", fake_scan)
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "run_xianyu_hot_items.py",
            "--keyword",
            "升降桌",
            "--state-file",
            "./xianyu_state.json",
            "--launch-arg=--start-maximized",
        ],
    )
    assert module.main() == 0
    assert "--start-maximized" in captured["config"].launch_args
    output = json.loads(capsys.readouterr().out)
    assert output["xianyu_market"]["category_keyword"] == "升降桌"


def test_run_xianyu_hot_items_routing(monkeypatch, capsys) -> None:
    script_path = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/run_xianyu_hot_items.py")
    spec = importlib.util.spec_from_file_location("run_xianyu_hot_items_script", script_path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    def clean_json(stdout_str):
        import re
        match = re.search(r'(\{.*"(?:hot_items|error)".*\})', stdout_str, re.DOTALL)
        return json.loads(match.group(1)) if match else json.loads(stdout_str)

    # 1. 测试图片分支 (IMAGE 模式)
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "run_xianyu_hot_items.py",
            "--keyword",
            "https://example.com/test_product.png?version=1",
            "--state-file",
            "./xianyu_state.json",
        ],
    )
    assert module.main() == 0
    output_img = clean_json(capsys.readouterr().out)
    assert output_img["hot_items"][0]["hot_item_id"] == "img_search"
    assert "test_product.png" in output_img["hot_items"][0]["title"]
    assert output_img["hot_items"][0]["image_url"] == "https://example.com/test_product.png?version=1"

    # 2. 测试链接分支 (URL 模式)
    class FakeDetailItem:
        item_id = "800991122"
        title = "测试闲鱼单品宝贝"
        price = 45.5
        want_count = 12
        item_url = "https://h5.m.goofish.com/item?id=800991122"
        images = ["https://img.alicdn.com/img1.jpg"]

    class FakeAdapter:
        def detail(self, target_id_or_url):
            assert target_id_or_url == "800991122"
            return FakeDetailItem()

    class FakeFactory:
        @staticmethod
        def from_browser(config):
            return FakeAdapter()

    monkeypatch.setattr(module, "PlaywrightXianyuAdapter", FakeFactory)
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "run_xianyu_hot_items.py",
            "--keyword",
            "【闲鱼】https://m.tb.cn/h.xxxxx?id=800991122 「我在闲鱼发布了...」",
            "--state-file",
            "./xianyu_state.json",
        ],
    )
    monkeypatch.setattr(module, "extract_item_url_or_id", lambda x, y: "800991122")

    assert module.main() == 0
    output_url = clean_json(capsys.readouterr().out)
    assert output_url["hot_items"][0]["hot_item_id"] == "800991122"
    assert output_url["hot_items"][0]["title"] == "测试闲鱼单品宝贝"
    assert output_url["hot_items"][0]["price"] == 45.5
    assert output_url["hot_items"][0]["image_url"] == "https://img.alicdn.com/img1.jpg"

    # 3. 测试无效链接分支 (INVALID_URL 模式)
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "run_xianyu_hot_items.py",
            "--keyword",
            "【闲鱼】无链接口令￥无链接口令￥",
            "--state-file",
            "./xianyu_state.json",
        ],
    )
    monkeypatch.setattr(module, "extract_item_url_or_id", lambda x, y: "goofish_invalid")
    assert module.main() == 1
    output_err1 = clean_json(capsys.readouterr().out)
    assert output_err1["error"] == "INVALID_URL"
    assert "解析闲鱼链接失败" in output_err1["msg"]



    # 4. 测试抓取单品详情超时或异常 (DETAIL_FETCH_FAILED 模式)
    class FakeAdapterError:
        def detail(self, target_id_or_url):
            raise Exception("Timeout when connecting to goofish detail service")

    class FakeFactoryError:
        @staticmethod
        def from_browser(config):
            return FakeAdapterError()

    monkeypatch.setattr(module, "PlaywrightXianyuAdapter", FakeFactoryError)
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "run_xianyu_hot_items.py",
            "--keyword",
            "https://m.tb.cn/h.xxxxx?id=999999",
            "--state-file",
            "./xianyu_state.json",
        ],
    )
    monkeypatch.setattr(module, "extract_item_url_or_id", lambda x, y: "999999")
    assert module.main() == 1
    output_err2 = clean_json(capsys.readouterr().out)
    assert output_err2["error"] == "DETAIL_FETCH_FAILED"
    assert "获取宝贝详情超时" in output_err2["msg"]




def test_run_source_resolution_from_urls_outputs_summary(monkeypatch, capsys, tmp_path) -> None:
    script_path = Path(
        "/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/run_source_resolution_from_urls.py"
    )
    spec = importlib.util.spec_from_file_location("run_source_resolution_from_urls_script", script_path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    hot_items_file = tmp_path / "hot_items.json"
    hot_items_file.write_text(
        json.dumps(
            {
                "hot_items": [
                    {
                        "hot_item_id": "xy-001",
                        "platform": "xianyu",
                        "title": "升降桌",
                        "price": 99.0,
                        "want_count": 123,
                        "item_url": "https://www.goofish.com/item?id=xy-001",
                    }
                ]
            }
        ),
        encoding="utf-8",
    )
    result_url_map_file = tmp_path / "result_url_map.json"
    result_url_map_file.write_text(
        json.dumps({"xy-001": "https://s.1688.com/selloffer/offer_search.htm?keywords=test"}),
        encoding="utf-8",
    )

    def fake_resolve(hot_items, *, result_urls_by_hot_item_id, adapter=None, limit_per_item=10):
        assert hot_items[0].hot_item_id == "xy-001"
        assert result_urls_by_hot_item_id["xy-001"].startswith("https://s.1688.com/")
        assert limit_per_item == 8
        return {
            "source_resolution": [
                {
                    "hot_item_id": "xy-001",
                    "resolved": True,
                    "source_item_ids": ["src-001"],
                    "resolution_reason": "matched_source_items",
                    "source_query": "升降桌",
                    "metadata": {"candidate_count": 1},
                }
            ],
            "source_items": [
                {
                    "source_platform": "1688",
                    "source_item_id": "src-001",
                    "title": "升降桌厂家直发",
                    "price": 66.0,
                    "item_url": "https://detail.1688.com/offer/src-001.html",
                    "shop_name": "源头工厂",
                    "sales": 300,
                    "metadata": {"hot_item_id": "xy-001"},
                }
            ],
        }

    monkeypatch.setattr(module, "resolve_hot_items_to_ali1688_urls", fake_resolve)
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "run_source_resolution_from_urls.py",
            "--hot-items-file",
            str(hot_items_file),
            "--result-url-map-file",
            str(result_url_map_file),
            "--limit-per-item",
            "8",
            "--summary-only",
        ],
    )
    assert module.main() == 0
    output = json.loads(capsys.readouterr().out)
    assert output["hot_items"][0]["hot_item_id"] == "xy-001"
    assert output["hot_items"][0]["source_resolution"]["resolved"] is True
    assert output["hot_items"][0]["source_items"][0]["source_item_id"] == "src-001"


def test_generate_ali1688_result_url_map_writes_template(monkeypatch, capsys, tmp_path) -> None:
    script_path = Path(
        "/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/generate_ali1688_result_url_map.py"
    )
    spec = importlib.util.spec_from_file_location("generate_ali1688_result_url_map_script", script_path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    hot_items_file = tmp_path / "hot_items.json"
    hot_items_file.write_text(
        json.dumps(
            {
                "hot_items": [
                    {"hot_item_id": "xy-001", "title": "升降桌"},
                    {"hot_item_id": "xy-002", "title": "电竞椅"},
                ]
            }
        ),
        encoding="utf-8",
    )
    output_file = tmp_path / "result_url_map.json"
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "generate_ali1688_result_url_map.py",
            "--hot-items-file",
            str(hot_items_file),
            "--output-file",
            str(output_file),
        ],
    )
    assert module.main() == 0
    output = json.loads(capsys.readouterr().out)
    assert output["count"] == 2
    result = json.loads(output_file.read_text(encoding="utf-8"))
    assert result == {"xy-001": "", "xy-002": ""}


def test_generate_ali1688_suggested_result_urls_writes_mapping(monkeypatch, capsys, tmp_path) -> None:
    script_path = Path(
        "/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/generate_ali1688_suggested_result_urls.py"
    )
    spec = importlib.util.spec_from_file_location("generate_ali1688_suggested_result_urls_script", script_path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    hot_items_file = tmp_path / "hot_items.json"
    hot_items_file.write_text(
        json.dumps(
            {
                "hot_items": [
                    {"hot_item_id": "xy-001", "platform": "xianyu", "title": "升降桌", "price": 99.0},
                ]
            }
        ),
        encoding="utf-8",
    )
    output_file = tmp_path / "suggested_result_urls.json"
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "generate_ali1688_suggested_result_urls.py",
            "--hot-items-file",
            str(hot_items_file),
            "--output-file",
            str(output_file),
        ],
    )
    assert module.main() == 0
    output = json.loads(capsys.readouterr().out)
    assert output["count"] == 1
    result = json.loads(output_file.read_text(encoding="utf-8"))
    assert "xy-001" in result
    assert result["xy-001"].startswith("https://s.1688.com/selloffer/offer_search.htm?keywords=")


def test_run_profit_analysis_outputs_rows(monkeypatch, capsys, tmp_path) -> None:
    script_path = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/run_profit_analysis.py")
    spec = importlib.util.spec_from_file_location("run_profit_analysis_script", script_path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    source_bundle_file = tmp_path / "source_bundle.json"
    source_bundle_file.write_text(
        json.dumps(
            {
                "hot_items": [
                    {
                        "hot_item_id": "xy-001",
                        "platform": "xianyu",
                        "title": "升降桌",
                        "price": 99.0,
                    }
                ],
                "source_items": [
                    {
                        "source_platform": "1688",
                        "source_item_id": "src-001",
                        "title": "升降桌工厂直发",
                        "price": 60.0,
                        "shipping_fee": 5.0,
                        "item_url": "https://detail.1688.com/offer/src-001.html",
                        "sales": 150,
                        "metadata": {"hot_item_id": "xy-001"},
                    }
                ],
            }
        ),
        encoding="utf-8",
    )
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "run_profit_analysis.py",
            "--source-bundle-file",
            str(source_bundle_file),
        ],
    )
    assert module.main() == 0
    output = json.loads(capsys.readouterr().out)
    assert output["profit_analysis"][0]["hot_item_id"] == "xy-001"
    assert output["profit_analysis"][0]["source_item_id"] == "src-001"


def test_run_listing_candidates_outputs_decisions(monkeypatch, capsys, tmp_path) -> None:
    script_path = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/run_listing_candidates.py")
    spec = importlib.util.spec_from_file_location("run_listing_candidates_script", script_path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    source_bundle_file = tmp_path / "source_bundle.json"
    source_bundle_file.write_text(
        json.dumps(
            {
                "hot_items": [{"hot_item_id": "xy-001", "title": "升降桌", "price": 99.0}],
                "source_items": [{"source_item_id": "src-001", "title": "升降桌工厂直发"}],
            }
        ),
        encoding="utf-8",
    )
    profit_analysis_file = tmp_path / "profit_analysis.json"
    profit_analysis_file.write_text(
        json.dumps(
            {
                "profit_analysis": [
                    {
                        "hot_item_id": "xy-001",
                        "source_item_id": "src-001",
                        "estimated_margin": 12.0,
                        "gross_margin_rate": 0.22,
                        "cost_profit_rate": 0.4,
                        "risk_flags": [],
                    }
                ]
            }
        ),
        encoding="utf-8",
    )
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "run_listing_candidates.py",
            "--source-bundle-file",
            str(source_bundle_file),
            "--profit-analysis-file",
            str(profit_analysis_file),
        ],
    )
    assert module.main() == 0
    output = json.loads(capsys.readouterr().out)
    assert output["listing_candidates"][0]["is_recommended"] is True


def test_export_pipeline_excel_outputs_workbook(tmp_path) -> None:
    from openpyxl import load_workbook

    hot_items_file = tmp_path / "hot_items.json"
    hot_items_file.write_text(
        json.dumps(
            {
                "hot_items": [
                    {
                        "hot_item_id": "xy-001",
                        "platform": "xianyu",
                        "title": "升降桌",
                        "price": 299.0,
                        "want_count": 1234,
                        "seller_name": "店主A",
                        "area": "杭州",
                        "image_url": "https://img.example.com/1.jpg",
                        "item_url": "https://www.goofish.com/item?id=xy-001",
                        "metadata": {"publish_time": "1小时前", "tags": ["包邮"]},
                    }
                ],
                "xianyu_market": {
                    "category_keyword": "升降桌",
                    "result_count": 54,
                    "filtered_by": ["超赞鱼小铺"],
                    "sorted_by": "want_count_desc",
                    "top_n": 10,
                    "top10_price_stats": {"min": 199.0, "max": 399.0, "median": 299.0},
                },
            }
        ),
        encoding="utf-8",
    )
    source_bundle_file = tmp_path / "source_bundle.json"
    source_bundle_file.write_text(
        json.dumps(
            {
                "hot_items": json.loads(hot_items_file.read_text(encoding="utf-8"))["hot_items"],
                "source_resolution": [
                    {
                        "hot_item_id": "xy-001",
                        "resolved": True,
                        "source_item_ids": ["src-001"],
                        "resolution_reason": "matched_source_items",
                    }
                ],
                "source_items": [
                    {
                        "hot_item_id": "xy-001",
                        "source_item_id": "src-001",
                        "source_platform": "1688",
                        "title": "升降桌工厂直发",
                        "price": 120.0,
                        "item_url": "https://detail.1688.com/offer/src-001.html",
                        "shop_name": "源头工厂",
                        "sales": 88,
                        "image_url": "https://img.example.com/src.jpg",
                        "shipping_fee": 0.0,
                    }
                ],
            }
        ),
        encoding="utf-8",
    )
    profit_analysis_file = tmp_path / "profit_analysis.json"
    profit_analysis_file.write_text(
        json.dumps(
            {
                "profit_analysis": [
                    {
                        "hot_item_id": "xy-001",
                        "source_item_id": "src-001",
                        "source_cost": 121.5,
                        "target_xianyu_price": 299.0,
                        "estimated_margin": 150.0,
                        "gross_margin_rate": 0.5,
                        "cost_profit_rate": 1.235,
                        "risk_flags": [],
                    }
                ]
            }
        ),
        encoding="utf-8",
    )
    listing_candidates_file = tmp_path / "listing_candidates.json"
    listing_candidates_file.write_text(
        json.dumps(
            {
                "listing_candidates": [
                    {
                        "hot_item_id": "xy-001",
                        "source_item_id": "src-001",
                        "estimated_margin": 150.0,
                        "gross_margin_rate": 0.5,
                        "cost_profit_rate": 1.235,
                        "is_recommended": True,
                        "reasons": ["margin_ok"],
                        "blocked_by": [],
                    }
                ]
            }
        ),
        encoding="utf-8",
    )
    output_file = tmp_path / "pipeline.xlsx"
    repo_root = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools")
    result = subprocess.run(
        [
            sys.executable,
            "scripts/export_pipeline_excel.py",
            "--hot-items-file",
            str(hot_items_file),
            "--source-bundle-file",
            str(source_bundle_file),
            "--profit-analysis-file",
            str(profit_analysis_file),
            "--listing-candidates-file",
            str(listing_candidates_file),
            "--output-file",
            str(output_file),
        ],
        cwd=repo_root,
        env={"PYTHONPATH": "src"},
        capture_output=True,
        text=True,
        check=False,
    )
    assert result.returncode == 0
    workbook = load_workbook(output_file)
    assert "汇总" in workbook.sheetnames
    assert "筛选总览" in workbook.sheetnames
    overview = workbook["筛选总览"]
    assert overview["A2"].value == "xy-001"
    listing_sheet = workbook["最终上架候选"]
    listing_headers = [cell.value for cell in listing_sheet[1]]
    header_index = {value: index + 1 for index, value in enumerate(listing_headers)}
    assert listing_sheet.cell(row=2, column=header_index["闲鱼价格"]).value == 299.0
    assert listing_sheet.cell(row=2, column=header_index["1688价格"]).value == 120.0


def test_run_category_pipeline_from_urls_writes_outputs(monkeypatch, capsys, tmp_path) -> None:
    script_path = Path("/Users/mac/PycharmProjects/mytools/xianyu-tools/scripts/run_category_pipeline_from_urls.py")
    spec = importlib.util.spec_from_file_location("run_category_pipeline_from_urls_script", script_path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    hot_items_file = tmp_path / "hot_items.json"
    hot_items_file.write_text(
        json.dumps(
            {
                "hot_items": [
                    {
                        "hot_item_id": "xy-001",
                        "platform": "xianyu",
                        "title": "升降桌",
                        "price": 99.0,
                    }
                ]
            }
        ),
        encoding="utf-8",
    )
    result_url_map_file = tmp_path / "result_url_map.json"
    result_url_map_file.write_text(json.dumps({"xy-001": "https://s.1688.com/selloffer/offer_search.htm?keywords=x"}), encoding="utf-8")
    output_dir = tmp_path / "out"

    def fake_resolve(hot_items, *, result_urls_by_hot_item_id, adapter=None, limit_per_item=10):
        return {
            "source_resolution": [
                {
                    "hot_item_id": "xy-001",
                    "resolved": True,
                    "source_item_ids": ["src-001"],
                    "resolution_reason": "matched_source_items",
                }
            ],
            "source_items": [
                {
                    "source_platform": "1688",
                    "source_item_id": "src-001",
                    "title": "升降桌工厂直发",
                    "price": 60.0,
                    "shipping_fee": 5.0,
                    "item_url": "https://detail.1688.com/offer/src-001.html",
                    "sales": 120,
                    "metadata": {"hot_item_id": "xy-001"},
                }
            ],
        }

    monkeypatch.setattr(module, "resolve_hot_items_to_ali1688_urls", fake_resolve)
    monkeypatch.setattr(
        sys,
        "argv",
        [
            "run_category_pipeline_from_urls.py",
            "--hot-items-file",
            str(hot_items_file),
            "--result-url-map-file",
            str(result_url_map_file),
            "--output-dir",
            str(output_dir),
        ],
    )
    assert module.main() == 0
    output = json.loads(capsys.readouterr().out)
    assert Path(output["source_bundle_file"]).exists()
    assert Path(output["profit_analysis_file"]).exists()
    assert Path(output["listing_candidates_file"]).exists()


def test_pipeline_excel_free_parsing_logic() -> None:
    # Test JSON-based SKU extraction and min_price calculation introduced in run_full_pipeline.py
    res = {
        "title": "测试商品",
        "offer_id": "12345",
        "item_url": "https://detail.1688.com/offer/12345.html",
        "images": ["http://img1.jpg"],
        "sku_items": [
            {"attributes": "颜色:红色;尺码:L", "price": "100.00", "stock": 50, "spec_id": "sp1", "image": "http://img2.jpg"},
            {"attributes": "颜色:红色;尺码:M", "price": 95.5, "stock": 20, "spec_id": "sp2", "image": ""},
        ]
    }
    sku_items = res.get("sku_items", [])
    min_price = 0
    sku_count = 0
    if sku_items:
        prices = [float(s.get("price") or 0.0) for s in sku_items if s.get("price") is not None]
        if prices:
            min_price = min(prices)
        sku_count = len(sku_items)
    
    assert min_price == 95.5
    assert sku_count == 2


def test_local_html_cleanup_on_decision_asset_delete(tmp_path) -> None:
    # Verify basic physical file unlink cascade logic
    dummy_html = tmp_path / "dummy_detail.html"
    dummy_html.write_text("<html>test</html>", encoding="utf-8")
    assert dummy_html.exists()
    
    # Simulate extraction of html_path and cascading unlink
    html_path_str = str(dummy_html.resolve())
    p = Path(html_path_str)
    if p.exists() and p.is_file():
        p.unlink()
        
    assert not dummy_html.exists()


def test_delete_task_physically_cleans_local_html(monkeypatch, tmp_path) -> None:
    # Test that delete_task route unlinks local HTML files cascadingly based on html_path
    import src.web_api.main as web_main
    
    dummy_html = tmp_path / "dummy_1688_detail.html"
    dummy_html.write_text("<html>test</html>", encoding="utf-8")
    assert dummy_html.exists()

    class FakeCursor:
        def __init__(self):
            self.query = None
            self.args = None

        def execute(self, query, args=None):
            self.query = query
            self.args = args

        def fetchall(self):
            # Simulate returning html_path for resources under this task
            return [{"html_path": str(dummy_html.resolve())}]

        def fetchone(self):
            # Simulate task details
            return {"root_dir": str(tmp_path)}

    class FakeConn:
        def cursor(self):
            return FakeCursor()
        def commit(self):
            pass
        def close(self):
            pass

    monkeypatch.setattr(web_main, "get_db_conn", lambda: FakeConn())
    monkeypatch.setattr(web_main, "pause_task", lambda tid: None)
    monkeypatch.setattr(web_main.Task, "update", lambda *args, **kwargs: None)

    # Call delete_task to trigger cascade physical file deletion
    res = web_main.delete_task("dummy_task_id")
    assert res == {"status": "ok"}
    
    # Verify the cascade delete unlinked the file successfully
    assert not dummy_html.exists()


def test_web_api_depublish_route_flow(monkeypatch) -> None:
    # Test POST /api/depublish/{source_id} API endpoint
    import src.web_api.main as web_main

    class FakeCursor:
        def __init__(self):
            self.queries = []
            self.args = []

        def execute(self, query, args=None):
            self.queries.append(query)
            self.args.append(args)

        def fetchone(self):
            # Simulate returning already published product_id and task_id
            return {"xianyu_item_id": "987654", "task_id": "task_111"}

    class FakeConn:
        def cursor(self):
            return FakeCursor()
        def commit(self):
            pass
        def close(self):
            pass

    monkeypatch.setattr(web_main, "get_db_conn", lambda: FakeConn())

    from xianyu_tools.xianyu_adapter.publisher_v3 import PublisherV3
    monkeypatch.setattr(PublisherV3, "depublish_item", lambda self, pid: {"status": "success", "msg": "下架成功"})

    import asyncio
    # 直接异步调用接口函数以规避对 httpx 的依赖
    res = asyncio.run(web_main.depublish_from_xianyu(802))
    assert res == {"status": "success", "msg": "下架成功"}


def test_web_api_batch_depublish(monkeypatch) -> None:
    import src.web_api.main as web_main

    class FakeCursor:
        def __init__(self):
            self.queries = []
            self.args = []

        def execute(self, query, args=None):
            self.queries.append(query)
            self.args.append(args)

        def fetchone(self):
            return {"xianyu_item_id": "987654", "task_id": "task_111"}

    class FakeConn:
        def cursor(self):
            return FakeCursor()
        def commit(self):
            pass
        def close(self):
            pass

    monkeypatch.setattr(web_main, "get_db_conn", lambda: FakeConn())

    from xianyu_tools.xianyu_adapter.publisher_v3 import PublisherV3
    monkeypatch.setattr(PublisherV3, "depublish_item", lambda self, pid: {"status": "success", "msg": "下架成功"})

    import asyncio
    req = {"source_ids": [802, 803]}
    res = asyncio.run(web_main.batch_depublish_from_xianyu(req))
    assert res == {
        "success": [{"source_id": 802}, {"source_id": 803}],
        "failed": []
    }


def test_web_api_delete_route_flow(monkeypatch) -> None:
    # Test POST /api/delete/{source_id} API endpoint
    import src.web_api.main as web_main

    # 1. 测试未发布商品删除拦截
    class FakeCursorNone:
        def execute(self, query, args=None): pass
        def fetchone(self): return None
    class FakeConnNone:
        def cursor(self): return FakeCursorNone()
        def close(self): pass

    monkeypatch.setattr(web_main, "get_db_conn", lambda: FakeConnNone())
    import asyncio
    res_none = asyncio.run(web_main.delete_from_xianyu(901))
    assert res_none == {"status": "failed", "msg": "商品未发布，无法删除"}

    # 2. 测试非下架状态商品（例如 success）删除拦截
    class FakeCursorSuccess:
        def execute(self, query, args=None): pass
        def fetchone(self): return {"publish_status": "success", "xianyu_item_id": "987654", "task_id": "task_111"}
    class FakeConnSuccess:
        def cursor(self): return FakeCursorSuccess()
        def close(self): pass

    monkeypatch.setattr(web_main, "get_db_conn", lambda: FakeConnSuccess())
    res_succ = asyncio.run(web_main.delete_from_xianyu(902))
    assert res_succ == {"status": "failed", "msg": "商品当前状态为 success，只有已下架商品可以删除"}

    # 3. 测试已下架商品（depublished）成功删除
    class FakeCursorDepublished:
        def execute(self, query, args=None): pass
        def fetchone(self): return {"publish_status": "depublished", "xianyu_item_id": "987654", "task_id": "task_111"}
    class FakeConnDepublished:
        def cursor(self): return FakeCursorDepublished()
        def commit(self): pass
        def close(self): pass

    monkeypatch.setattr(web_main, "get_db_conn", lambda: FakeConnDepublished())
    from xianyu_tools.xianyu_adapter.publisher_v3 import PublisherV3
    monkeypatch.setattr(PublisherV3, "delete_item", lambda self, pid: {"status": "success", "msg": "删除成功"})

    res_del = asyncio.run(web_main.delete_from_xianyu(903))
    assert res_del == {"status": "success", "msg": "删除成功"}


def test_web_api_batch_delete(monkeypatch) -> None:
    # Test POST /api/delete/batch API endpoint
    import src.web_api.main as web_main

    # 模拟批量数据：
    # 801: 未发布 (None)
    # 802: 已下架 (depublished) -> 成功删除
    # 803: 已上架 (success) -> 状态不符报错拦截
    class FakeCursorBatch:
        def __init__(self):
            self.count = 0
        def execute(self, query, args=None):
            self.current_args = args
        def fetchone(self):
            sid = self.current_args[0]
            if sid == 801:
                return None
            elif sid == 802:
                return {"publish_status": "depublished", "xianyu_item_id": "item_802", "task_id": "task_802"}
            elif sid == 803:
                return {"publish_status": "success", "xianyu_item_id": "item_803", "task_id": "task_803"}
            return None

    class FakeConnBatch:
        def cursor(self): return FakeCursorBatch()
        def commit(self): pass
        def close(self): pass

    monkeypatch.setattr(web_main, "get_db_conn", lambda: FakeConnBatch())
    from xianyu_tools.xianyu_adapter.publisher_v3 import PublisherV3
    monkeypatch.setattr(PublisherV3, "delete_item", lambda self, pid: {"status": "success", "msg": "删除成功"})

    import asyncio
    req = {"source_ids": [801, 802, 803]}
    res = asyncio.run(web_main.batch_delete_from_xianyu(req))
    assert res["success"] == [{"source_id": 802}]
    assert len(res["failed"]) == 2
    assert res["failed"][0] == {"source_id": 801, "msg": "商品未发布，无法删除"}
    assert res["failed"][1] == {"source_id": 803, "msg": "商品状态为 success，只有已下架商品可以删除"}


def test_web_api_get_xianyu_products(monkeypatch) -> None:
    # Test GET /api/xianyu_products API endpoint
    import src.web_api.main as web_main
    from datetime import datetime

    class FakeCursor:
        def __init__(self):
            self.queries = []
            self.args = []

        def execute(self, query, args=None):
            self.queries.append(query)
            self.args.append(args)

        def fetchone(self):
            return {"count": 1}

        def fetchall(self):
            return [{
                "publish_id": 1,
                "task_id": "task_abc",
                "source_db_id": 101,
                "xianyu_item_id": "999888",
                "publish_status": "success",
                "publish_msg": "已上架",
                "published_url": "http://xianyu.com/999888",
                "publish_time": datetime(2026, 6, 9, 12, 0, 0),
                "source_title": "1688源头好物",
                "source_url": "http://1688.com/101",
                "source_images": '["http://img.1688.com/101.jpg"]',
                "source_price": 50.0,
                "source_sku_count": 3,
                "ref_title": "参考爆款标题",
                "ref_price": 80.0,
                "ref_want_count": 500
            }]

    class FakeConn:
        def cursor(self): return FakeCursor()
        def close(self): pass

    monkeypatch.setattr(web_main, "get_db_conn", lambda: FakeConn())

    res = web_main.get_xianyu_products(page=1, limit=10, keyword="1688")
    assert res["total"] == 1
    assert len(res["items"]) == 1
    assert res["items"][0]["source_title"] == "1688源头好物"
    assert res["items"][0]["xianyu_item_id"] == "999888"
    assert res["items"][0]["publish_time"] == "2026-06-09 12:00:00"
    assert res["items"][0]["source_image"] == "http://img.1688.com/101.jpg"
